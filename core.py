# -*- coding: utf-8 -*-
"""Боевая логика Турналога (без GUI): обработка, история, сопоставления."""
import io
import logging
import os
import re
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from agents import default_agents, match_agent, norm_key
from schemas import DNA_DATA_SCHEMA, TCH_DATA_SCHEMA, TCH_ROOMS_SCHEMA
from tax import TaxParams, parse_price

log = logging.getLogger('turnalog')

FILL_GUEST = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
FILL_DATE = PatternFill(start_color="FFCC99", end_color="FFCC99", fill_type="solid")
FILL_PRICE = PatternFill(start_color="99CCFF", end_color="99CCFF", fill_type="solid")
FILL_ROOM = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
FILL_DUP = PatternFill(start_color="CC99FF", end_color="CC99FF", fill_type="solid")
FILL_TOTALS = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
DATE_FMT = 'DD.MM.YYYY HH:MM'
ROOM_LABEL_TEXT_PREFIX = '__TL_ROOM_LABEL__:'

def norm_name(s):
    """Нормализация ФИО: регистр, ё->е, лишние пробелы."""
    return re.sub(r'\s+', ' ', str(s).lower().replace('ё', 'е')).strip()


def clean_guest_names(text):
    if pd.isna(text):
        return ""
    return re.sub(r'\d+\.\s*', '', str(text)).strip()


_EXTRA_AMOUNT_KEYS = {
    'допуслуги', 'суммадопуслуг', 'стоимостьдопуслуг',
    'суммадополнительныхуслуг', 'стоимостьдополнительныхуслуг',
    'extras', 'extrasamount', 'additionalservicesamount',
}


def _is_extra_amount_header(label):
    return norm_key(label) in _EXTRA_AMOUNT_KEYS


def _is_service_name_header(label):
    return norm_key(label) in {
        'наименованиедопуслуг', 'названиедопуслуг',
        'описаниедопуслуг', 'additionalservicenames',
    }


def _is_extension_amount_header(label):
    return norm_key(label) in {
        'суммапродленияномера', 'продлениеномерасумма',
    }


def _service_name_parts(value):
    for part in str(value or '').split(';'):
        match = re.match(r'^(.*?)\s+[—–-]\s+(.*)$', part.strip())
        if match:
            yield match.group(1).strip(), match.group(2).strip()
        else:
            yield part.strip(), None


def _contains_named_service(value, service_name):
    target = norm_key(service_name)
    return any(norm_key(name) == target
               for name, _amount in _service_name_parts(value))


def _named_service_amount(value, service_name):
    """Извлекает суммы из колонок вида «Услуга — 1 000,00 ₽; ...»."""
    target = norm_key(service_name)
    total = 0.0
    for name, amount_text in _service_name_parts(value):
        if amount_text is not None and norm_key(name) == target:
            amount = parse_price(amount_text)
            if amount is not None:
                total += amount
    return round(total, 2)


class WorkbookReadError(Exception):
    """Файл Excel не читается: понятное пользователю описание + что делать."""


# Кэш починок: abspath -> (mtime, fixed_path, notes). Временные копии лежат
# в системном temp, переиспользуются, пока исходник не изменился, и не
# удаляются (их чистит сама ОС) — repair может вызываться многократно.
_REPAIR_CACHE = {}

_ILLEGAL_XML_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _fix_short_rgb(text):
    def fix_rgb(m):
        v = m.group(1)
        return f'rgb="{v.rjust(8, "0")}"' if len(v) < 8 else m.group(0)
    return re.sub(r'rgb="([0-9A-Fa-f]{1,8})"', fix_rgb, text)


def _build_styles(numfmts, numfmt_ids):
    """Минимальный валидный styles.xml с заданными форматами чисел.

    Чтение ЗНАЧЕНИЙ (включая распознавание дат) зависит только от
    соответствия xf -> numFmtId, поэтому точность данных сохраняется.
    numfmts: [(id>=164, raw_formatCode)], numfmt_ids: [numFmtId по порядку xf].
    """
    numfmts_xml = ''.join(
        f'<numFmt numFmtId="{i}" formatCode="{c}"/>' for i, c in numfmts)
    cellxfs_xml = ''.join(
        f'<xf numFmtId="{n}" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        for n in numfmt_ids)
    out = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<numFmts count="{len(numfmts)}">{numfmts_xml}</numFmts>'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/>'
        '</border></borders>'
        '<cellStyleXfs count="1">'
        '<xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        f'<cellXfs count="{len(numfmt_ids)}">{cellxfs_xml}</cellXfs>'
        '<cellStyles count="1">'
        '<cellStyle name="Обычный" xfId="0" builtinId="0"/></cellStyles>'
        '</styleSheet>')
    ET.fromstring(out)
    return out.encode('utf-8')


def _max_sheet_xf(work):
    """Максимальный индекс стиля s="N" в ячейках листов (для tier 4–5)."""
    top = 0
    for n, data in work.items():
        if n.startswith('xl/worksheets/') and n.endswith('.xml'):
            try:
                txt = data.decode('utf-8')
            except UnicodeDecodeError:
                continue
            for mm in re.finditer(r'<c\b[^>]*?\ss="(\d+)"', txt):
                top = max(top, int(mm.group(1)))
    return min(top, 4096)


def _salvage_parts(styles_text):
    """Вытянуть из битого styles.xml (numfmts, numfmt_ids) или None."""
    m = re.search(r'<cellXfs[^>]*>(.*?)</cellXfs>', styles_text, re.S)
    if not m:
        return None
    xf_attrs = re.findall(r'<xf\b([^>]*?)/?>', m.group(1))
    if not xf_attrs:
        return None
    numfmt_ids = []
    for a in xf_attrs:
        mm = re.search(r'numFmtId="(\d+)"', a)
        numfmt_ids.append(int(mm.group(1)) if mm else 0)
    numfmts = []
    for fm in re.finditer(r'<numFmt\b([^>]*?)/?>', styles_text):
        idm = re.search(r'numFmtId="(\d+)"', fm.group(1))
        cdm = re.search(r'formatCode="([^"]*)"', fm.group(1))
        # встроенные форматы (<164) openpyxl знает сам, спасаем только свои
        if idm and cdm and int(idm.group(1)) >= 164:
            numfmts.append((int(idm.group(1)), cdm.group(1)))
    return numfmts, numfmt_ids


def repair_workbook(path):
    """Вернуть (читаемый_путь, notes) для битого xlsx.

    Уровни (каждый следующий включает предыдущие правки):
    1. короткий rgb в цветах; 2. чистка styles.xml от управляющих символов;
    3. чистка всех xml-частей; 4. пересборка styles.xml с сохранением
    форматов чисел/дат; 5. минимальные стили (даты могли пострадать).
    Уровни 1–3 без потерь и молчаливые (только в лог), 4–5 дают заметку
    пользователю. Бросает WorkbookReadError, если файл безнадёжен.
    """
    apath = os.path.abspath(path)
    try:
        mtime = os.path.getmtime(apath)
    except OSError:
        raise WorkbookReadError(
            f'Файл «{os.path.basename(path)}» не найден или недоступен.')
    cached = _REPAIR_CACHE.get(apath)
    if cached and cached[0] == mtime and os.path.exists(cached[1]):
        return cached[1], cached[2]
    if not zipfile.is_zipfile(apath):
        raise WorkbookReadError(
            f'Файл «{os.path.basename(path)}» не является книгой Excel (.xlsx): '
            'внутри нет zip-архива. Возможно, это старый .xls, HTML-выгрузка или CSV '
            'с расширением .xlsx. Откройте файл в Excel и сохраните как '
            '«Книга Excel (*.xlsx)», затем повторите.')
    try:
        with zipfile.ZipFile(apath) as z:
            items = {n: z.read(n) for n in z.namelist()}
    except Exception as e:
        raise WorkbookReadError(
            f'Не удалось открыть «{os.path.basename(path)}» как архив: {e}')
    base = os.path.basename(path)
    first_err = None

    def probe(items_map):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            for name, data in items_map.items():
                z.writestr(name, data)
        raw = buf.getvalue()
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        wb.close()
        return raw

    def finish(raw, notes):
        fd, tmp = tempfile.mkstemp(prefix='Turnalog_fix_', suffix='.xlsx')
        with os.fdopen(fd, 'wb') as f:
            f.write(raw)
        res = (tmp, notes)
        _REPAIR_CACHE[apath] = (mtime, tmp, notes)
        return res

    def attempt(work):
        nonlocal first_err
        try:
            return probe(work)
        except Exception as e:
            first_err = first_err or e
            return None

    sname = next((n for n in items if n.lower().endswith('styles.xml')), None)
    work = dict(items)

    # Tier 1: короткий rgb (старый кейс PMS-выгрузок).
    if sname:
        try:
            work[sname] = _fix_short_rgb(work[sname].decode('utf-8')).encode('utf-8')
        except UnicodeDecodeError:
            pass
    raw = attempt(work)
    if raw is not None:
        log.info('Починка %s: короткий rgb', base)
        return finish(raw, [])

    # Tier 2: запрещённые управляющие символы в styles.xml.
    if sname:
        try:
            txt = work[sname].decode('utf-8')
        except UnicodeDecodeError:
            txt = None
        if txt is not None:
            work[sname] = _ILLEGAL_XML_CHARS.sub('', txt).encode('utf-8')
            raw = attempt(work)
            if raw is not None:
                log.info('Починка %s: чистка styles.xml', base)
                return finish(raw, [])

    # Tier 3: чистка всех xml-частей архива.
    for n in list(work):
        if n.lower().endswith('.xml'):
            try:
                work[n] = _ILLEGAL_XML_CHARS.sub(
                    '', work[n].decode('utf-8')).encode('utf-8')
            except UnicodeDecodeError:
                pass
    raw = attempt(work)
    if raw is not None:
        log.info('Починка %s: чистка всех xml', base)
        return finish(raw, [])

    # Tier 4: пересборка styles.xml с сохранением форматов чисел/дат.
    if sname:
        try:
            parts = _salvage_parts(items[sname].decode('utf-8'))
        except (UnicodeDecodeError, ValueError):
            parts = None
        if parts is not None:
            numfmts, ids = parts
            need = _max_sheet_xf(work) + 1
            if len(ids) < need:  # страховка от рассинхрона индексов
                ids = ids + [0] * (need - len(ids))
            try:
                work[sname] = _build_styles(numfmts, ids)
            except Exception as e:
                first_err = first_err or e
                work[sname] = items[sname]
                raw = None
            else:
                raw = attempt(work)
            if raw is not None:
                log.warning('Починка %s: стили пересобраны', base)
                return finish(raw, [
                    f'Файл «{base}» открыт после восстановления '
                    '(стили были повреждены). Даты и числа сохранены; '
                    'рекомендуем пересохранить файл в Excel.'])

    # Tier 5: минимальные стили — даты могли прочитаться как числа.
    if sname:
        need = _max_sheet_xf(work) + 1
        work[sname] = _build_styles([], [0] * need)
        raw = attempt(work)
        if raw is not None:
            log.warning('Починка %s: минимальные стили', base)
            return finish(raw, [
                f'Файл «{base}» СИЛЬНО повреждён: открыт с потерей стилей, '
                'даты могли прочитаться как числа — ПРОВЕРЬТЕ результат. Обязательно '
                'пересохраните файл в Excel (Файл → Сохранить как) и обработайте заново.'])

    raise WorkbookReadError(
        f'Файл «{base}» повреждён и не открывается. Откройте его в Excel — обычно Excel '
        'предлагает восстановить такой файл, — сохраните заново (Файл → Сохранить как → '
        f'Книга Excel *.xlsx) и повторите. Подробности: {first_err}')


def repair_notes_for(paths):
    """Заметки о восстановлении файлов (для финального диалога)."""
    out = []
    for p in paths:
        try:
            apath = os.path.abspath(p)
        except Exception:
            continue
        cached = _REPAIR_CACHE.get(apath)
        if cached:
            for n in cached[2]:
                if n not in out:
                    out.append(n)
    return out


def excel_serial_to_dt(v):
    """Excel-сериал (float при сломанных стилях) -> datetime.

    Нечисла возвращаются как есть; числа вне диапазона 2015–2035
    считаются мусором и дают None.
    """
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return v
    try:
        dt = from_excel(float(v))
    except Exception:
        return None
    return dt if 2015 <= dt.year <= 2035 else None


def parse_usd_dt(v):
    """Дата ячейки усадебного отчёта: datetime / Excel-сериал / текст."""
    if isinstance(v, datetime):
        return v
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return excel_serial_to_dt(v)
    s = str(v).strip()
    for fmt in ('%d.%m.%Y %H:%M', '%d.%m.%Y %H:%M:%S', '%d.%m.%Y'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def read_table(path, nrows=None):
    """Чтение .xlsx/.csv с автоподбором кодировки для CSV."""
    if str(path).lower().endswith('.csv'):
        last_err = None
        for enc in ('utf-8-sig', 'utf-8', 'cp1251', 'cp866'):
            try:
                return pd.read_csv(path, encoding=enc, nrows=nrows)
            except UnicodeDecodeError as e:
                last_err = e
        raise ValueError(f"Не удалось определить кодировку CSV: {last_err}")
    try:
        return pd.read_excel(path, nrows=nrows)
    except WorkbookReadError:
        raise
    except Exception:
        fixed, _notes = repair_workbook(path)
        return pd.read_excel(fixed, nrows=nrows)


def read_headers(path):
    return [str(c) for c in read_table(path, nrows=0).columns]


def guess_mapping(headers, schema):
    """Автоподбор колонок файла под поля приложения (точное, затем по подстроке)."""
    nh = {h: norm_key(h) for h in headers}
    taken, mapping = set(), {}
    for f in schema:  # pass 1: точное совпадение синонима
        for h, k in nh.items():
            if h in taken:
                continue
            if any(norm_key(a) == k for a in f['aliases']):
                mapping[f['key']] = h
                taken.add(h)
                break
    for f in schema:  # pass 2: подстрока (самый длинный синоним)
        if f['key'] in mapping:
            continue
        best, bestlen = None, 0
        for h, k in nh.items():
            if h in taken or not k:
                continue
            for a in f['aliases']:
                na = norm_key(a)
                if na and (na in k or k in na) and len(na) > bestlen:
                    best, bestlen = h, len(na)
        if best:
            mapping[f['key']] = best
            taken.add(best)
    return mapping


def mapping_valid(mapping, headers, schema):
    if not mapping:
        return False
    hs = set(headers)
    return all(mapping.get(f['key']) in hs for f in schema if f['required'])


def complete_mapping(mapping, headers, schema):
    """Дополняет старую сохранённую схему новыми необязательными полями.

    Явно сохранённое ``None`` не перезаписывается: автоподбор применяется
    только к ключам, которых ещё не было в конфигурации.
    """
    result = dict(mapping or {})
    guessed = guess_mapping(headers, schema)
    for field in schema:
        if field['key'] not in result:
            result[field['key']] = guessed.get(field['key'])
    return result


def open_in_system(path):
    try:
        if os.name == 'nt':
            os.startfile(path)
        elif sys.platform == 'darwin':
            subprocess.Popen(['open', path])
        else:
            subprocess.Popen(['xdg-open', path])
    except Exception as e:
        log.exception("open_in_system")
        try:
            from tkinter import messagebox  # лениво: core импортируется и без GUI
            messagebox.showerror("Ошибка", f"Не удалось открыть файл:\n{e}")
        except Exception:
            pass


def pack_originals(paths):
    if len(paths) == 1:
        p = paths[0]
        return os.path.basename(p), open(p, 'rb').read()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            z.write(p, os.path.basename(p))
    return 'исходные_файлы.zip', buf.getvalue()



def process_tchaikovsky( data_path, rooms_path, output_path,
                        map_data=None, map_rooms=None,
                        extra_data=None, dup_checker=None, tax=None,
                        agents=None, source_map=None):
    """Возвращает (errors_summary, bookings)."""
    tax_cfg = dict(tax or {})
    # Для Чайковского сохраняем включение всех допуслуг по умолчанию.
    tax_cfg.setdefault('include_extras_in_tax', True)
    tp = TaxParams(tax_cfg)
    log.info("Чайковский: %s / %s", data_path, rooms_path)
    df_data = read_table(data_path)
    df_rooms = read_table(rooms_path)

    headers_d = [str(c) for c in df_data.columns]
    headers_r = [str(c) for c in df_rooms.columns]
    map_data = map_data or guess_mapping(headers_d, TCH_DATA_SCHEMA)
    map_rooms = map_rooms or guess_mapping(headers_r, TCH_ROOMS_SCHEMA)
    miss = [f['label'] for f in TCH_DATA_SCHEMA if f['required'] and not map_data.get(f['key'])]
    if miss:
        raise ValueError(f"Не сопоставлены колонки: {', '.join(miss)}")
    has_room_name = bool(map_rooms.get('fio')
                         and map_rooms['fio'] in df_rooms.columns)
    has_room_booking = bool(map_rooms.get('bron')
                            and map_rooms['bron'] in df_rooms.columns)
    if (not map_rooms.get('room') or map_rooms['room'] not in df_rooms.columns
            or not (has_room_name or has_room_booking)):
        raise ValueError(
            'В файле «Номера» сопоставьте № комнаты и хотя бы один ключ: ФИО '
            'или № брони')
    extra_data = [c for c in (extra_data or []) if c in df_data.columns]
    extra_amount_col = next(
        (c for c in df_data.columns if _is_extra_amount_header(c)), None)
    service_name_col = next(
        (c for c in df_data.columns if _is_service_name_header(c)), None)
    extension_amount_col = next(
        (c for c in df_data.columns if _is_extension_amount_header(c)), None)
    for special_col in (extra_amount_col, service_name_col, extension_amount_col):
        if special_col is not None and special_col not in extra_data:
            extra_data.append(special_col)
    report_extra_data = [
        c for c in extra_data
        if c not in (extra_amount_col, service_name_col, extension_amount_col)
    ]

    room_mapping = {}
    room_mapping_by_bron = {}
    booking_date_candidates = {}
    room_file_has_booking_ids = False
    has_corpus_mapping = bool(
        map_rooms.get('corpus') and map_rooms['corpus'] in df_rooms.columns)
    has_direct_room_mapping = bool(
        map_data.get('room') and map_data['room'] in df_data.columns)
    has_direct_corpus_mapping = bool(
        map_data.get('corpus') and map_data['corpus'] in df_data.columns)

    def blank(value):
        if value is None:
            return True
        try:
            if pd.isna(value):
                return True
        except (TypeError, ValueError):
            pass
        return not str(value).strip()

    def clean_room(value):
        if blank(value):
            return None
        text = str(value).strip()
        if text.startswith(ROOM_LABEL_TEXT_PREFIX):
            return text[len(ROOM_LABEL_TEXT_PREFIX):]
        try:
            number = float(text)
            if number.is_integer():
                return text if re.fullmatch(r'0\d+', text) else int(number)
        except (ValueError, TypeError, OverflowError):
            pass
        return value

    def booking_key(value):
        if blank(value):
            return ''
        text = str(value).strip().casefold()
        if re.fullmatch(r'\d+\.0+', text):
            text = text.split('.', 1)[0]
        return re.sub(r'\s+', '', text)

    def stay_date_key(value):
        if blank(value):
            return None
        try:
            parsed = pd.to_datetime(
                excel_serial_to_dt(value), dayfirst=True, errors='coerce')
        except (TypeError, ValueError, OverflowError):
            return None
        return None if pd.isna(parsed) else parsed.date()

    guest_candidates = {}
    booking_candidates = {}
    booking_col = map_rooms.get('bron')
    room_checkin_col = map_rooms.get('checkin')
    room_checkout_col = map_rooms.get('checkout')
    for _, rr in df_rooms.iterrows():
        fio_col = map_rooms.get('fio')
        fio = rr[fio_col] if fio_col and fio_col in df_rooms.columns else None
        guest_key = '' if blank(fio) else norm_name(fio)
        room = clean_room(rr[map_rooms['room']])
        corpus_col = map_rooms.get('corpus')
        corpus = (rr[corpus_col] if corpus_col and corpus_col in df_rooms.columns
                  else None)
        corpus = None if blank(corpus) else corpus
        info = (room, corpus) if room is not None else None
        if guest_key and info is not None:
            guest_candidates.setdefault(guest_key, []).append(info)

        bron = (rr[booking_col]
                if booking_col and booking_col in df_rooms.columns else None)
        bron_key = booking_key(bron)
        if bron_key:
            room_file_has_booking_ids = True
            booking_candidates.setdefault(bron_key, []).append(info)
            room_checkin = (stay_date_key(rr[room_checkin_col])
                            if room_checkin_col in df_rooms.columns else None)
            room_checkout = (stay_date_key(rr[room_checkout_col])
                             if room_checkout_col in df_rooms.columns else None)
            if room_checkin is not None and room_checkout is not None:
                date_key = (bron_key, room_checkin, room_checkout)
                booking_date_candidates.setdefault(date_key, []).append(info)

    # Only use a guest/booking key when it resolves to one unambiguous room.
    for key, candidates in guest_candidates.items():
        first = candidates[0]
        if all(candidate == first for candidate in candidates[1:]):
            room_mapping[key] = first
    for key, candidates in booking_candidates.items():
        if candidates and candidates[0] is not None and all(
                candidate == candidates[0] for candidate in candidates[1:]):
            room_mapping_by_bron[key] = candidates[0]

    def find_room_info(guest_string, booking_value=None,
                       checkin_value=None, checkout_value=None):
        bron_key = booking_key(booking_value)
        if bron_key:
            if bron_key in room_mapping_by_bron:
                return room_mapping_by_bron[bron_key]
            # For a multi-room booking, an exact booking + stay-date match can
            # distinguish room changes without guessing from the room type.
            checkin_key = stay_date_key(checkin_value)
            checkout_key = stay_date_key(checkout_value)
            if checkin_key is not None and checkout_key is not None:
                date_candidates = booking_date_candidates.get(
                    (bron_key, checkin_key, checkout_key), [])
                if date_candidates:
                    first = date_candidates[0]
                    if first is not None and all(
                            candidate == first
                            for candidate in date_candidates[1:]):
                        return first
            # The room file carries booking IDs, so an absent/blank match is
            # authoritative; don't borrow another booking's room by guest name.
            if room_file_has_booking_ids:
                candidates = booking_candidates.get(bron_key)
                if not candidates or any(candidate is None
                                         for candidate in candidates):
                    return None, None
                # A multi-room booking may have several different rooms; then
                # fall through to the guest key, which must also be unique.
        if not guest_string:
            return None, None
        for guest in re.split(r'[,\n;]+', guest_string):
            key = norm_name(guest)
            if key in room_mapping:
                return room_mapping[key]
        return None, None

    def get(row, key):
        col = map_data.get(key)
        return row[col] if col and col in row.index else None

    records = []
    for _, row in df_data.iterrows():
        guest = clean_guest_names(get(row, 'guest'))

        checkin = pd.to_datetime(excel_serial_to_dt(get(row, 'checkin')),
                                 dayfirst=True, errors='coerce')
        checkout = pd.to_datetime(excel_serial_to_dt(get(row, 'checkout')),
                                  dayfirst=True, errors='coerce')
        checkin = None if pd.isna(checkin) else checkin.to_pydatetime()
        checkout = None if pd.isna(checkout) else checkout.to_pydatetime()

        nights_v = get(row, 'nights')
        nights = None
        if nights_v is not None and not pd.isna(nights_v):
            try:
                nights = int(float(nights_v))
            except (ValueError, TypeError):
                nights = None
        if nights is None and checkin and checkout:
            nights = max(1, (checkout.date() - checkin.date()).days)

        category = get(row, 'category')
        category = '' if category is None or pd.isna(category) else str(category).strip()
        num_rooms = max(1, len([x for x in category.split('\n') if x.strip()])) if category else 1

        price = parse_price(get(row, 'price'))

        bron_v = get(row, 'bron')
        bron = '' if bron_v is None or pd.isna(bron_v) else str(bron_v).strip()

        source_v = get(row, 'source')
        source = '' if source_v is None or pd.isna(source_v) else str(source_v).strip()

        sm = (source_map or {}).get(source)
        if sm:
            commission, agent = float(sm[0]), sm[1]
        else:
            commission, agent = match_agent(
                source, agents or default_agents(),
                date=checkin.date() if checkin else None, hotel='tch')

        extras_amount = parse_price(row[extra_amount_col]) if extra_amount_col else 0.0
        extras_amount = 0.0 if extras_amount is None else round(extras_amount, 2)
        names_value = row[service_name_col] if service_name_col else ''
        service_names = ('' if names_value is None or pd.isna(names_value)
                         else str(names_value).strip())
        extra_vals = []
        for c in report_extra_data:
            value = row[c]
            if pd.isna(value):
                value = ''
            extra_vals.append(value)

        if has_direct_room_mapping:
            room = clean_room(get(row, 'room'))
            corpus = (get(row, 'corpus') if has_direct_corpus_mapping else None)
            corpus = None if blank(corpus) else corpus
            # A present-but-empty direct column is not a usable corpus mapping.
            row_has_corpus_mapping = corpus is not None
            # TravelLine exports include direct room/corpus columns even when
            # the API could not assign a room. In that case, retain the
            # traditional guest-to-room lookup from the separate rooms file.
            if room is None or corpus is None:
                manual_room, manual_corpus = find_room_info(
                    guest, get(row, 'bron'), checkin, checkout)
                if room is None:
                    room = manual_room
                if corpus is None and manual_corpus is not None:
                    corpus = manual_corpus
                    row_has_corpus_mapping = True
        else:
            room, corpus = find_room_info(
                guest, get(row, 'bron'), checkin, checkout)
            row_has_corpus_mapping = has_corpus_mapping
        records.append({
            'bron': bron, 'num_rooms': num_rooms, 'guest': guest, 'price': price,
            'commission': commission, 'agent': agent,
            'unknown_source': ((not agent) and bool(source)
                               and source not in (source_map or {})),
            'checkin': checkin, 'nights': nights, 'checkout': checkout,
            'source': source, 'category': category, 'room': room,
            'corpus': corpus,
            'has_corpus_mapping': row_has_corpus_mapping,
            'extras_amount': extras_amount, 'service_names': service_names,
            'extra_vals': extra_vals, 'dup_note': None,
        })

    if dup_checker:
        notes = dup_checker([{
            'bron': r['bron'], 'guest_norm': norm_name(r['guest']),
            'checkin': r['checkin'].date() if r['checkin'] else None,
            'checkout': r['checkout'].date() if r['checkout'] else None,
        } for r in records])
        for r, n in zip(records, notes):
            r['dup_note'] = n

    records.sort(key=lambda x: norm_name(x['guest']))

    def corpus_number(value):
        if blank(value):
            return None
        try:
            number = float(value)
            if number.is_integer():
                return int(number)
        except (ValueError, TypeError, OverflowError):
            pass
        match = re.search(r'\d+', str(value))
        return int(match.group()) if match else None

    errors_summary = {
        "Отсутствует № комнаты (КРАСНЫЙ)": 0,
        "Отсутствует корпус (КРАСНЫЙ)": 0,
        "Не указан Гость (ЖЕЛТЫЙ)": 0,
        "Отсутствуют Даты (ОРАНЖЕВЫЙ)": 0,
        "Отсутствует Стоимость (ГОЛУБОЙ)": 0,
        "Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)": 0,
        "Переход кварталов (ФИОЛЕТОВЫЙ)": 0,
    }

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    headers = [
        '№ брони', 'Номеров', 'Гость', 'Стоимость без комиссии', 'Комиссия %',
        'Агент', 'Стоимость с комиссией', 'Допуслуги', 'Стоимость для налога',
        'Наименование допуслуг', 'Заезд', 'Ночей', 'Сут.ставка', 'Метод',
        'Налог', 'Выезд', 'Источник', 'Категория номера', '№ комнаты',
        'корпус', 'Примечание',
    ] + report_extra_data
    base = 'G' if tp.use_full else 'D'
    n_cols = len(headers)
    font_bold = Font(bold=True)

    def write_sheet(name, recs, count_errors=False):
        ws = wb.create_sheet(name)
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = font_bold
        for i, rec in enumerate(recs):
            r = i + 2
            ws.cell(row=r, column=1, value=rec['bron'])
            ws.cell(row=r, column=2, value=rec['num_rooms'])

            c_guest = ws.cell(row=r, column=3, value=rec['guest'])
            c_guest.alignment = Alignment(wrap_text=True, vertical='top')
            if not rec['guest']:
                c_guest.fill = FILL_GUEST
                if count_errors: errors_summary["Не указан Гость (ЖЕЛТЫЙ)"] += 1

            c_price = ws.cell(row=r, column=4, value=rec['price'])
            c_price.number_format = '#,##0.00'
            if rec['price'] is None:
                c_price.fill = FILL_PRICE
                if count_errors: errors_summary["Отсутствует Стоимость (ГОЛУБОЙ)"] += 1

            c_comm = ws.cell(row=r, column=5, value=rec['commission'])
            c_comm.number_format = '0.##'
            c_agent = ws.cell(row=r, column=6, value=rec['agent'] or '')
            if rec['unknown_source']:
                c_comm.fill = FILL_GUEST
                c_agent.fill = FILL_GUEST
                if count_errors:
                    errors_summary["Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)"] += 1

            c_full = ws.cell(row=r, column=7,
                             value=f'=ROUND(D{r}/(100-E{r})*100,2)')
            c_full.number_format = '#,##0.00'

            c_extras = ws.cell(row=r, column=8, value=rec['extras_amount'])
            c_extras.number_format = '#,##0.00'
            extra_term = f'+H{r}' if tp.include_extras_in_tax else ''
            c_tax_base = ws.cell(
                row=r, column=9, value=f'=ROUND({base}{r}{extra_term},2)')
            c_tax_base.number_format = '#,##0.00'
            c_names = ws.cell(row=r, column=10, value=rec['service_names'])
            c_names.alignment = Alignment(wrap_text=True, vertical='top')

            c_in = ws.cell(row=r, column=11, value=rec['checkin'])
            if rec['checkin']:
                c_in.number_format = DATE_FMT
            else:
                c_in.fill = FILL_DATE
                if count_errors: errors_summary["Отсутствуют Даты (ОРАНЖЕВЫЙ)"] += 1

            ws.cell(row=r, column=12, value=rec['nights'])
            c_rate = ws.cell(row=r, column=13,
                             value=tp.f_daily('I', 'L', r))
            c_rate.number_format = '#,##0.00'
            ws.cell(row=r, column=14, value=tp.f_method('M', r))
            c_tax = ws.cell(row=r, column=15,
                            value=tp.f_tax('N', 'L', 'I', r))
            c_tax.number_format = '#,##0.00'

            c_out = ws.cell(row=r, column=16, value=rec['checkout'])
            if rec['checkout']:
                c_out.number_format = DATE_FMT
            else:
                c_out.fill = FILL_DATE
                if count_errors: errors_summary["Отсутствуют Даты (ОРАНЖЕВЫЙ)"] += 1

            ws.cell(row=r, column=17, value=rec['source'])
            cat = ws.cell(row=r, column=18, value=rec['category'])
            cat.alignment = Alignment(wrap_text=True, vertical='top')

            c_room = ws.cell(row=r, column=19, value=rec['room'])
            if rec['room'] is None:
                c_room.fill = FILL_ROOM
                if count_errors: errors_summary["Отсутствует № комнаты (КРАСНЫЙ)"] += 1

            corpus = corpus_number(rec.get('corpus'))
            c_corpus = ws.cell(row=r, column=20, value=corpus)
            if corpus is None and rec['room'] is not None:
                if rec.get('has_corpus_mapping'):
                    c_corpus.fill = FILL_ROOM
                    if count_errors:
                        errors_summary["Отсутствует корпус (КРАСНЫЙ)"] += 1
                else:
                    inferred_corpus = corpus_of(rec)
                    if inferred_corpus is None:
                        c_corpus.fill = FILL_ROOM
                        if count_errors:
                            errors_summary["Отсутствует корпус (КРАСНЫЙ)"] += 1
                    else:
                        # Store the inferred value, not a formula without a
                        # cached result: previews/non-Excel readers then see it.
                        c_corpus.value = inferred_corpus

            c_note = ws.cell(row=r, column=21, value=rec['dup_note'] or '')
            if rec['dup_note']:
                c_note.fill = FILL_DUP
                if count_errors: errors_summary["Переход кварталов (ФИОЛЕТОВЫЙ)"] += 1

            for k, v in enumerate(rec['extra_vals']):
                col_idx = 22 + k
                c_extra = ws.cell(row=r, column=col_idx, value=v)
                label = report_extra_data[k]
                if _is_extra_amount_header(label):
                    c_extra.number_format = '#,##0.00'
                elif (_is_service_name_header(label) or any(
                        word in norm_key(label)
                        for word in ('наимен', 'назван', 'servicename'))):
                    c_extra.alignment = Alignment(wrap_text=True, vertical='top')

        if recs:
            last = len(recs) + 1
            tr = last + 1
            ws.cell(row=tr, column=3, value="ИТОГО:")
            for col in (2, 4, 7, 8, 9, 12, 15):
                letter = get_column_letter(col)
                ws.cell(row=tr, column=col,
                        value=f'=SUM({letter}2:{letter}{last})')
            for c in range(1, n_cols + 1):
                cell = ws.cell(row=tr, column=c)
                cell.font = font_bold
                cell.fill = FILL_TOTALS

        widths = {
            'A': 24, 'B': 9, 'C': 35, 'D': 18, 'E': 10, 'F': 32,
            'G': 20, 'H': 16, 'I': 24, 'J': 48, 'K': 12, 'L': 10,
            'M': 16, 'N': 12, 'O': 12, 'P': 12, 'Q': 25, 'R': 30,
            'S': 10, 'T': 8, 'U': 45,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for k, label in enumerate(report_extra_data):
            col_letter = get_column_letter(22 + k)
            if _is_extra_amount_header(label):
                ws.column_dimensions[col_letter].width = 16
            elif (_is_service_name_header(label) or any(
                    word in norm_key(label)
                    for word in ('наимен', 'назван', 'servicename'))):
                ws.column_dimensions[col_letter].width = 48
            elif col_letter not in widths:
                ws.column_dimensions[col_letter].width = 20
        return ws

    def corpus_of(rec):
        explicit = corpus_number(rec.get('corpus'))
        if explicit is not None:
            return explicit
        if rec.get('has_corpus_mapping'):
            return None
        try:
            return 1 if float(rec['room']) < 50 else 2
        except (ValueError, TypeError):
            return None

    write_sheet('1&2', records, count_errors=True)
    write_sheet('нет корпуса', [x for x in records if corpus_of(x) is None])
    write_sheet('1', [x for x in records if corpus_of(x) == 1])
    write_sheet('2', [x for x in records if corpus_of(x) == 2])
    for corpus in sorted({corpus_of(x) for x in records
                          if corpus_of(x) not in (None, 1, 2)}):
        write_sheet(str(corpus), [x for x in records if corpus_of(x) == corpus])

    wb.save(output_path)
    log.info("Чайковский сохранён: %s", output_path)

    bookings = [{
        'guest': r['guest'], 'guest_norm': norm_name(r['guest']), 'bron': r['bron'],
        'checkin': r['checkin'].date().isoformat() if r['checkin'] else None,
        'checkout': r['checkout'].date().isoformat() if r['checkout'] else None,
        'price': r['price'],
        'price_full': round(r['price'] / (100 - r['commission']) * 100, 2)
                      if r['price'] is not None else None,
        'nights': r['nights'], 'dup_note': r['dup_note'],
    } for r in records]
    return errors_summary, bookings


def process_dona(data_path, output_path, map_data=None,
                 extra_data=None, dup_checker=None, tax=None,
                 agents=None, source_map=None):
    """Возвращает (summary, bookings)."""
    tax_cfg = dict(tax or {})
    tax_cfg.setdefault('include_extras_in_tax', True)
    tp = TaxParams(tax_cfg)
    log.info("Дона: %s", data_path)
    df_data = read_table(data_path)
    headers = [str(c) for c in df_data.columns]
    map_data = map_data or guess_mapping(headers, DNA_DATA_SCHEMA)
    miss = [f['label'] for f in DNA_DATA_SCHEMA if f['required'] and not map_data.get(f['key'])]
    if miss:
        raise ValueError(f"Не сопоставлены колонки: {', '.join(miss)}")
    extra_data = [c for c in (extra_data or []) if c in df_data.columns]
    extra_amount_col = next(
        (c for c in df_data.columns if _is_extra_amount_header(c)), None)
    service_name_col = next(
        (c for c in df_data.columns if _is_service_name_header(c)), None)
    extension_amount_col = next(
        (c for c in df_data.columns if _is_extension_amount_header(c)), None)
    for special_col in (extra_amount_col, service_name_col, extension_amount_col):
        if special_col is not None and special_col not in extra_data:
            extra_data.append(special_col)
    report_extra_data = [
        c for c in extra_data
        if c not in (extra_amount_col, service_name_col, extension_amount_col)
    ]

    def get(row, key):
        col = map_data.get(key)
        return row[col] if col else None

    records = []
    for _, row in df_data.iterrows():
        guest = clean_guest_names(get(row, 'guest'))

        checkin = pd.to_datetime(excel_serial_to_dt(get(row, 'checkin')),
                                 dayfirst=True, errors='coerce')
        checkout = pd.to_datetime(excel_serial_to_dt(get(row, 'checkout')),
                                  dayfirst=True, errors='coerce')
        checkin = None if pd.isna(checkin) else checkin.to_pydatetime()
        checkout = None if pd.isna(checkout) else checkout.to_pydatetime()

        nights_v = get(row, 'nights')
        nights = None
        if nights_v is not None and not pd.isna(nights_v):
            try:
                nights = int(float(nights_v))
            except (ValueError, TypeError):
                nights = None
        if nights is None and checkin and checkout:
            nights = max(1, (checkout.date() - checkin.date()).days)

        category = get(row, 'category')
        category = '' if category is None or pd.isna(category) else str(category).strip()
        num_rooms = max(1, len([x for x in category.split('\n') if x.strip()])) if category else 1

        price = parse_price(get(row, 'price'))

        bron_v = get(row, 'bron')
        bron = '' if bron_v is None or pd.isna(bron_v) else str(bron_v).strip()

        source_v = get(row, 'source')
        source = '' if source_v is None or pd.isna(source_v) else str(source_v).strip()

        sm = (source_map or {}).get(source)
        if sm:
            commission, agent = float(sm[0]), sm[1]
        else:
            # комиссия берётся на дату заезда (плавающие проценты)
            commission, agent = match_agent(
                source, agents or default_agents(),
                date=checkin.date() if checkin else None, hotel='dna')
        names_value = row[service_name_col] if service_name_col else ''
        service_names = ('' if names_value is None or pd.isna(names_value)
                         else str(names_value).strip())
        extras_amount = (parse_price(row[extra_amount_col])
                         if extra_amount_col is not None else 0.0)
        extras_amount = 0.0 if extras_amount is None else round(extras_amount, 2)

        extension_amount = 0.0
        if _contains_named_service(service_names, 'Продление номера'):
            extension_amount = (
                parse_price(row[extension_amount_col])
                if extension_amount_col is not None else None)
            described_amount = _named_service_amount(
                service_names, 'Продление номера')
            if extension_amount is None or (
                    extension_amount == 0 and described_amount > 0):
                extension_amount = described_amount
            extension_amount = (0.0 if extension_amount is None
                                else round(extension_amount, 2))

        extra_vals = [row[c] if not pd.isna(row[c]) else ''
                      for c in report_extra_data]

        records.append({
            'bron': bron, 'num_rooms': num_rooms, 'guest': guest, 'price': price,
            'commission': commission, 'agent': agent,
            'unknown_source': ((not agent) and bool(source)
                               and source not in (source_map or {})),
            'checkin': checkin, 'nights': nights, 'checkout': checkout,
            'source': source, 'category': category,
            'extras_amount': extras_amount,
            'extension_amount': extension_amount,
            'service_names': service_names,
            'extra_vals': extra_vals, 'dup_note': None,
        })

    if dup_checker:
        notes = dup_checker([{
            'bron': r['bron'], 'guest_norm': norm_name(r['guest']),
            'checkin': r['checkin'].date() if r['checkin'] else None,
            'checkout': r['checkout'].date() if r['checkout'] else None,
        } for r in records])
        for r, n in zip(records, notes):
            r['dup_note'] = n

    records.sort(key=lambda x: norm_name(x['guest']))

    summary = {
        "Не указан Гость (ЖЕЛТЫЙ)": 0,
        "Отсутствуют Даты (ОРАНЖЕВЫЙ)": 0,
        "Отсутствует Стоимость (ГОЛУБОЙ)": 0,
        "Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)": 0,
        "Переход кварталов (ФИОЛЕТОВЫЙ)": 0,
    }

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    headers_out = [
        '№ брони', 'Номеров', 'Гость', 'Стоимость без комиссии', 'Комиссия %',
        'Агент', 'Стоимость с комиссией', 'Допуслуги',
        'Стоимость для налога', 'Сумма продления номера',
        'Наименование допуслуг', 'Заезд', 'Ночей', 'Сут.ставка', 'Метод',
        'Налог', 'Выезд', 'Источник', 'Категория номера', 'Примечание',
    ] + report_extra_data
    n_cols = len(headers_out)
    font_bold = Font(bold=True)
    base = 'G' if tp.use_full else 'D'

    ws = wb.create_sheet('Дона')
    for c, h in enumerate(headers_out, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = font_bold

    for i, rec in enumerate(records):
        r = i + 2
        ws.cell(row=r, column=1, value=rec['bron'])
        ws.cell(row=r, column=2, value=rec['num_rooms'])

        c_guest = ws.cell(row=r, column=3, value=rec['guest'])
        c_guest.alignment = Alignment(wrap_text=True, vertical='top')
        if not rec['guest']:
            c_guest.fill = FILL_GUEST
            summary["Не указан Гость (ЖЕЛТЫЙ)"] += 1

        c_price = ws.cell(row=r, column=4, value=rec['price'])
        c_price.number_format = '#,##0.00'
        if rec['price'] is None:
            c_price.fill = FILL_PRICE
            summary["Отсутствует Стоимость (ГОЛУБОЙ)"] += 1

        c_comm = ws.cell(row=r, column=5, value=rec['commission'])
        c_comm.number_format = '0'
        c_agent = ws.cell(row=r, column=6, value=rec['agent'] or '')
        if rec['unknown_source']:
            c_comm.fill = FILL_GUEST
            c_agent.fill = FILL_GUEST
            summary["Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)"] += 1

        c_full = ws.cell(row=r, column=7,
                         value=f'=ROUND(D{r}/(100-E{r})*100,2)')
        c_full.number_format = '#,##0.00'
        c_extras = ws.cell(row=r, column=8, value=rec['extras_amount'])
        c_extras.number_format = '#,##0.00'
        extra_term = f'+J{r}' if tp.include_extras_in_tax else ''
        c_tax_base = ws.cell(
            row=r, column=9, value=f'=ROUND({base}{r}{extra_term},2)')
        c_tax_base.number_format = '#,##0.00'
        c_extension = ws.cell(row=r, column=10, value=rec['extension_amount'])
        c_extension.number_format = '#,##0.00'
        c_names = ws.cell(row=r, column=11, value=rec['service_names'])
        c_names.alignment = Alignment(wrap_text=True, vertical='top')

        c_in = ws.cell(row=r, column=12, value=rec['checkin'])
        if rec['checkin']:
            c_in.number_format = DATE_FMT
        else:
            c_in.fill = FILL_DATE
            summary["Отсутствуют Даты (ОРАНЖЕВЫЙ)"] += 1

        ws.cell(row=r, column=13, value=rec['nights'])
        pcol = 'I'
        c_rate = ws.cell(row=r, column=14,
                         value=tp.f_daily(pcol, 'M', r))
        c_rate.number_format = '#,##0.00'
        ws.cell(row=r, column=15, value=tp.f_method('N', r))
        c_tax = ws.cell(row=r, column=16,
                        value=tp.f_tax('O', 'M', pcol, r))
        c_tax.number_format = '#,##0.00'

        c_out = ws.cell(row=r, column=17, value=rec['checkout'])
        if rec['checkout']:
            c_out.number_format = DATE_FMT
        else:
            c_out.fill = FILL_DATE
            summary["Отсутствуют Даты (ОРАНЖЕВЫЙ)"] += 1

        ws.cell(row=r, column=18, value=rec['source'])
        cat = ws.cell(row=r, column=19, value=rec['category'])
        cat.alignment = Alignment(wrap_text=True, vertical='top')

        c_note = ws.cell(row=r, column=20, value=rec['dup_note'] or '')
        if rec['dup_note']:
            c_note.fill = FILL_DUP
            summary["Переход кварталов (ФИОЛЕТОВЫЙ)"] += 1

        for k, v in enumerate(rec['extra_vals']):
            ws.cell(row=r, column=21 + k, value=v)

    if records:
        last = len(records) + 1
        tr = last + 1
        ws.cell(row=tr, column=3, value="ИТОГО:")
        for col in (2, 4, 7, 8, 9, 10, 13, 16):
            letter = get_column_letter(col)
            ws.cell(row=tr, column=col,
                    value=f'=SUM({letter}2:{letter}{last})')
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=tr, column=c)
            cell.font = font_bold
            cell.fill = FILL_TOTALS

    widths = {
        'A': 24, 'B': 9, 'C': 35, 'D': 18, 'E': 11, 'F': 30,
        'G': 20, 'H': 16, 'I': 24, 'J': 24, 'K': 48, 'L': 12,
        'M': 10, 'N': 16, 'O': 12, 'P': 12, 'Q': 12, 'R': 35,
        'S': 30, 'T': 45,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for k, label in enumerate(report_extra_data):
        col_letter = get_column_letter(21 + k)
        if _is_extra_amount_header(label):
            ws.column_dimensions[col_letter].width = 16
        elif (_is_service_name_header(label) or any(
                word in norm_key(label)
                for word in ('наимен', 'назван', 'servicename'))):
            ws.column_dimensions[col_letter].width = 48
        elif col_letter not in widths:
            ws.column_dimensions[col_letter].width = 20

    wb.save(output_path)
    log.info("Дона сохранена: %s", output_path)

    bookings = [{
        'guest': r['guest'], 'guest_norm': norm_name(r['guest']), 'bron': r['bron'],
        'checkin': r['checkin'].date().isoformat() if r['checkin'] else None,
        'checkout': r['checkout'].date().isoformat() if r['checkout'] else None,
        'price': r['price'],
        'price_full': round(r['price'] / (100 - r['commission']) * 100, 2)
                      if r['price'] is not None else None,
        'nights': r['nights'], 'dup_note': r['dup_note'],
    } for r in records]
    return summary, bookings


def load_workbook_safe(path):
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except WorkbookReadError:
        raise
    except Exception:
        fixed, _notes = repair_workbook(path)
        return openpyxl.load_workbook(fixed, data_only=True)


def process_usadba( report_path, output_path, dup_checker=None,
                   tax=None, agents=None, source_map=None):
    """Возвращает (stats, bookings)."""
    tp = TaxParams(tax)
    log.info("Усадьба: %s", report_path)
    wb_in = load_workbook_safe(report_path)
    ws_in = wb_in[wb_in.sheetnames[0]]

    valid_states = ('Проживал', 'Сейчас проживает')


    records = []
    rooms = set()
    current_room = None
    stats = {'rows': 0, 'rooms': 0, 'zero_price': 0, 'bad_dates': 0,
             'still_living': 0, 'unknown_source': 0}

    for r in range(1, ws_in.max_row + 1):
        a_val = ws_in.cell(row=r, column=1).value
        if isinstance(a_val, str) and a_val.strip().startswith('Номер:'):
            current_room = a_val.split(':', 1)[1].strip()
            rooms.add(current_room)
            continue

        state = ws_in.cell(row=r, column=5).value
        fio = ws_in.cell(row=r, column=6).value
        if not fio or not state or str(state).strip() not in valid_states:
            continue

        checkin_raw = ws_in.cell(row=r, column=7).value
        checkout_raw = ws_in.cell(row=r, column=8).value
        price = parse_price(ws_in.cell(row=r, column=9).value)
        tariff = ws_in.cell(row=r, column=14).value
        org = ws_in.cell(row=r, column=16).value
        org_s = '' if org is None else str(org).strip()

        dt_in = parse_usd_dt(checkin_raw)
        dt_out = parse_usd_dt(checkout_raw)

        if dt_in and dt_out:
            nights = max(1, (dt_out.date() - dt_in.date()).days)
        else:
            nights = 1
            stats['bad_dates'] += 1

        if str(state).strip() == 'Сейчас проживает':
            stats['still_living'] += 1
        if not price:
            stats['zero_price'] += 1

        sm = (source_map or {}).get(org_s)
        if sm:
            commission, agent = float(sm[0]), sm[1]
        else:
            commission, agent = match_agent(
                org_s, agents or default_agents(),
                date=dt_in.date() if dt_in else None, hotel='usd')
        if not agent and org_s:
            stats['unknown_source'] += 1

        records.append({
            'fio': str(fio).strip(),
            'checkin': dt_in if dt_in else (str(checkin_raw or '').strip()),
            'checkout': dt_out if dt_out else (str(checkout_raw or '').strip()),
            'nights': nights,
            'price': price if price is not None else 0,
            'commission': commission, 'agent': agent,
            'unknown_source': ((not agent) and bool(org_s)
                               and org_s not in (source_map or {})),
            'tariff': str(tariff).strip() if tariff else '',
            'org': org_s,
            'room': current_room,
            'dup_note': None,
        })

    if not records:
        raise ValueError("В файле не найдено ни одной строки гостя.\n"
                         "Убедитесь, что выбран отчёт «Доход по номеру» из UCS Shelter.")

    if dup_checker:
        notes = dup_checker([{
            'bron': '', 'guest_norm': norm_name(rec['fio']),
            'checkin': rec['checkin'].date() if isinstance(rec['checkin'], datetime) else None,
            'checkout': rec['checkout'].date() if isinstance(rec['checkout'], datetime) else None,
        } for rec in records])
        for rec, n in zip(records, notes):
            rec['dup_note'] = n
    stats['dups'] = sum(1 for rec in records if rec['dup_note'])

    stats['rows'] = len(records)
    stats['rooms'] = len(rooms)

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = 'Page 1'

    headers = ['Фамилия Имя Отчество', 'Заселение', 'ночей', 'Выселение ',
               'Стоимость без комиссии', 'Комиссия %', 'Агент', 'Стоимость с комиссией',
               'Сут.ставка', 'Метод', 'Налог', 'Тариф',
               'Организация\nзаказчик', 'Примечание']
    font_bold = Font(bold=True)
    for c, h in enumerate(headers, start=1):
        cell = ws_out.cell(row=1, column=c, value=h)
        cell.font = font_bold
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    base = 'H' if tp.use_full else 'E'

    for i, rec in enumerate(records):
        row = i + 2
        ws_out.cell(row=row, column=1, value=rec['fio'])

        c_in = ws_out.cell(row=row, column=2, value=rec['checkin'])
        c_out = ws_out.cell(row=row, column=4, value=rec['checkout'])
        if isinstance(rec['checkin'], datetime):
            c_in.number_format = DATE_FMT
        else:
            c_in.fill = FILL_DATE
        if isinstance(rec['checkout'], datetime):
            c_out.number_format = DATE_FMT
        else:
            c_out.fill = FILL_DATE

        ws_out.cell(row=row, column=3, value=rec['nights'])

        c_price = ws_out.cell(row=row, column=5, value=rec['price'])
        c_price.number_format = r'#,##0.00\ "₽"'
        if not rec['price']:
            c_price.fill = FILL_PRICE

        c_comm = ws_out.cell(row=row, column=6, value=rec['commission'])
        c_comm.number_format = '0.##'
        c_agent = ws_out.cell(row=row, column=7, value=rec['agent'] or '')
        if rec['unknown_source']:
            c_comm.fill = FILL_GUEST
            c_agent.fill = FILL_GUEST

        c_full = ws_out.cell(row=row, column=8,
                             value=f'=ROUND(E{row}/(100-F{row})*100,2)')
        c_full.number_format = '#,##0.00'

        c_rate = ws_out.cell(row=row, column=9, value=tp.f_daily(base, 'C', row))
        c_rate.number_format = '#,##0.00'
        ws_out.cell(row=row, column=10, value=tp.f_method('I', row))
        c_tax = ws_out.cell(row=row, column=11, value=tp.f_tax('J', 'C', base, row))
        c_tax.number_format = '#,##0.00'
        ws_out.cell(row=row, column=12, value=rec['tariff'])
        ws_out.cell(row=row, column=13, value=rec['org'])

        c_note = ws_out.cell(row=row, column=14, value=rec['dup_note'] or '')
        if rec['dup_note']:
            c_note.fill = FILL_DUP

    last_row = len(records) + 1
    total_row = last_row + 1

    ws_out.cell(row=total_row, column=1, value="ИТОГО:")
    ws_out.cell(row=total_row, column=5, value=f'=SUM(E2:E{last_row})').number_format = r'#,##0.00\ "₽"'
    ws_out.cell(row=total_row, column=8, value=f'=SUM(H2:H{last_row})')
    ws_out.cell(row=total_row, column=11, value=f'=SUM(K2:K{last_row})')
    for c in range(1, 15):
        cell = ws_out.cell(row=total_row, column=c)
        cell.font = font_bold
        cell.fill = FILL_TOTALS

    widths = {'A': 38, 'B': 18, 'C': 8, 'D': 18, 'E': 14, 'F': 10, 'G': 30,
              'H': 18, 'I': 12, 'J': 12, 'K': 12, 'L': 10, 'M': 40, 'N': 45}
    for col, w in widths.items():
        ws_out.column_dimensions[col].width = w

    ws2 = wb_out.create_sheet('Лист1')
    ws2.cell(row=1, column=1, value='Заселение').font = font_bold
    ws2.cell(row=1, column=2, value='Выселение ').font = font_bold
    for i, rec in enumerate(records):
        row = i + 2
        c_a = ws2.cell(row=row, column=1, value=rec['checkin'])
        c_b = ws2.cell(row=row, column=2, value=rec['checkout'])
        if isinstance(rec['checkin'], datetime):
            c_a.number_format = DATE_FMT
        if isinstance(rec['checkout'], datetime):
            c_b.number_format = DATE_FMT
        ws2.cell(row=row, column=4, value=f'=INT(A{row})')
        ws2.cell(row=row, column=5, value=f'=INT(B{row})')
        ws2.cell(row=row, column=6, value=f'=E{row}-D{row}')
    ws2.column_dimensions['A'].width = 18
    ws2.column_dimensions['B'].width = 18

    wb_out.save(output_path)
    log.info("Усадьба сохранена: %s", output_path)

    bookings = [{
        'guest': rec['fio'], 'guest_norm': norm_name(rec['fio']), 'bron': '',
        'checkin': rec['checkin'].date().isoformat() if isinstance(rec['checkin'], datetime) else None,
        'checkout': rec['checkout'].date().isoformat() if isinstance(rec['checkout'], datetime) else None,
        'price': rec['price'],
        'price_full': round(rec['price'] / (100 - rec['commission']) * 100, 2)
                      if rec['price'] else None,
        'nights': rec['nights'], 'dup_note': rec['dup_note'],
    } for rec in records]
    return stats, bookings


def usadba_sources(path):
    """Колонка «Организация-заказчик» из отчёта UCS Shelter (для диалога источников)."""
    wb = load_workbook_safe(path)
    ws = wb[wb.sheetnames[0]]
    out = []
    for r in range(1, ws.max_row + 1):
        state = ws.cell(row=r, column=5).value
        org = ws.cell(row=r, column=16).value
        if org and state and str(state).strip() in ('Проживал', 'Сейчас проживает'):
            out.append(str(org))
    return out
