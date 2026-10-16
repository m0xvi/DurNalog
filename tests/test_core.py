# -*- coding: utf-8 -*-
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from agents import match_dona_agent
from tax import parse_price, dates_overlap, quarter_of, tax_formula_method
from db import HistoryDB

D = date


def test_match_agents():
    assert match_dona_agent("Ostrovok.ru (Emerging Travel Group)") == (15, "Островок")
    assert match_dona_agent("Bronevik.com") == (20, "Броневик")
    assert match_dona_agent("Cuva") == (10, "ИП лагун = КУВА")
    assert match_dona_agent("Т-Путешествия") == (15, "Тинькофф путешествия")
    assert match_dona_agent("yandex.travel") == (20, "Яндекс")
    assert match_dona_agent("прямая бронь (сайт)") == (0, None)


def test_parse_price():
    assert parse_price(5000) == 5000.0
    assert parse_price("5 000,00") == 5000.0
    assert parse_price("12 345.6 руб.") == 12345.6
    assert parse_price(None) is None
    assert parse_price("") is None
    assert parse_price("нет") is None


def test_dates_and_quarter():
    assert dates_overlap(D(2026, 3, 30), D(2026, 4, 2), D(2026, 3, 30), D(2026, 4, 2))
    assert dates_overlap(D(2026, 3, 28), D(2026, 4, 1), D(2026, 3, 30), D(2026, 4, 5))
    # смежные ночи — НЕ пересечение (две разные брони подряд)
    assert not dates_overlap(D(2026, 3, 30), D(2026, 4, 2), D(2026, 4, 2), D(2026, 4, 5))
    assert not dates_overlap(None, D(2026, 4, 2), D(2026, 3, 30), D(2026, 4, 2))
    assert quarter_of(D(2026, 1, 1)) == "2026-Q1"
    assert quarter_of(D(2026, 4, 1)) == "2026-Q2"
    assert quarter_of(D(2026, 12, 31)) == "2026-Q4"
    assert "5100" in tax_formula_method("J", 2)


def _mk_db(tmp_path):
    return HistoryDB(os.path.join(str(tmp_path), "test.db"))


def _booking(guest="Иванов Иван", bron="", ci="2026-03-30", co="2026-04-02"):
    return {"guest": guest, "guest_norm": guest.lower(), "bron": bron,
            "checkin": ci, "checkout": co, "price": 9000, "price_full": None,
            "nights": 3, "dup_note": None}


def test_db_duplicates(tmp_path):
    db = _mk_db(tmp_path)
    rid1 = db.add_report("Дона", "2026-Q1", "q1.csv", b"orig1", "q1.xlsx", b"out1", 1, 0)
    db.add_bookings(rid1, [_booking()])

    # та же бронь (ФИО+даты) во втором квартале -> дубль
    notes = db.check_duplicates("Дона", [{
        "bron": "", "guest_norm": "иванов иван",
        "checkin": D(2026, 3, 30), "checkout": D(2026, 4, 2)}])
    assert notes[0] and "2026-Q1" in notes[0]

    # другой гость -> не дубль
    notes = db.check_duplicates("Дона", [{
        "bron": "", "guest_norm": "петров",
        "checkin": D(2026, 3, 30), "checkout": D(2026, 4, 2)}])
    assert notes[0] is None

    # тот же гость, но другие (несмежные) даты -> не дубль
    notes = db.check_duplicates("Дона", [{
        "bron": "", "guest_norm": "иванов иван",
        "checkin": D(2026, 5, 1), "checkout": D(2026, 5, 3)}])
    assert notes[0] is None

    # по № брони — даже без пересечения дат
    rid2 = db.add_report("Дона", "2026-Q2", "q2.csv", b"o2", "q2.xlsx", b"o2", 1, 0)
    db.add_bookings(rid2, [_booking(guest="Сидоров", bron="D-777",
                                   ci="2026-04-10", co="2026-04-12")])
    notes = db.check_duplicates("Дона", [{
        "bron": "D-777", "guest_norm": "другой",
        "checkin": D(2026, 6, 1), "checkout": D(2026, 6, 2)}])
    assert notes[0] and "2026-Q2" in notes[0]

    # другой отель не пересекается
    notes = db.check_duplicates("Чайковский", [{
        "bron": "", "guest_norm": "иванов иван",
        "checkin": D(2026, 3, 30), "checkout": D(2026, 4, 2)}])
    assert notes[0] is None

    # файлы отчёта достаются из БД
    name, blob = db.get_file(rid1, "orig")
    assert name == "q1.csv" and blob == b"orig1"

    # отчёт того же квартала находится и удаляется
    assert db.find_same("Дона", "2026-Q1")["id"] == rid1
    db.delete_report(rid1)
    assert db.find_same("Дона", "2026-Q1") is None
    db.close()


def test_dona_end_to_end(tmp_path):
    import pandas as pd
    import openpyxl
    import core

    src = os.path.join(str(tmp_path), "dona.csv")
    out = os.path.join(str(tmp_path), "dona.xlsx")
    pd.DataFrame([
        ["Иванов Иван", 8500, "30.03.2026", "02.04.2026", "Ostrovok.ru", "D-1", 3, "Стандарт"],
        ["Иванов Иван", 8500, "30.03.2026", "02.04.2026", "Ostrovok.ru", "D-1", 3, "Стандарт"],
    ], columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник", "№ брони",
                "Ночей", "Категория номера"]).to_csv(src, index=False)

    # второй «отчёт» с той же бронью -> checker помечает дубль
    def checker(rows):
        return ["Переход кварталов: тест" if r["guest_norm"] == "иванов иван" else None
                for r in rows]

    summary, bookings = core.process_dona(src, out, dup_checker=checker)
    assert summary["Переход кварталов (ФИОЛЕТОВЫЙ)"] == 2
    assert all(b["dup_note"] for b in bookings)
    ws = openpyxl.load_workbook(out)["Дона"]
    assert ws["T2"].value == "Переход кварталов: тест"
    assert ws["T2"].fill.start_color.rgb.endswith("CC99FF")
    assert ws["E2"].value == 15
    assert ws["G2"].value == "=ROUND(D2/(100-E2)*100,2)"


def test_tchaikovsky_end_to_end(tmp_path):
    import pandas as pd
    import openpyxl
    import core

    src = os.path.join(str(tmp_path), "tch.csv")
    rooms = os.path.join(str(tmp_path), "rooms.csv")
    out = os.path.join(str(tmp_path), "tch.xlsx")
    pd.DataFrame([
        ["Иванов Иван", "5 000,00", "01.08.2026", "04.08.2026", "Сайт", "Стандарт", "B-1", 3],
    ], columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник",
                "Категория номера", "№ брони", "Ночей"]).to_csv(src, index=False)
    pd.DataFrame([["Иванов Иван", 12]], columns=["ФИО", "№ комнаты"]).to_csv(rooms, index=False)

    summary, bookings = core.process_tchaikovsky(src, rooms, out)
    assert summary["Отсутствует Стоимость (ГОЛУБОЙ)"] == 0  # "5 000,00" распознано
    ws = openpyxl.load_workbook(out)["1&2"]
    assert ws["D2"].value == 5000.0
    assert ws["G2"].value == '=ROUND(D2/(100-E2)*100,2)'
    assert ws["I1"].value == "Стоимость для налога"
    assert ws["I2"].value == "=ROUND(G2+H2,2)"
    assert ws["N2"].value == '=IF(M2>=5100,"ставка","мин.ставка")'
    assert ws["O2"].value == '=ROUND(IF(N2="мин.ставка",L2*100,I2/102*2),2)'
    assert ws["U1"].value == "Примечание"
    assert bookings[0]["price"] == 5000.0


def test_travelline_tax_base_override_does_not_reapply_agent_commission(tmp_path):
    import pandas as pd
    import openpyxl
    import core

    src = tmp_path / "tl-tax-base.csv"
    rooms = tmp_path / "rooms.csv"
    out = tmp_path / "tl-tax-base.xlsx"
    tax_base_header = "Налоговая база TravelLine, ₽"
    pd.DataFrame([[
        "Иванов Иван", 88448.20, "22.08.2026", "27.08.2026",
        "TravelLine YAX", "Стандарт", "TL-8186", 5, 89000,
    ]], columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник",
                "Категория номера", "№ брони", "Ночей", tax_base_header]
                 ).to_csv(src, index=False)
    pd.DataFrame([["Иванов Иван", 12]],
                 columns=["ФИО", "№ комнаты"]).to_csv(rooms, index=False)

    _, bookings = core.process_tchaikovsky(
        str(src), str(rooms), str(out), extra_data=[tax_base_header],
        source_map={"TravelLine YAX": (18, "Яндекс.Путешествия")},
        tax={"include_extras_in_tax": False})
    ws = openpyxl.load_workbook(out)["1&2"]
    assert ws["D2"].value == 88448.2
    assert ws["E2"].value == 18
    assert ws["F2"].value == "Яндекс.Путешествия"
    assert ws["G2"].value == "=ROUND(D2/(100-E2)*100,2)"
    assert ws["V1"].value == tax_base_header
    assert ws["V2"].value == 89000.0
    assert ws["I2"].value == "=ROUND(IF(ISNUMBER(V2),V2,G2),2)"
    assert bookings[0]["price"] == 88448.2


def test_tchaikovsky_tax_includes_extras_after_commission_and_can_disable(tmp_path):
    import pandas as pd
    import openpyxl
    import core

    src = tmp_path / "tch-extras.csv"
    rooms = tmp_path / "rooms.csv"
    pd.DataFrame([[
        "Иванов Иван", 10000, "01.08.2026", "03.08.2026", "OTA",
        "Стандарт", "B-EXTRA", 2, 1500, "Завтрак",
    ]], columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник",
                "Категория номера", "№ брони", "Ночей", "Допуслуги",
                "Наименование допуслуг"]).to_csv(src, index=False)
    pd.DataFrame([["Иванов Иван", 12]],
                 columns=["ФИО", "№ комнаты"]).to_csv(rooms, index=False)
    extras = ["Допуслуги", "Наименование допуслуг"]
    source_map = {"OTA": (15, "Тестовый агент")}

    with_extras = tmp_path / "tch-extras-on.xlsx"
    core.process_tchaikovsky(
        str(src), str(rooms), str(with_extras), extra_data=extras,
        source_map=source_map, tax={"include_extras_in_tax": True})
    ws = openpyxl.load_workbook(with_extras)["1&2"]
    assert ws["G1"].value == "Стоимость с комиссией"
    assert ws["H1"].value == "Допуслуги"
    assert ws["I1"].value == "Стоимость для налога"
    assert ws["G2"].value == "=ROUND(D2/(100-E2)*100,2)"
    assert ws["H2"].value == 1500
    assert ws["I1"].value == "Стоимость для налога"
    assert ws["I2"].value == "=ROUND(G2+H2,2)"
    assert ws["J2"].value == "Завтрак"
    assert ws["M2"].value == "=I2/L2"
    assert ws["O2"].value == '=ROUND(IF(N2="мин.ставка",L2*100,I2/102*2),2)'

    without_extras = tmp_path / "tch-extras-off.xlsx"
    core.process_tchaikovsky(
        str(src), str(rooms), str(without_extras), extra_data=extras,
        source_map=source_map, tax={"include_extras_in_tax": False})
    ws_off = openpyxl.load_workbook(without_extras)["1&2"]
    assert ws_off["I1"].value == "Стоимость для налога"
    assert ws_off["I2"].value == "=ROUND(G2,2)"
    assert ws_off["M2"].value == "=I2/L2"
    assert ws_off["O2"].value == '=ROUND(IF(N2="мин.ставка",L2*100,I2/102*2),2)'


def test_dona_travelline_tax_base_override_keeps_discounted_exception(tmp_path):
    import pandas as pd
    import openpyxl
    import core

    src = tmp_path / "dona-tl-tax-base.csv"
    out = tmp_path / "dona-tl-tax-base.xlsx"
    tax_base_header = "Налоговая база TravelLine, ₽"
    pd.DataFrame([[
        "Иванов Иван", 88448.20, "22.08.2026", "27.08.2026",
        "TravelLine YAX", 88448.20,
    ]], columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник",
                tax_base_header]).to_csv(src, index=False)

    core.process_dona(
        str(src), str(out), extra_data=[tax_base_header],
        source_map={"TravelLine YAX": (18, "Яндекс.Путешествия")},
        tax={"include_extras_in_tax": False})
    ws = openpyxl.load_workbook(out)["Дона"]
    assert ws["E2"].value == 18
    assert ws["G2"].value == "=ROUND(D2/(100-E2)*100,2)"
    assert ws["I2"].value == "=ROUND(IF(ISNUMBER(U2),U2,G2),2)"
    assert ws["U1"].value == tax_base_header
    assert ws["U2"].value == 88448.20


def test_dona_does_not_add_extras_to_tax_base(tmp_path):
    import pandas as pd
    import openpyxl
    import core

    src = tmp_path / "dona-extras.csv"
    out = tmp_path / "dona-extras.xlsx"
    pd.DataFrame([[
        "Иванов Иван", 10000, "01.08.2026", "03.08.2026", "Сайт", 1500,
    ]], columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник",
                "Допуслуги"]).to_csv(src, index=False)
    core.process_dona(
        str(src), str(out), extra_data=["Допуслуги"],
        tax={"include_extras_in_tax": True})
    ws = openpyxl.load_workbook(out)["Дона"]
    assert ws["G1"].value == "Стоимость с комиссией"
    assert ws["H1"].value == "Допуслуги"
    assert ws["I1"].value == "Стоимость для налога"
    assert ws["G2"].value == "=ROUND(D2/(100-E2)*100,2)"
    assert ws["H2"].value == 1500
    assert ws["I2"].value == "=ROUND(G2+J2,2)"
    assert ws["J2"].value == 0
    assert ws["P2"].value == '=ROUND(IF(O2="мин.ставка",M2*100,I2/102*2),2)'


def test_dona_tax_base_includes_only_named_room_extension(tmp_path):
    import pandas as pd
    import openpyxl
    import core

    src = tmp_path / "dona-extension.csv"
    out = tmp_path / "dona-extension.xlsx"
    pd.DataFrame([[
        "Иванов Иван", 10000, "01.08.2026", "03.08.2026", "Сайт", 2500,
        "Продление номера — 1 000,00 ₽; Завтрак — 1 500,00 ₽", 1000,
    ]], columns=[
        "Гость", "Стоимость", "Заезд", "Выезд", "Источник", "Допуслуги",
        "Наименование допуслуг", "Сумма продления номера",
    ]).to_csv(src, index=False)

    core.process_dona(str(src), str(out),
                      tax={"include_extras_in_tax": True})
    ws = openpyxl.load_workbook(out)["Дона"]
    assert ws["H2"].value == 2500
    assert ws["I1"].value == "Стоимость для налога"
    assert ws["I2"].value == "=ROUND(G2+J2,2)"
    assert ws["J2"].value == 1000
    assert ws["K2"].value == (
        "Продление номера — 1 000,00 ₽; Завтрак — 1 500,00 ₽")
    assert ws["P2"].value == '=ROUND(IF(O2="мин.ставка",M2*100,I2/102*2),2)'

    disabled = tmp_path / "dona-extension-off.xlsx"
    core.process_dona(str(src), str(disabled),
                      tax={"include_extras_in_tax": False})
    ws_off = openpyxl.load_workbook(disabled)["Дона"]
    assert ws_off["J2"].value == 1000  # для контроля сумма показана и без включения
    assert ws_off["I2"].value == "=ROUND(G2,2)"

    # Older/manual exports may only have the service description with amount.
    fallback_src = tmp_path / "dona-extension-no-amount-column.csv"
    fallback_out = tmp_path / "dona-extension-no-amount-column.xlsx"
    pd.DataFrame([[
        "Иванов Иван", 10000, "01.08.2026", "03.08.2026", "Сайт", 2500,
        "Продление номера — 1 000,00 ₽; Завтрак — 1 500,00 ₽",
    ]], columns=[
        "Гость", "Стоимость", "Заезд", "Выезд", "Источник", "Допуслуги",
        "Наименование допуслуг",
    ]).to_csv(fallback_src, index=False)
    core.process_dona(str(fallback_src), str(fallback_out))
    ws_fallback = openpyxl.load_workbook(fallback_out)["Дона"]
    assert ws_fallback["J2"].value == 1000
    assert ws_fallback["I2"].value == "=ROUND(G2+J2,2)"


def test_guess_mapping_renamed_headers():
    from schemas import TCH_DATA_SCHEMA
    from core import guess_mapping, mapping_valid
    headers = ["ФИО гостя", "Сумма", "Дата заезда", "Дата выезда",
               "Источник бронирования", "Категория", "ID брони", "Кол-во ночей",
               "Лишняя колонка"]
    m = guess_mapping(headers, TCH_DATA_SCHEMA)
    assert m['guest'] == "ФИО гостя"
    assert m['price'] == "Сумма"
    assert m['checkin'] == "Дата заезда"
    assert m['checkout'] == "Дата выезда"
    assert m['source'] == "Источник бронирования"
    assert m['bron'] == "ID брони"
    assert m['nights'] == "Кол-во ночей"
    assert mapping_valid(m, headers, TCH_DATA_SCHEMA)


def test_dona_extra_columns_and_custom_mapping(tmp_path):
    import pandas as pd
    import openpyxl
    import core

    src = os.path.join(str(tmp_path), "dona2.csv")
    out = os.path.join(str(tmp_path), "dona2.xlsx")
    pd.DataFrame([
        ["Иванов Иван", 8500, "01.08.2026", "04.08.2026", "островок", "Менеджер"],
    ], columns=["Гость", "Цена", "Прибытие", "Убытие", "Канал", "Менеджер"]
    ).to_csv(src, index=False)

    mapping = {"guest": "Гость", "price": "Цена", "checkin": "Прибытие",
               "checkout": "Убытие", "source": "Канал", "category": None,
               "bron": None, "nights": None}
    summary, bookings = core.process_dona(src, out, map_data=mapping,
                                              extra_data=["Менеджер"])
    ws = openpyxl.load_workbook(out)["Дона"]
    assert ws["U1"].value == "Менеджер"          # доп. колонка в шапке
    assert ws["U2"].value == "Менеджер"          # значение перенесено
    assert ws["D2"].value == 8500
    assert ws["E2"].value == 16  # островок распознан: 16% с 01.08.2026 (заезд 01.08)
    assert summary["Отсутствует Стоимость (ГОЛУБОЙ)"] == 0


def test_tax_params_custom():
    from tax import TaxParams
    t = TaxParams()  # значения по умолчанию = прежние формулы
    assert t.include_extras_in_tax is False
    assert t.f_daily('D', 'F', 2) == '=D2/F2'
    assert t.f_method('G', 2) == '=IF(G2>=5100,"ставка","мин.ставка")'
    assert t.f_tax('H', 'F', 'D', 2) == '=ROUND(IF(H2="мин.ставка",F2*100,D2/102*2),2)'
    t3 = TaxParams({'rate': 3, 'min_per_night': 200})
    assert t3.f_tax('H', 'F', 'D', 2) == '=ROUND(IF(H2="мин.ставка",F2*200,D2/103*3),2)'
    assert abs(t3.auto_threshold() - 200 * 103 / 3) < 1e-9
    t_not_incl = TaxParams({'incl': False, 'auto_threshold': True})
    assert t_not_incl.f_tax('H', 'F', 'D', 2) == '=ROUND(IF(H2="мин.ставка",F2*100,D2/100*2),2)'
    # свой шаблон формулы
    t_tpl = TaxParams({'tpl_tax': '=ROUND({N}*{rate},2)'})
    assert t_tpl.f_tax('H', 'F', 'D', 5) == '=ROUND(F5*2,2)'


def test_agent_overrides():
    from agents import match_dona_agent
    assert match_dona_agent("островок", {"Островок": 25}) == (25, "Островок")
    assert match_dona_agent("островок") == (15, "Островок")


def test_dona_custom_tax_in_output(tmp_path):
    import pandas as pd
    import openpyxl
    import core
    src = os.path.join(str(tmp_path), "d3.csv")
    out = os.path.join(str(tmp_path), "d3.xlsx")
    pd.DataFrame([["Иванов Иван", 8500, "01.08.2026", "04.08.2026", "островок"]],
                 columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник"]
                 ).to_csv(src, index=False)
    _, _ = core.process_dona(src, out, tax={'rate': 3})
    ws = openpyxl.load_workbook(out)["Дона"]
    assert ws["P2"].value == '=ROUND(IF(O2="мин.ставка",M2*100,I2/103*3),2)'


def test_match_cross_language_and_typos():
    from agents import match_dona_agent as m
    # другой язык (транслитерация)
    assert m("Ostrovok.ru")[1] == "Островок"
    assert m("ostrovok")[1] == "Островок"
    assert m("SUTOCHNO")[1] == "Суточно.ру"
    assert m("Yandex Travel")[1] == "Яндекс"
    assert m("Bronevik")[1] == "Броневик"
    # опечатка / сокращение
    assert m("Ostrovk.ru")[1] == "Островок"
    assert m("Янт.Мозаика")[1] == "Янт.Мозаика"
    # исправленная комиссия АВС Трэвел
    assert m("АВС Трэвел") == (10, "АВС Трэвел")
    # по-прежнему не распознаёт постороннее
    assert m("прямая бронь (сайт)") == (0, None)


def test_match_user_aliases():
    from agents import match_dona_agent as m
    assert m("ООО Ромашка", None, {"Островок": ["ромашка"]}) == (15, "Островок")
    assert m("ООО Ромашка") == (0, None)


def test_direct_channels_and_custom_agents():
    from agents import match_dona_agent as m
    direct = "Прямые бронирования (сайт/стойка/звонки/email)"
    assert m("Официальный сайт\nageevgroup.ru") == (0, direct)
    assert m("Бронирование из шахматки (От стойки)") == (0, direct)
    assert m("Бронирование из шахматки (По звонку)") == (0, direct)
    assert m("Бронирование из шахматки (Email)") == (0, direct)
    # действительно неизвестные — остаются нераспознанными (пойдут в диалог)
    # агенты из отчёта Q3 2026 — встроены в таблицу Доны
    assert m("Ozon (new)") == (15, 'Ozon Travel')
    assert m("Ozon Travel") == (15, 'Ozon Travel')
    assert m("HotelBook") == (15, 'HotelBook')
    assert m('Acase.ru (ООО "АОН.РУ")') == (20, 'Академсервис')
    assert m("Roomlink (Zabroniryi.ru)") == (18, 'Эй энд Эй (ЦБП)')
    assert m("OneTwoTrip") == (15, 'OneTwoTrip')
    assert m("Trip.com Group") == (15, 'Trip.com')
    assert m("Агрегатор каналов Comfort Booking") == (15, 'Туту.ру')
    # юрнаименования из того же отчёта — синонимы существующих агентов
    assert m('ООО "Система Алеан"') == (0, 'Алеан Отдых')  # СКИДКА
    assert m('ООО "Пегас СПб"') == (0, 'Хомлин Лэнд (Pegas)')  # СКИДКА
    assert m('ООО "Туроператор Дельфин"') == (0, 'Туроператор Дельфин')  # СКИДКА
    assert m('ООО "Бронирование гостиниц"') == (15, 'Островок')
    assert m('ООО "Компания Броневик"') == (20, 'Броневик')
    assert m('ООО "Академ Онлайн"') == (20, 'Академсервис')
    assert m("Бронирование из шахматки (Туроператор)")[1] is None
    # кастомный агент, добавленный пользователем
    cust = [{'name': 'Extra Channel', 'commission': 12, 'aliases': ['extraxyz']}]
    assert m("ExtraXYZ booking", None, None, cust) == (12, 'Extra Channel')


def test_floating_rates_by_date():
    from agents import default_agents, match_agent, rate_for
    agents = default_agents()
    ost = next(a for a in agents if a['name'] == 'Островок')
    ost['rates'] = [{'from': '', 'to': '31.03.2026', 'pct': 15},
                    {'from': '01.04.2026', 'to': '', 'pct': 17}]
    from datetime import date
    assert rate_for(ost, date(2026, 3, 30)) == 15
    assert rate_for(ost, date(2026, 4, 1)) == 17
    assert rate_for(ost, None) == 15  # без даты — базовый (первый открытый? первый)
    c, n = match_agent("Ostrovok.ru", agents, date=date(2026, 5, 10))
    assert (c, n) == (17, 'Островок')
    c, n = match_agent("Ostrovok.ru", agents, date=date(2026, 1, 10))
    assert (c, n) == (15, 'Островок')


def test_agents_hotel_filter():
    from agents import default_agents, match_agent
    agents = default_agents()
    # агент Усадьбы не распознаётся в Доне и наоборот
    assert match_agent("Свена Тур", agents, hotel='dna')[1] is None
    assert match_agent("Свена Тур", agents, hotel='usd')[1] is not None
    # а «Комфорт Букинг» в Доне резолвится через Туту.ру (агрегатор каналов)
    assert match_agent("Комфорт Букинг", agents, hotel='dna')[1] == 'Туту.ру'
    # у Чайковского свой Островок
    assert match_agent("Ostrovok.ru", agents, hotel='tch')[1] is not None
    assert match_agent("Ostrovok.ru", agents, hotel='dna')[1] == 'Островок'


def test_agent_zero_commission_override_preserves_configured_rates():
    from agents import default_agents, match_agent

    agents = default_agents()
    agent = next(a for a in agents if a['name'] == 'Островок')
    configured_rates = [dict(rate) for rate in agent['rates']]
    assert match_agent('Ostrovok.ru', agents, hotel='dna') == (15, 'Островок')
    agent['commission_zero'] = True
    assert match_agent('Ostrovok.ru', agents, hotel='dna') == (0, 'Островок')
    assert agent['rates'] == configured_rates
    agent['commission_zero'] = False
    assert match_agent('Ostrovok.ru', agents, hotel='dna') == (15, 'Островок')


def test_import_export_roundtrip(tmp_path):
    from agents import default_agents, export_agents, import_agents
    agents = default_agents()
    agents.append({'name': 'Extra Channel', 'aliases': ['extraxyz'], 'hotels': ['dna'],
                   'rates': [{'from': '', 'to': '31.03.2026', 'pct': 12},
                             {'from': '01.04.2026', 'to': '', 'pct': 15}],
                   'commission_zero': True})
    p_x = str(tmp_path / "agents.xlsx")
    p_c = str(tmp_path / "agents.csv")
    export_agents(agents, p_x)
    export_agents(agents, p_c)
    for p in (p_x, p_c):
        back = import_agents(p)
        oz = next(a for a in back if a['name'] == 'Extra Channel')
        assert len(oz['rates']) == 2 and oz['rates'][1]['pct'] == 15
        assert 'extraxyz' in oz['aliases']
        assert oz['commission_zero'] is True
        ost = next(a for a in back if a['name'] == 'Островок')
        assert ost['rates'][0]['pct'] == 15


def test_users_db(tmp_path):
    from db import HistoryDB
    db = HistoryDB(os.path.join(str(tmp_path), "u.db"))
    db.ensure_default_admin()
    assert db.check_user('admin', 'admin') is not None
    assert db.check_user('admin', 'nope') is None
    assert db.add_user('dona_user', 'pw', 'hotel', 'dna')
    assert not db.add_user('dona_user', 'x', 'hotel', 'dna')  # логин занят
    u = db.check_user('dona_user', 'pw')
    assert u['role'] == 'hotel' and u['hotel'] == 'dna'
    admin_id = [x for x in db.list_users() if x['login'] == 'admin'][0]['id']
    assert db.del_user(admin_id) is False  # последний админ неприкосновенен
    du = [x for x in db.list_users() if x['login'] == 'dona_user'][0]['id']
    assert db.del_user(du) is True
    db.set_password(admin_id, 'newpass')
    assert db.check_user('admin', 'newpass') is not None
    db.close()


def test_use_full_off(tmp_path):
    import pandas as pd
    import openpyxl
    import core
    src = os.path.join(str(tmp_path), "uf.csv")
    out = os.path.join(str(tmp_path), "uf.xlsx")
    pd.DataFrame([["Иванов Иван", 8500, "01.08.2026", "04.08.2026", "островок"]],
                 columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник"]
                 ).to_csv(src, index=False)
    _, _ = core.process_dona(src, out, tax={'use_full': False})
    ws = openpyxl.load_workbook(out)["Дона"]
    # налог от исходной стоимости D, а не от пересчитанной G
    assert ws["P2"].value == '=ROUND(IF(O2="мин.ставка",M2*100,I2/102*2),2)'
    assert ws["N2"].value == '=I2/M2'
    assert ws["I2"].value == '=ROUND(D2+J2,2)'


def test_source_map_none_no_highlight(tmp_path):
    """Источник, оставленный «— не указано —» (source_map с 0%), и пустой
    источник не должны подсвечиваться жёлтым как нераспознанные."""
    import pandas as pd
    import openpyxl
    import core
    src = os.path.join(str(tmp_path), "dnone.csv")
    out = os.path.join(str(tmp_path), "dnone.xlsx")
    pd.DataFrame([
        ["Иванов Иван", 8500, "01.08.2026", "04.08.2026", "Mystery Channel"],
        ["Петров Петр", 9000, "01.08.2026", "04.08.2026", ""],
    ], columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник"]
    ).to_csv(src, index=False)
    sm = {"Mystery Channel": (0, "")}
    summary, _ = core.process_dona(src, out, source_map=sm)
    assert summary["Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)"] == 0
    ws = openpyxl.load_workbook(out)["Дона"]
    for row in (2, 3):
        assert ws.cell(row=row, column=5).fill.start_color.rgb in (
            None, "00000000")
    # а без source_map тот же неизвестный источник подсвечивается
    summary2, _ = core.process_dona(src, out)
    assert summary2["Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)"] == 1


def test_source_dialog_search():
    from dialogs_common import (NOT_SET, NEW_AGENT, filter_agent_options,
                                resolve_agent_choice)
    names = ['Островок', 'Отель Онлайн', 'Броневик', 'Яндекс']
    # фильтр: подстрока без учёта регистра, служебные пункты на месте
    assert filter_agent_options('', names) == [NOT_SET] + names + [NEW_AGENT]
    assert filter_agent_options('ост', names) == [NOT_SET, 'Островок', NEW_AGENT]
    assert filter_agent_options('ОТЕЛ', names) == [NOT_SET, 'Отель Онлайн', NEW_AGENT]
    assert filter_agent_options('xxx', names) == [NOT_SET, NEW_AGENT]
    many = ['Агент %03d' % i for i in range(100)]
    assert filter_agent_options('агент', many) == [NOT_SET] + many[:30] + [NEW_AGENT]
    assert filter_agent_options('', many) == [NOT_SET] + many + [NEW_AGENT]
    # разрешение выбора
    assert resolve_agent_choice(NOT_SET, names) == ('none',)
    assert resolve_agent_choice('  ', names) == ('none',)
    assert resolve_agent_choice(NEW_AGENT, names) == ('new',)
    assert resolve_agent_choice('Броневик', names) == ('agent', 'Броневик')
    assert resolve_agent_choice('броневик', names) == ('agent', 'Броневик')
    assert resolve_agent_choice('Яндекс', names) == ('agent', 'Яндекс')
    assert resolve_agent_choice('бронь', names)[0] == 'unknown'
    # «о» подходит сразу нескольким -> неоднозначность, а не угадайка
    assert resolve_agent_choice('о', ['Островок', 'Броневик'])[0] == 'unknown'


def test_default_tax_base_is_full_price(tmp_path):
    """Дефолт: цена в файле — нетто; налог от 100%-й «Стоимости с комиссией»."""
    import pandas as pd
    import openpyxl
    import core
    from tax import TaxParams
    assert TaxParams().use_full is True
    src = os.path.join(str(tmp_path), "dd.csv")
    out = os.path.join(str(tmp_path), "dd.xlsx")
    pd.DataFrame([["Иванов Иван", 8500, "01.08.2026", "04.08.2026", "островок"]],
                 columns=["Гость", "Стоимость", "Заезд", "Выезд", "Источник"]
                 ).to_csv(src, index=False)
    _, _ = core.process_dona(src, out)
    ws = openpyxl.load_workbook(out)["Дона"]
    assert ws["D1"].value == 'Стоимость без комиссии'
    assert ws["G1"].value == 'Стоимость с комиссией'
    assert ws["G2"].value == '=ROUND(D2/(100-E2)*100,2)'
    assert ws["P2"].value == '=ROUND(IF(O2="мин.ставка",M2*100,I2/102*2),2)'
    assert ws["N2"].value == '=I2/M2'
    assert ws["I2"].value == '=ROUND(G2+J2,2)'


def test_discount_zero_and_ostrovok_float():
    """СКИДКА по договору -> 0% во всех отелях; Островок +1 п.п. с 01.08.2026."""
    from agents import default_agents, match_agent
    from datetime import date
    agents = default_agents()
    assert match_agent("Pegas", agents, hotel="dna") == (0, "Хомлин Лэнд (Pegas)")
    assert match_agent("Пегас", agents, hotel="tch")[0] == 0
    assert match_agent("PEGAS", agents, hotel="usd") == (0, "PEGAS [Усадьба]")
    assert match_agent("Алеан", agents, hotel="dna") == (0, "Алеан Отдых")
    assert match_agent("Алеан", agents, hotel="tch")[0] == 0
    assert match_agent("Алеан", agents, hotel="usd") == (0, "ООО Система Алеан [Усадьба]")
    assert match_agent("ООО Туроператор Дельфин", agents, hotel="dna") == (0, "Туроператор Дельфин")
    assert match_agent("ООО Туроператор Дельфин", agents, hotel="tch")[0] == 0
    assert match_agent("Дельфин", agents, hotel="usd") == (0, "ООО Туроператор Дельфин [Усадьба]")
    # обычный Дельфин (Дельфин-Центр) остался платным
    assert match_agent("Дельфин", agents, hotel="dna") == (15, "Дельфин")
    assert match_agent("Дельфин-Центр", agents, hotel="tch") == (15, 'ООО "Дельфин-Центр" [Чайковский]')
    # Островок: 15 до 31.07.2026, 16 с 01.08.2026 (Дона и Чайковский)
    assert match_agent("Ostrovok.ru", agents, date=date(2026, 7, 31), hotel="dna") == (15, "Островок")
    assert match_agent("Ostrovok.ru", agents, date=date(2026, 8, 1), hotel="dna") == (16, "Островок")
    assert match_agent("Ostrovok.ru", agents, date=date(2026, 7, 15), hotel="tch")[0] == 15
    assert match_agent("Ostrovok.ru", agents, date=date(2026, 8, 15), hotel="tch")[0] == 16
    # Усадьба — свой договор 18%, документ его не меняет
    assert match_agent("Ostrovok.ru", agents, date=date(2026, 8, 15), hotel="usd")[0] == 18


def test_db_path_override(tmp_path):
    from db import HistoryDB, set_db_path, default_db_path
    p = os.path.join(str(tmp_path), "shared", "x.db")
    os.makedirs(os.path.dirname(p))
    set_db_path(p)
    try:
        assert default_db_path() == p
        db = HistoryDB()
        db.ensure_default_admin()
        db.close()
        assert os.path.exists(p)
        db2 = HistoryDB(p)
        assert db2.check_user('admin', 'admin') is not None
        db2.close()
    finally:
        set_db_path(None)


def test_travelline_channel_codes_map_to_existing_chai_agents():
    from agents import default_agents, match_agent
    agents = default_agents()
    when = date(2026, 8, 15)
    expected = {
        'TravelLine ANA': 18,
        'TravelLine TIO': 18,
        'TravelLine PGS': 0,
        'TravelLine TNK': 15,
        'TravelLine RTT': 15,
        'TravelLine ALN': 0,
        'TravelLine OL2': 20,
        'TravelLine ACS': 15,
        'TravelLine BGL': 15,
        'TravelLine YAX': 18,
        'TravelLine UTS': 20,
        'TravelLine OZ2': 17,
        'TravelLine H1C': 17,
        'TravelLine OTT': 15,
        'TravelLine CTP': 15,
        'TravelLine CBC': 20,
        'TravelLine OTK': 16,
        'TravelLine BR2': 20,
        'TravelLine JLB': 15,
        'TravelLine сайт': 4,
        'TravelLine ресепшен': 0,
        'TravelLine ресепшн': 0,  # прежнее написание из выгрузок v3.22
    }
    for source, rate in expected.items():
        assert match_agent(source, agents, date=when, hotel='tch')[0] == rate, source
    # CMB is a multi-channel aggregator: don't fuzzy-map it to CBC or another OTA.
    assert match_agent('TravelLine CMB', agents, date=when,
                       hotel='tch') == (0, None)


def test_travelline_mappings_merge_into_saved_agent_list():
    from agents import default_agents, migrate_agents, match_agent
    saved = default_agents()
    direct = {'TravelLine Booking Engine [Чайковский]',
              'Прямые бронирования (стойка) [Чайковский]'}
    for agent in saved:
        agent['aliases'] = [a for a in agent['aliases']
                            if not a.startswith('TravelLine ')]
    saved = [a for a in saved if a['name'] not in direct]
    yandex = next(a for a in saved
                  if a['name'] == 'Яндекс.Путешествия [Чайковский]')
    yandex['rates'] = [{'from': '', 'to': '', 'pct': 12}]

    merged = migrate_agents({'agents': saved})
    assert match_agent('TravelLine YAX', merged, date=date(2026, 8, 15),
                       hotel='tch') == (12, 'Яндекс.Путешествия [Чайковский]')
    assert match_agent('TravelLine сайт', merged, hotel='tch') == (
        4, 'TravelLine Booking Engine [Чайковский]')
    assert match_agent('TravelLine ресепшен', merged, hotel='tch') == (
        0, 'Прямые бронирования (стойка) [Чайковский]')
    assert match_agent('TravelLine ресепшн', merged, hotel='tch') == (
        0, 'Прямые бронирования (стойка) [Чайковский]')


def test_travelline_cache_survives_connection_restart(tmp_path):
    path = str(tmp_path / 'cache-persistent.db')
    key = 'tlcache-v4|tch|8186|2026-07-01|2026-10-01|2026-07-01'
    payload = '{"rows":[{"Гость":"Иван Петров"}],"report":{}}'
    with HistoryDB(path) as db:
        db.put_travelline_cache(
            key, 'tch', '8186', '2026-07-01', '2026-10-01',
            '2026-07-01', 1, payload)
    with HistoryDB(path) as db:
        saved = db.get_travelline_cache(key)
        assert saved['payload_json'] == payload
        assert saved['row_count'] == 1


def test_travelline_cache_crud_and_refresh(tmp_path):
    path = str(tmp_path / 'cache.db')
    key = 'tlcache-v4|tch|8186|2026-07-01|2026-10-01|2026-07-01'
    with HistoryDB(path) as db:
        assert db.get_travelline_cache(key) is None
        db.put_travelline_cache(
            key, 'tch', '8186', '2026-07-01', '2026-10-01',
            '2026-07-01', 2, '{"rows":[1,2],"report":{}}')
        first = db.get_travelline_cache(key)
        assert first['hotel_code'] == 'tch'
        assert first['property_id'] == '8186'
        assert first['period_start'] == '2026-07-01'
        assert first['period_end'] == '2026-10-01'
        assert first['modified_since'] == '2026-07-01'
        assert first['row_count'] == 2
        assert first['payload_json'] == '{"rows":[1,2],"report":{}}'
        assert first['fetched_at']

        # Повторная загрузка обновляет одну запись, не создавая дубликат.
        db.put_travelline_cache(
            key, 'tch', '8186', '2026-07-01', '2026-10-01',
            '2026-07-01', 1, '{"rows":[3],"report":{"updated":true}}')
        refreshed = db.get_travelline_cache(key)
        assert refreshed['row_count'] == 1
        assert refreshed['payload_json'] == (
            '{"rows":[3],"report":{"updated":true}}')
        count = db.con.execute(
            'SELECT COUNT(*) FROM travelline_cache').fetchone()[0]
        assert count == 1

        db.delete_travelline_cache(key)
        assert db.get_travelline_cache(key) is None


def test_travelline_period_snapshot_is_persistent_covering_and_replaceable(tmp_path):
    path = str(tmp_path / 'travelline-periods.db')
    def booking(number, arrival, departure, status='Active'):
        return {
            'number': number, 'status': status,
            'roomStays': [{'stayDates': {
                'arrivalDateTime': f'{arrival}T14:00:00',
                'departureDateTime': f'{departure}T12:00:00'},
                'total': {'priceAfterTax': 1000}}],
        }

    quarter = [
        booking('TL-JUL', '2026-07-10', '2026-07-12'),
        booking('TL-AUG', '2026-08-10', '2026-08-12'),
        booking('TL-SEP', '2026-09-10', '2026-09-12'),
    ]
    with HistoryDB(path) as db:
        assert db.get_travelline_archive_snapshot_bookings(
            'tch', '8186', '2026-08-01', '2026-09-01') is None
        saved = db.replace_travelline_archive_snapshot(
            'tch', '8186', '2026-07-01', '2026-10-01', quarter)
        assert saved['saved'] == 3 and saved['skipped'] == 0
        assert [item['number'] for item in
                db.get_travelline_archive_snapshot_bookings(
                    'tch', '8186', '2026-08-01', '2026-09-01')] == ['TL-AUG']
        stats = db.get_travelline_archive_stats(
            'tch', '8186', '2026-08-01', '2026-09-01')
        assert stats['period_cached'] is True
        assert stats['period_booking_count'] == 1
        assert stats['archive_period_count'] == 1
        assert stats['archive_bytes'] > 0 and stats['database_bytes'] > 0

    # Refreshing the same query replaces the snapshot, not appends stale bookings.
    with HistoryDB(path) as db:
        db.replace_travelline_archive_snapshot(
            'tch', '8186', '2026-07-01', '2026-10-01', [
                booking('TL-JUL', '2026-07-10', '2026-07-12', 'Cancelled'),
                booking('TL-SEP', '2026-09-10', '2026-09-12'),
            ])
        assert db.get_travelline_archive_snapshot_bookings(
            'tch', '8186', '2026-08-01', '2026-09-01') == []
        stats = db.get_travelline_archive_stats(
            'tch', '8186', '2026-08-01', '2026-09-01')
        assert stats['period_cached'] is True
        assert stats['period_booking_count'] == 0


def test_travelline_full_archive_persists_compressed_bookings_and_sync_state(tmp_path):
    path = str(tmp_path / 'travelline-archive.db')
    booking = {
        'number': 'TL-2026-1', 'status': 'Active', 'guest': 'Иван Петров',
        'roomStays': [{
            'stayDates': {'arrivalDateTime': '2026-08-01T14:00:00',
                          'departureDateTime': '2026-08-03T12:00:00'},
            'total': {'priceAfterTax': 12000},
        }],
    }
    inventory = {'accommodationInventory': {'buildings': [{'name': 'Корпус 1'}]}}
    with HistoryDB(path) as db:
        db.begin_travelline_archive_sync('tch', '8186', full=True)
        saved = db.upsert_travelline_archive_bookings('tch', '8186', [booking])
        assert saved == {'saved': 1, 'skipped': 0}
        db.put_travelline_inventory('tch', '8186', inventory)
        partial = db.get_travelline_archive_sync_state('tch', '8186')
        assert partial['full_sync_complete'] == 0
    with HistoryDB(path) as db:
        db.finish_travelline_archive_sync('tch', '8186', full=True)
    with HistoryDB(path) as db:
        stats = db.get_travelline_archive_stats('tch', '8186')
        assert stats['booking_count'] == 1
        assert stats['archive_bytes'] > 0 and stats['database_bytes'] > 0
        assert stats['full_sync_complete'] is True
        archived = db.get_travelline_archive_bookings(
            'tch', '8186', '2026-08-02', '2026-08-03')
        assert archived == [booking]
        assert db.get_travelline_archive_bookings(
            'tch', '8186', '2026-08-03', '2026-08-04') == []
        assert db.get_travelline_inventory('tch', '8186') == inventory
        old_full_sync = stats['last_full_sync_at']
        changed = dict(booking, status='Cancelled')
        db.begin_travelline_archive_sync('tch', '8186', full=False)
        db.upsert_travelline_archive_bookings('tch', '8186', [changed])
        db.finish_travelline_archive_sync('tch', '8186', full=False)
        updated = db.get_travelline_archive_stats('tch', '8186')
        assert updated['booking_count'] == 1
        assert updated['full_sync_complete'] is True
        assert updated['last_full_sync_at'] == old_full_sync
        assert db.get_travelline_archive_bookings(
            'tch', '8186', '2026-08-01', '2026-08-04')[0]['status'] == 'Cancelled'
