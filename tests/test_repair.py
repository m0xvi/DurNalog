# -*- coding: utf-8 -*-
"""Починка битых xlsx: синтетические поломки styles.xml и не-xlsx файлы."""
import os
import sys
import zipfile
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import (WorkbookReadError, _REPAIR_CACHE, excel_serial_to_dt,
                  load_workbook_safe, parse_usd_dt, read_table,
                  repair_notes_for, repair_workbook)

DT = datetime(2026, 7, 15, 14, 0)


def _build_valid(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Гость'
    ws['B1'] = 'Заезд'
    ws['C1'] = 'Сумма'
    ws['A2'] = 'Иванов Иван Иванович'
    ws['B2'] = DT
    ws['B2'].number_format = 'DD.MM.YYYY HH:MM'
    ws['C2'] = 12500.5
    for c in 'ABC':
        ws[f'{c}1'].font = openpyxl.styles.Font(bold=True, color='FF112233')
        ws[f'{c}1'].fill = openpyxl.styles.PatternFill('solid', start_color='FFD9EAD3')
    wb.save(path)
    return path


def _corrupt_styles(src, dst, func):
    with zipfile.ZipFile(src) as z:
        items = {n: z.read(n) for n in z.namelist()}
    assert 'xl/styles.xml' in items
    items['xl/styles.xml'] = func(items['xl/styles.xml'])
    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, d in items.items():
            z.writestr(n, d)
    return dst


def _raw_load_fails(path):
    try:
        openpyxl.load_workbook(path, data_only=True).close()
    except Exception:
        return True
    return False


def _check_values(path):
    wb = load_workbook_safe(path)
    ws = wb.active
    assert ws['A2'].value == 'Иванов Иван Иванович'
    assert ws['C2'].value == 12500.5
    return ws['B2'].value


def test_valid_needs_no_repair(tmp_path):
    p = _build_valid(str(tmp_path / 'ok.xlsx'))
    fixed, notes = repair_workbook(p)
    assert notes == [] and fixed.endswith('.xlsx')
    assert _check_values(fixed) == DT
    assert repair_notes_for([p]) == []
    assert _check_values(p) == DT


def test_short_rgb(tmp_path):
    p = _corrupt_styles(_build_valid(str(tmp_path / 'o.xlsx')),
                        str(tmp_path / 'rgb.xlsx'),
                        lambda b: b.replace(b'FF112233', b'123'))
    assert _raw_load_fails(p)  # premise: openpyxl падает на коротком rgb
    assert _check_values(p) == DT  # tier 1: чинится молча
    assert repair_notes_for([p]) == []


def test_control_char_in_styles(tmp_path):
    p = _corrupt_styles(_build_valid(str(tmp_path / 'o.xlsx')),
                        str(tmp_path / 'ctl.xlsx'),
                        lambda b: b.replace(b'Calibri', b'Cali\x07bri'))
    assert _raw_load_fails(p)
    assert _check_values(p) == DT  # tier 2: чистка без потерь
    assert repair_notes_for([p]) == []


def test_control_char_in_sheet(tmp_path):
    src = _build_valid(str(tmp_path / 'o.xlsx'))
    with zipfile.ZipFile(src) as z:
        items = {n: z.read(n) for n in z.namelist()}
    # ломаем имя гостя в той части, где оно лежит, — styles.xml цел, нужен tier 3
    target = next(n for n in items
                  if n.endswith('.xml') and 'Иванов'.encode() in items[n])
    items[target] = items[target].replace('Иванов'.encode(), 'Ива\x07нов'.encode())
    p = str(tmp_path / 'sheet.xlsx')
    with zipfile.ZipFile(p, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, d in items.items():
            z.writestr(n, d)
    assert _raw_load_fails(p)
    wb = load_workbook_safe(p)  # tier 3: мусорный символ вычищен из данных
    assert wb.active['A2'].value == 'Иванов Иван Иванович'
    assert wb.active['B2'].value == DT
    assert repair_notes_for([p]) == []


def test_broken_tag_salvages_formats(tmp_path):
    p = _corrupt_styles(_build_valid(str(tmp_path / 'o.xlsx')),
                        str(tmp_path / 'tag.xlsx'),
                        lambda b: b.replace(b'</fonts>', b'</fonts_typo>'))
    assert _raw_load_fails(p)
    assert _check_values(p) == DT  # tier 4: даты распознаны, заметка есть
    notes = repair_notes_for([p])
    assert len(notes) == 1 and 'восстановления' in notes[0]


def test_garbage_styles_minimal(tmp_path):
    p = _corrupt_styles(_build_valid(str(tmp_path / 'o.xlsx')),
                        str(tmp_path / 'garbage.xlsx'),
                        lambda b: b'<!!! not xml at all')
    assert _raw_load_fails(p)
    v = _check_values(p)  # tier 5: значения целы, дата стала сериалом
    assert isinstance(v, float)
    assert parse_usd_dt(v) == DT  # страховочная ветка сериалов спасает
    notes = repair_notes_for([p])
    assert len(notes) == 1 and 'СИЛЬНО' in notes[0]


def test_not_a_zip(tmp_path):
    p = str(tmp_path / 'fake.xlsx')
    open(p, 'w', encoding='utf-8').write('<html>not excel</html>')
    try:
        load_workbook_safe(p)
    except WorkbookReadError as e:
        assert 'не является книгой Excel' in str(e)
    else:
        raise AssertionError('псевдо-xlsx обязан дать WorkbookReadError')


def test_zip_without_workbook(tmp_path):
    p = str(tmp_path / 'empty.zip.xlsx')
    with zipfile.ZipFile(p, 'w') as z:
        z.writestr('dummy.txt', b'x')
    try:
        load_workbook_safe(p)
    except WorkbookReadError as e:
        assert 'повреждён' in str(e) and 'Сохранить как' in str(e)
    else:
        raise AssertionError('архив без книги обязан дать WorkbookReadError')


def test_missing_file(tmp_path):
    try:
        load_workbook_safe(str(tmp_path / 'nope.xlsx'))
    except WorkbookReadError as e:
        assert 'не найден' in str(e)
    else:
        raise AssertionError('отсутствующий файл обязан дать WorkbookReadError')


def test_repair_cached(tmp_path):
    p = _corrupt_styles(_build_valid(str(tmp_path / 'o.xlsx')),
                        str(tmp_path / 'c.xlsx'),
                        lambda b: b.replace(b'Calibri', b'Cali\x07bri'))
    n0 = len(_REPAIR_CACHE)
    f1, _ = repair_workbook(p)
    f2, _ = repair_workbook(p)
    assert f1 == f2 and len(_REPAIR_CACHE) == n0 + 1


def test_read_table_repairs(tmp_path):
    p = _corrupt_styles(_build_valid(str(tmp_path / 'o.xlsx')),
                        str(tmp_path / 't.xlsx'),
                        lambda b: b.replace(b'Calibri', b'Cali\x07bri'))
    df = read_table(p)
    assert df.iloc[0, 0] == 'Иванов Иван Иванович'
    assert df.iloc[0, 2] == 12500.5


def test_excel_serial():
    assert excel_serial_to_dt(44927.0) == datetime(2023, 1, 1)
    assert excel_serial_to_dt(5000.0) is None  # цена, а не дата
    assert excel_serial_to_dt(999999.0) is None
    assert excel_serial_to_dt('01.02.2023') == '01.02.2023'  # нечисла как есть
    assert excel_serial_to_dt(True) is True
    assert excel_serial_to_dt(None) is None


def test_parse_usd_dt():
    assert parse_usd_dt(datetime(2026, 1, 2, 3, 4)) == datetime(2026, 1, 2, 3, 4)
    assert parse_usd_dt('15.08.2026 14:00') == datetime(2026, 8, 15, 14, 0)
    assert parse_usd_dt('15.08.2026') == datetime(2026, 8, 15)
    assert parse_usd_dt(45658.0) == datetime(2025, 1, 1)  # сериал
    assert parse_usd_dt(5.0) is None
    assert parse_usd_dt(None) is None
    assert parse_usd_dt('мусор') is None
