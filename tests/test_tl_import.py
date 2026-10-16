# -*- coding: utf-8 -*-
"""Импорт TravelLine: разбор броней в строки Excel (фикстуры, без сети)."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import complete_mapping, guess_mapping, mapping_valid, read_table
from schemas import DNA_DATA_SCHEMA, TCH_DATA_SCHEMA, TCH_ROOMS_SCHEMA
from tl_client import (TL_DATA_MAP, TL_EXTRA, TL_ROOMS_MAP,
                       COL_TL_DISCOUNT, COL_TL_DISCOUNT_PERCENT,
                       COL_TL_DISCOUNT_DATES, COL_TL_DAILY_SUM,
                       COL_TL_DAILY_CHECK, COL_TL_DISCOUNT_CAMPAIGN,
                       COL_TL_DISCOUNT_ADDED, COL_TL_DISCOUNT_RETAINED,
                       COL_TL_DISCOUNT_RULE, COL_TL_TAX_BASE,
                       COL_TL_PERIOD_PRICE, COL_TL_PERIOD_STATUS,
                       ROOM_LABEL_TEXT_PREFIX, inclusive_period_bounds,
                       accommodation_room_map, bookings_to_rows, channel_label,
                       is_tl_export, write_api_inspection_workbook,
                       write_data_workbook, write_rooms_workbook)


def _service(value):
    if isinstance(value, dict):
        service = dict(value)
        amount = service.pop('amount', None)
        if amount is not None:
            service.setdefault('total', {'priceAfterTax': amount})
        return service
    return {'total': {'priceAfterTax': value}}


def _stay(arr='2026-07-10T14:00', dep='2026-07-12T12:00', room=10000.0,
          services=(), guests=(('Иван', 'Петров'),), rt='Стандарт',
          rp='Базовый'):
    return {
        'stayDates': {'arrivalDateTime': arr, 'departureDateTime': dep},
        'total': {'priceAfterTax': room},
        'services': [_service(s) for s in services],
        'extraStayCharges': {'earlyArrival': None, 'lateDeparture': None},
        'guests': [{'firstName': f, 'lastName': l} for f, l in guests],
        'roomType': {'name': rt}, 'ratePlans': [{'name': rp}],
    }


def _book(number='20260701-8186-1', status='Active', stays=None,
          book_services=(), total=10000.0, source=None):
    return {'number': number, 'status': status,
            'roomStays': stays if stays is not None else [_stay()],
            'services': [_service(s) for s in book_services],
            'total': {'priceAfterTax': total},
            'source': source or {'type': 'Channel', 'code': 'YAX'}}

QS, QE = '2026-07-01', '2026-10-01'


def test_plain_booking_no_services():
    rows, rep = bookings_to_rows([_book()], QS, QE)
    assert rep['votes'] == {'A': 0, 'B': 0, '?': 0}  # неинформативна
    assert len(rows) == 1
    r = rows[0]
    assert r['Стоимость'] == 10000.0 and r['Допуслуги'] == 0.0
    assert r['Гость'] == 'Иван Петров' and r['Ночей'] == 2
    assert r['Источник'] == 'TravelLine YAX'
    assert r['№ брони'] == '20260701-8186-1'


def test_discount_night_is_inferred_only_from_unique_daily_rate_match(tmp_path):
    stay = _stay(arr='2026-08-22T14:00', dep='2026-08-27T12:00',
                 room=88448.20)
    stay['dailyRates'] = [
        {'date': '2026-08-22', 'priceBeforeTax': 17248.20},
        {'date': '2026-08-23', 'priceBeforeTax': 17800.00},
        {'date': '2026-08-24', 'priceBeforeTax': 17800.00},
        {'date': '2026-08-25', 'priceBeforeTax': 17800.00},
        {'date': '2026-08-26', 'priceBeforeTax': 17800.00},
    ]
    stay['total'] = {
        'priceBeforeTax': 88448.20,
        'priceAfterTax': 88448.20,
        'discounts': [{
            'amount': 551.80, 'percent': 3.1,
            'discountType': 'Percentage', 'name': 'Скидка первая ночь',
        }],
    }
    booking = _book(stays=[stay], total=88448.20)
    rows, report = bookings_to_rows([booking], QS, QE)

    row = rows[0]
    assert row['Стоимость'] == 88448.20  # исходная цена проживания со скидкой
    assert row['Итог брони'] == 88448.20  # сумма API остаётся без изменений
    assert row[COL_TL_DISCOUNT] == 551.80
    assert row[COL_TL_DISCOUNT_PERCENT] == '3.1%'
    assert row[COL_TL_DISCOUNT_DATES] == '2026-08-22'
    assert row[COL_TL_DAILY_SUM] == 88448.20
    assert row[COL_TL_DAILY_CHECK] == 'Сошлось по ночным тарифам'
    assert row[COL_TL_DISCOUNT_CAMPAIGN] == 'Скидка первая ночь'
    assert row[COL_TL_DISCOUNT_ADDED] == 551.80
    assert row[COL_TL_DISCOUNT_RETAINED] == 0.0
    assert row[COL_TL_DISCOUNT_RULE] == 'Скидка добавлена к налоговой базе'
    assert row[COL_TL_TAX_BASE] == 89000.0
    audit = report['discount_audit']
    assert audit['stays_with_discount'] == 1
    assert audit['nights_with_discount_identified'] == 1
    assert audit['discounts_without_night_attribution'] == 0
    assert audit['discount_base_added'] == 551.80
    assert audit['discount_exception_retained'] == 0.0
    assert audit['daily_rates_matched'] == 1

    import openpyxl

    path = str(tmp_path / 'inspection.xlsx')
    write_api_inspection_workbook(rows, report, path, details=[booking])
    workbook = openpyxl.load_workbook(path, data_only=True)
    data = workbook['Данные']
    assert data['Q1'].value == COL_TL_DISCOUNT
    assert data['Q2'].value == 551.80
    assert data['S2'].value == '2026-08-22'
    assert data['T2'].value == 88448.20
    assert data['V2'].value == 'Скидка первая ночь'
    assert data['W2'].value == 551.80
    assert data['Z1'].value == COL_TL_TAX_BASE
    assert data['Z2'].value == 89000.0
    summary = workbook['Сводка']
    summary_values = {summary.cell(row, 1).value: summary.cell(row, 2).value
                      for row in range(1, summary.max_row + 1)}
    discount_row = next(row for row in range(1, summary.max_row + 1)
                        if summary.cell(row, 1).value == 'Общая скидка TravelLine, ₽')
    assert summary_values['Общая скидка TravelLine, ₽'] == 551.80
    assert summary.cell(discount_row, 2).number_format == '#,##0.00'
    assert summary_values['Добавлено к стоимости для налоговой базы, ₽'] == 551.80
    stays = workbook['Проживания API']
    headers = [cell.value for cell in stays[1]]
    discounts_col = headers.index('Скидки API: название/тип/процент/сумма') + 1
    assert 'Скидка первая ночь' in stays.cell(2, discounts_col).value
    raw_stay = stays['K2'].value
    assert 'Иван' not in raw_stay and 'Петров' not in raw_stay
    assert '<скрыто>' in raw_stay


def test_imported_discount_reaches_tax_base_without_agent_commission(tmp_path):
    import openpyxl
    import core
    from openpyxl.utils import get_column_letter

    stay = _stay(arr='2026-08-22T14:00', dep='2026-08-27T12:00',
                 room=88448.20)
    stay['total'] = {
        'priceBeforeTax': 88448.20, 'priceAfterTax': 88448.20,
        'discounts': [{'amount': 551.80, 'percent': 3.1,
                       'discountType': 'Percentage',
                       'name': 'Скидка первая ночь'}],
    }
    rows, _ = bookings_to_rows(
        [_book(stays=[stay], total=88448.20)], QS, QE)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'TL_tch_2026III_rooms.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)
    write_rooms_workbook(rows, rooms_path)

    core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP, TL_ROOMS_MAP,
        extra_data=TL_EXTRA,
        source_map={'TravelLine YAX': (18, 'Яндекс.Путешествия')},
        tax={'include_extras_in_tax': False})
    report = openpyxl.load_workbook(output_path, data_only=False)['1&2']
    headers = [cell.value for cell in report[1]]
    tax_base_col = get_column_letter(headers.index(COL_TL_TAX_BASE) + 1)
    actual_total_col = get_column_letter(headers.index('Итог брони') + 1)
    assert report['D2'].value == 88448.20
    assert report['E2'].value == 18
    assert report['G2'].value == '=ROUND(D2/(100-E2)*100,2)'
    assert report['I2'].value == (
        f'=ROUND(IF(ISNUMBER({tax_base_col}2),{tax_base_col}2,G2),2)')
    assert report[f'{actual_total_col}2'].value == 88448.20
    assert report[f'{tax_base_col}2'].value == 89000.0


def test_campaign_exception_keeps_discount_in_tax_base():
    stay = _stay(arr='2026-08-22T14:00', dep='2026-08-27T12:00',
                 room=88448.20)
    stay['dailyRates'] = [
        {'date': '2026-08-22', 'priceBeforeTax': 17248.20},
        {'date': '2026-08-23', 'priceBeforeTax': 17800.00},
        {'date': '2026-08-24', 'priceBeforeTax': 17800.00},
        {'date': '2026-08-25', 'priceBeforeTax': 17800.00},
        {'date': '2026-08-26', 'priceBeforeTax': 17800.00},
    ]
    stay['total'] = {
        'priceBeforeTax': 88448.20, 'priceAfterTax': 88448.20,
        'discounts': [{
            'amount': 551.80, 'percent': 3.1,
            'discountType': 'Percentage', 'name': 'Скидка первая ночь',
        }],
    }
    rows, report = bookings_to_rows(
        [_book(stays=[stay], total=88448.20)], QS, QE,
        campaigns_with_discount=['скидка первая ночь'])

    row = rows[0]
    assert row['Стоимость'] == 88448.20
    assert row['Итог брони'] == 88448.20
    assert row[COL_TL_DISCOUNT_ADDED] == 0.0
    assert row[COL_TL_DISCOUNT_RETAINED] == 551.80
    assert row[COL_TL_TAX_BASE] == 88448.20
    assert row[COL_TL_DISCOUNT_RULE] == (
        'Кампания-исключение: скидка сохранена в налоговой базе')
    assert report['discount_audit']['discount_base_added'] == 0.0
    assert report['discount_audit']['discount_exception_retained'] == 551.80
    assert report['discount_audit']['discount_exception_stays'] == 1


def test_ambiguous_discount_is_reported_without_guessing_a_date():
    stay = _stay(arr='2026-07-10T14:00', dep='2026-07-13T12:00', room=180.0)
    stay['dailyRates'] = [
        {'date': '2026-07-10', 'priceBeforeTax': 90.0},
        {'date': '2026-07-11', 'priceBeforeTax': 45.0},
        {'date': '2026-07-12', 'priceBeforeTax': 45.0},
    ]
    stay['total'] = {
        'priceBeforeTax': 180.0, 'priceAfterTax': 180.0,
        'discounts': [{
            'amount': 10.0, 'percent': 10.0,
            'discountType': 'Percentage',
        }],
    }
    rows, report = bookings_to_rows([_book(stays=[stay], total=180.0)], QS, QE)

    row = rows[0]
    assert row[COL_TL_DISCOUNT] == 10.0
    assert row[COL_TL_DISCOUNT_DATES] == 'Не распределена по датам'
    assert report['discount_audit']['nights_with_discount_identified'] == 0
    assert report['discount_audit']['discounts_without_night_attribution'] == 1


def test_daily_rate_reconciliation_accounts_for_included_services():
    stay = _stay(arr='2026-07-10T14:00', dep='2026-07-12T12:00',
                 room=250.0, services=({'name': 'Завтрак', 'amount': 50.0},))
    stay['dailyRates'] = [
        {'date': '2026-07-10', 'priceBeforeTax': 100.0},
        {'date': '2026-07-11', 'priceBeforeTax': 100.0},
    ]
    stay['total'] = {'priceBeforeTax': 250.0, 'priceAfterTax': 250.0}
    rows, report = bookings_to_rows([_book(stays=[stay], total=250.0)], QS, QE)

    assert rows[0][COL_TL_DAILY_SUM] == 200.0
    assert rows[0][COL_TL_DAILY_CHECK] == 'Сошлось с услугами'
    assert report['discount_audit']['daily_rates_matched'] == 1
    assert report['discount_audit']['daily_rates_mismatched'] == 0


def test_daily_rates_are_not_compared_to_after_tax_when_before_tax_exists():
    stay = _stay(arr='2026-07-10T14:00', dep='2026-07-12T12:00', room=200.0)
    stay['dailyRates'] = [
        {'date': '2026-07-10', 'priceBeforeTax': 100.0},
        {'date': '2026-07-11', 'priceBeforeTax': 100.0},
    ]
    stay['total'] = {'priceBeforeTax': 220.0, 'priceAfterTax': 200.0}
    rows, report = bookings_to_rows([_book(stays=[stay], total=200.0)], QS, QE)

    assert rows[0][COL_TL_DAILY_CHECK] == 'Расхождение -20.00 ₽'
    assert report['discount_audit']['daily_rates_mismatched'] == 1


def test_mode_a_with_paid_services():
    b = _book(stays=[_stay(services=(1500.0,))], total=11500.0)
    rows, rep = bookings_to_rows([b], QS, QE)
    assert rep['votes']['A'] == 1 and not rep['mismatch']
    assert rows[0]['Стоимость'] == 10000.0
    assert rows[0]['Допуслуги'] == 1500.0


def test_service_names_are_exported_with_amounts():
    stay = _stay(services=({'name': 'Завтрак', 'amount': 1500.0},))
    stay['extraStayCharges'] = {
        'earlyArrival': {'total': {'priceAfterTax': 100.0}},
        'lateDeparture': None,
    }
    b = _book(stays=[stay],
              book_services=({'name': 'Парковка', 'amount': 900.0},),
              total=12500.0)
    rows, _ = bookings_to_rows([b], QS, QE)
    row = rows[0]
    assert row['Допуслуги'] == 2500.0
    names = row['Наименование допуслуг']
    assert 'Завтрак — 1 500,00 ₽' in names
    assert 'Парковка — 900,00 ₽' in names
    assert 'Ранний заезд — 100,00 ₽' in names
    assert row['Сумма продления номера'] == 0


def test_room_extension_amount_is_separated_from_other_extras():
    stay = _stay(services=(
        {'name': 'Продление номера', 'amount': 1200.0},
        {'name': 'Завтрак', 'amount': 800.0},
    ))
    b = _book(stays=[stay], total=12000.0)
    rows, _ = bookings_to_rows([b], QS, QE)
    assert rows[0]['Допуслуги'] == 2000.0
    assert rows[0]['Наименование допуслуг'] == (
        'Продление номера — 1 200,00 ₽; Завтрак — 800,00 ₽')
    assert rows[0]['Сумма продления номера'] == 1200.0


def test_mode_b_extras_inside():
    b = _book(stays=[_stay(room=11500.0, services=(1500.0,))], total=11500.0)
    rows, rep = bookings_to_rows([b], QS, QE)
    assert rep['votes']['B'] == 1
    assert rows[0]['Стоимость'] == 10000.0  # 11500 − 1500
    assert rows[0]['Допуслуги'] == 1500.0


def test_mismatch_flagged_base_kept():
    b = _book(stays=[_stay(services=(1500.0,))], total=99999.0)
    rows, rep = bookings_to_rows([b], QS, QE)
    assert rep['votes']['?'] == 1
    assert len(rep['mismatch']) == 1
    assert rep['mismatch'][0]['number'] == '20260701-8186-1'
    assert rows[0]['Стоимость'] == 10000.0  # база как есть + флаг


def test_inclusive_zero_neutral():
    b = _book(stays=[_stay(services=(0.0,))], total=10000.0)
    rows, rep = bookings_to_rows([b], QS, QE)
    assert rep['votes'] == {'A': 0, 'B': 0, '?': 0}  # неинформативна
    assert rows[0]['Стоимость'] == 10000.0


def test_cancelled_skipped_and_quarter_filter():
    cancelled = _book(number='c', status='Cancelled')
    out_q = _book(number='o', stays=[_stay(arr='2026-10-05T14:00',
                                           dep='2026-10-07T12:00')])
    rows, rep = bookings_to_rows([cancelled, out_q, _book()], QS, QE)
    assert rep['cancelled'] == 1 and rep['in_quarter'] == 1
    assert [r['№ брони'] for r in rows] == ['20260701-8186-1']


def test_group_booking_two_rows():
    b = _book(stays=[_stay(guests=(('А', 'А'),)),
                     _stay(room=8000.0, guests=(('Б', 'Б'),))],
              total=18000.0)
    rows, rep = bookings_to_rows([b], QS, QE)
    assert len(rows) == 2 and rep['rows'] == 2
    assert [r['Гость'] for r in rows] == ['А А', 'Б Б']


def test_group_booking_three_rooms_uses_each_room_stay_for_tax(tmp_path):
    room_values = [(23, 60876.0), (26, 61176.0), (36, 60876.0)]
    stays = []
    for room_number, amount in room_values:
        stay = _stay(
            arr='2026-08-22T14:00', dep='2026-08-28T12:00', room=amount,
            guests=(('Один', 'Гость'),))
        stay.update({'roomId': f'room-{room_number}',
                     'pmsRoomStayId': f'stay-{room_number}'})
        stay['total']['discounts'] = [{
            'amount': 15219.0, 'percent': 20.0,
            'discountType': 'Percentage', 'name': 'Group campaign',
        }]
        stays.append(stay)
    inventory = {'accommodationInventory': {'buildings': [{
        'id': 'building-1', 'name': 'Корпус 1', 'floors': [{
            'id': 'floor-1', 'rooms': [
                {'id': f'room-{number}', 'displayName': str(number)}
                for number, _ in room_values
            ],
        }],
    }]}}
    group_total = sum(amount for _, amount in room_values)
    rows, report = bookings_to_rows(
        [_book(number='group-3-rooms', stays=stays, total=group_total)],
        QS, QE, accommodation_inventory=inventory)

    assert len(rows) == 3 and report['rows'] == 3
    assert [row['№ комнаты'] for row in rows] == [23, 26, 36]
    assert [row['Стоимость'] for row in rows] == [60876.0, 61176.0, 60876.0]
    assert [row[COL_TL_PERIOD_PRICE] for row in rows] == [
        60876.0, 61176.0, 60876.0]
    assert [row[COL_TL_TAX_BASE] for row in rows] == [76095.0, 76395.0, 76095.0]
    assert [row['Итог брони'] for row in rows] == [group_total] * 3
    assert [row['Ночей'] for row in rows] == [6, 6, 6]

    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'TL_tch_2026III_rooms.xlsx')
    output_path = str(tmp_path / 'group-report.xlsx')
    write_data_workbook(rows, data_path)
    write_rooms_workbook(rows, rooms_path)
    import core
    core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP,
        {'fio': 'ФИО', 'room': '№ комнаты', 'corpus': 'Корпус',
         'bron': '№ брони'},
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    import openpyxl
    report_sheet = openpyxl.load_workbook(output_path, data_only=False)['1']
    assert report_sheet.max_row == 5  # заголовок, три номера и строка итога
    assert [report_sheet.cell(row, 19).value for row in (2, 3, 4)] == [23, 26, 36]
    assert all('IF(ISNUMBER(' in report_sheet.cell(row, 9).value
               for row in (2, 3, 4))


def test_inclusive_period_bounds_include_selected_end_date():
    assert inclusive_period_bounds('2026-09-15', '2026-09-30') == (
        '2026-09-15', '2026-10-01')
    try:
        inclusive_period_bounds('2026-10-01', '2026-09-30')
    except ValueError as exc:
        assert 'раньше даты начала' in str(exc)
    else:
        raise AssertionError('reversed period must fail')


def test_cross_period_stay_is_allocated_by_exact_nightly_rates():
    from datetime import date, timedelta

    stay = _stay(arr='2026-05-15T14:00', dep='2026-06-19T12:00',
                 room=3500.0)
    arrival = date(2026, 5, 15)
    stay['dailyRates'] = [
        {'date': (arrival + timedelta(days=offset)).isoformat(),
         'priceAfterTax': 100.0, 'priceBeforeTax': 100.0}
        for offset in range(35)
    ]
    stay['total'] = {'priceAfterTax': 3500.0, 'priceBeforeTax': 3500.0}
    rows, report = bookings_to_rows(
        [_book(number='cross-period', stays=[stay], total=3500.0)],
        '2026-06-01', '2026-07-01')

    assert len(rows) == 1
    row = rows[0]
    assert row['Ночей'] == 18  # ночи 1–18 июня; checkout 19-го не считается
    assert row['Заезд'].isoformat() == '2026-05-15T14:00:00'
    assert row['Стоимость'] == 3500.0  # API amount kept for commission
    assert row['Итог брони'] == 3500.0
    assert row[COL_TL_PERIOD_PRICE] == 1800.0
    assert row[COL_TL_TAX_BASE] == 1800.0
    assert 'распределено по dailyRates' in row[COL_TL_PERIOD_STATUS]
    assert report['period_allocation']['allocated_partial_stays'] == 1
    assert not report['period_warnings']


def test_cross_period_stay_without_complete_daily_rates_is_skipped_with_warning():
    stay = _stay(arr='2026-05-15T14:00', dep='2026-06-19T12:00',
                 room=3500.0)
    stay['dailyRates'] = [{'date': '2026-05-15', 'priceBeforeTax': 100.0}]
    rows, report = bookings_to_rows(
        [_book(number='missing-rates', stays=[stay], total=3500.0)],
        '2026-06-01', '2026-07-01')

    assert rows == []
    assert report['period_allocation']['skipped_partial_stays'] == 1
    assert report['period_warnings'][0]['number'] == 'missing-rates'
    assert 'dailyRates' in report['period_warnings'][0]['reason']


def test_partial_period_discount_is_restored_only_on_identified_night():
    stay = _stay(arr='2026-05-15T14:00', dep='2026-05-18T12:00', room=290.0)
    stay['dailyRates'] = [
        {'date': '2026-05-15', 'priceAfterTax': 90.0,
         'priceBeforeTax': 90.0},
        {'date': '2026-05-16', 'priceAfterTax': 100.0,
         'priceBeforeTax': 100.0},
        {'date': '2026-05-17', 'priceAfterTax': 100.0,
         'priceBeforeTax': 100.0},
    ]
    stay['total'] = {
        'priceAfterTax': 290.0, 'priceBeforeTax': 290.0,
        'discounts': [{'amount': 10.0, 'percent': 10.0,
                       'discountType': 'Percentage', 'name': 'First-night deal'}],
    }
    rows, report = bookings_to_rows(
        [_book(number='dated-discount', stays=[stay], total=290.0)],
        '2026-05-15', '2026-05-16')

    assert len(rows) == 1
    assert rows[0][COL_TL_PERIOD_PRICE] == 90.0
    assert rows[0][COL_TL_DISCOUNT_DATES] == '2026-05-15'
    assert rows[0][COL_TL_DISCOUNT_ADDED] == 10.0
    assert rows[0][COL_TL_TAX_BASE] == 100.0
    assert not report['period_warnings']


def test_partial_period_keeps_discounted_daily_rate_when_discount_date_unknown():
    stay = _stay(arr='2026-05-15T14:00', dep='2026-05-18T12:00', room=290.0)
    stay['dailyRates'] = [
        {'date': '2026-05-15', 'priceAfterTax': 90.0,
         'priceBeforeTax': 90.0},
        {'date': '2026-05-16', 'priceAfterTax': 100.0,
         'priceBeforeTax': 100.0},
        {'date': '2026-05-17', 'priceAfterTax': 100.0,
         'priceBeforeTax': 100.0},
    ]
    stay['total'] = {
        'priceAfterTax': 290.0, 'priceBeforeTax': 290.0,
        'discounts': [{'amount': 10.0, 'discountType': 'Fixed',
                       'name': 'Unattributed deal'}],
    }
    rows, report = bookings_to_rows(
        [_book(number='unknown-discount-date', stays=[stay], total=290.0)],
        '2026-05-15', '2026-05-16')

    assert len(rows) == 1
    assert rows[0][COL_TL_TAX_BASE] == 90.0
    assert rows[0][COL_TL_DISCOUNT_ADDED] == 0.0
    assert 'не восстановлена' in rows[0][COL_TL_PERIOD_STATUS]
    assert any(warning['type'] == 'скидка'
               for warning in report['period_warnings'])


def test_partial_period_includes_only_dated_extra_services():
    stay = _stay(arr='2026-05-15T14:00', dep='2026-05-17T12:00',
                 room=2000.0,
                 services=({'name': 'Завтрак', 'amount': 50.0,
                            'serviceDate': '2026-05-15'},
                           {'name': 'Услуга вне периода', 'amount': 60.0,
                            'serviceDate': '2026-05-16'},
                           {'name': 'Без даты', 'amount': 70.0}))
    stay['dailyRates'] = [
        {'date': '2026-05-15', 'priceAfterTax': 1000.0,
         'priceBeforeTax': 1000.0},
        {'date': '2026-05-16', 'priceAfterTax': 1000.0,
         'priceBeforeTax': 1000.0},
    ]
    rows, report = bookings_to_rows(
        [_book(number='dated-service', stays=[stay], total=2180.0)],
        '2026-05-15', '2026-05-16')

    assert len(rows) == 1
    assert rows[0]['Допуслуги'] == 50.0
    assert rows[0]['Сумма продления номера'] == 0.0
    assert 'Завтрак — 50,00 ₽' in rows[0]['Наименование допуслуг']
    assert rows[0][COL_TL_TAX_BASE] == 1000.0
    assert 'допуслуги без даты исключены' in rows[0][COL_TL_PERIOD_STATUS]
    assert report['period_allocation']['undated_services'] == 1


def test_channels():
    assert channel_label({'type': 'BookingEngine', 'code': 'PropertySite'}) \
        == 'TravelLine сайт'
    assert channel_label({'type': 'PMS', 'code': None}) == 'TravelLine ресепшен'
    assert channel_label({'type': 'Channel', 'code': 'OTK'}) == 'TravelLine OTK'
    assert channel_label({}) == 'TravelLine ?'


def test_accommodation_inventory_maps_room_display_names_to_buildings():
    inventory = {'accommodationInventory': {'buildings': [
        {'id': 'building-1', 'name': 'Корпус № 1', 'floors': [
            {'id': 'floor-1', 'name': '1 этаж', 'rooms': [
                {'id': 'room-111', 'displayName': '111',
                 'roomTypeId': 'type-1', 'floorId': 'floor-1'},
            ]},
        ]},
        {'id': 'building-2', 'name': 'Main building', 'floors': [
            {'id': 'floor-2', 'name': '2 этаж', 'rooms': [
                {'id': 'room-205', 'displayName': '205',
                 'roomTypeId': 'type-2', 'floorId': 'floor-2'},
                {'id': 'room-001', 'displayName': '001',
                 'roomTypeId': 'type-2', 'floorId': 'floor-2'},
                {'roomId': {'id': ' room-A12 '}, 'name': 'A-12'},
            ]},
        ]},
    ]}}

    mapped = accommodation_room_map(inventory)
    assert mapped['room-111'] == {
        'room': 111, 'corpus': 1, 'building_name': 'Корпус № 1'}
    # Arbitrary building names have no number: use their stable API position.
    assert mapped['room-205'] == {
        'room': 205, 'corpus': 2, 'building_name': 'Main building'}
    assert mapped['room-001']['room'] == '001'
    assert mapped['room-a12'] == {
        'room': 'A-12', 'corpus': 2, 'building_name': 'Main building'}
    assert 'room-111' not in accommodation_room_map({'buildings': []})


def test_flat_property_room_list_maps_api_room_and_display_name_fallback():
    from tl_client import bookings_to_rows

    inventory = {'rooms': [
        {'id': 'room-72', 'displayName': '72', 'floorId': 'floor-a'},
        {'id': 'room-A12', 'displayName': 'A-12', 'floorId': 'floor-b'},
    ]}
    mapped = accommodation_room_map(inventory)
    assert mapped['room-72'] == {
        'room': 72, 'corpus': None, 'building_name': ''}
    assert mapped['room-a12']['room'] == 'A-12'

    # Some TravelLine responses carry the actual room object in `room` while
    # `roomId` is a nonphysical identifier; try both exact API references.
    stay = _stay(guests=(('Иван', 'Петров'),))
    stay.update({'roomId': 'room-type-72',
                 'room': {'id': 'room-72', 'displayName': '72'}})
    rows, report = bookings_to_rows(
        [_book(stays=[stay])], QS, QE, accommodation_inventory=inventory)
    assert rows[0]['№ комнаты'] == 72
    assert report['room_assignments']['assigned'] == 1
    assert report['room_assignments']['assigned_by_room_id'] == 1

    upper_id_stay = _stay()
    upper_id_stay.update({'roomId': 'ROOM-A12'})
    rows, _ = bookings_to_rows(
        [_book(stays=[upper_id_stay])], QS, QE,
        accommodation_inventory=inventory)
    assert rows[0]['№ комнаты'] == 'A-12'

    # If only a unique display name is present, use it but never guess when
    # duplicate room labels exist in different buildings.
    label_stay = _stay(guests=(('Иван', 'Петров'),))
    label_stay.update({'roomId': 'not-an-inventory-id', 'room': 'A-12'})
    rows, report = bookings_to_rows(
        [_book(stays=[label_stay])], QS, QE, accommodation_inventory=inventory)
    assert rows[0]['№ комнаты'] == 'A-12'
    assert report['room_assignments']['assigned_by_display_name'] == 1

    duplicate_labels = {'buildings': [
        {'name': 'Корпус 1', 'rooms': [
            {'id': 'room-a', 'displayName': '101'},
        ]},
        {'name': 'Корпус 2', 'rooms': [
            {'id': 'room-b', 'displayName': '101'},
        ]},
    ]}
    ambiguous_stay = _stay()
    ambiguous_stay.update({'roomId': 'unknown', 'room': '101'})
    rows, report = bookings_to_rows(
        [_book(stays=[ambiguous_stay])], QS, QE,
        accommodation_inventory=duplicate_labels)
    assert rows[0]['№ комнаты'] is None
    assert report['room_assignments']['assigned'] == 0


def test_bookings_rows_keep_real_room_assignment_and_leave_missing_blank():
    assigned = _stay(guests=(('Иван', 'Петров'),))
    assigned.update({'roomId': 'room-72', 'pmsRoomStayId': 'stay-assigned'})
    unassigned = _stay(guests=(('Иван', 'Петров'),))
    unassigned.update({'roomId': None, 'pmsRoomStayId': 'stay-unassigned'})
    unknown = _stay(guests=(('Петр', 'Сидоров'),))
    unknown.update({'roomId': 'not-in-inventory', 'pmsRoomStayId': 'stay-unknown'})
    details = [_book(stays=[assigned, unassigned, unknown], total=30000.0)]
    inventory = {'accommodationInventory': {'buildings': [{
        'name': 'Корпус 1', 'floors': [{'rooms': [
            {'id': ' room-72 ', 'displayName': '', 'roomNumber': '72'},
        ]}],
    }]}}

    rows, report = bookings_to_rows(details, QS, QE, inventory)

    assert [row['№ комнаты'] for row in rows] == [72, None, None]
    assert [row['Корпус'] for row in rows] == [1, None, None]
    assert report['room_assignments']['inventory_rooms'] == 1
    assert report['room_assignments']['assigned'] == 1
    assert report['room_assignments']['unassigned'] == 2
    assert report['room_assignments']['room_id_present'] == 2
    assert report['room_assignments']['room_id_missing'] == 1
    assert report['room_assignments']['room_id_not_in_inventory'] == 1


def test_room_diagnostic_workbook_shows_api_ids_and_match_reasons(tmp_path):
    import openpyxl

    assigned = _stay()
    assigned.update({'roomId': 'room-72', 'pmsRoomStayId': 'stay-assigned'})
    missing = _stay()
    missing.update({'roomId': None, 'pmsRoomStayId': 'stay-missing'})
    unknown = _stay()
    unknown.update({'roomId': 'room-new', 'pmsRoomStayId': 'stay-unknown'})
    details = [_book(stays=[assigned, missing, unknown], total=30000.0)]
    inventory = {'accommodationInventory': {'buildings': [{
        'id': 'building-1', 'name': 'Корпус 1', 'floors': [{
            'id': 'floor-1', 'name': '1 этаж', 'rooms': [
                {'id': 'room-72', 'displayName': '72'},
            ],
        }],
    }]}}
    rows, report = bookings_to_rows(details, QS, QE, inventory)
    stats = report['room_assignments']
    assert stats['room_id_present'] == 2
    assert stats['room_id_missing'] == 1
    assert stats['room_id_not_in_inventory'] == 1
    assert stats['unknown_room_id_samples'] == ['room-new']

    path = str(tmp_path / 'inspection.xlsx')
    write_api_inspection_workbook(
        rows, report, path, details=details,
        accommodation_inventory=inventory)
    wb = openpyxl.load_workbook(path, data_only=True)
    stays = wb['Проживания API']
    headers = [cell.value for cell in stays[1]]
    status_col = headers.index('Результат сопоставления') + 1
    room_id_col = headers.index('roomId') + 1
    values = list(stays.iter_rows(min_row=2, values_only=True))
    assert [row[status_col - 1] for row in values] == [
        'Сопоставлено по ID', 'roomId/room отсутствует',
        'ID не найден в инвентаре']
    assert values[2][room_id_col - 1] == 'room-new'
    assert wb['Номерной фонд API']['F2'].value == 'room-72'
    assert wb['Данные']['N2'].value == 72


def test_tl_room_workbooks_drive_corpus_and_do_not_guess_unassigned_stay(tmp_path):
    import openpyxl
    import core

    first = _stay(guests=(('Иван', 'Петров'),))
    first.update({'roomId': 'room-72', 'pmsRoomStayId': 'stay-1'})
    second = _stay(guests=(('Ольга', 'Сидорова'),))
    second.update({'roomId': 'room-15', 'pmsRoomStayId': 'stay-2'})
    no_room = _stay(guests=(('Иван', 'Петров'),))
    no_room.update({'roomId': None, 'pmsRoomStayId': 'stay-3'})
    details = [
        _book(number='assigned-1', stays=[first]),
        _book(number='assigned-2', stays=[second]),
        _book(number='unassigned', stays=[no_room]),
    ]
    inventory = {'accommodationInventory': {'buildings': [
        {'name': 'Корпус 1', 'floors': [{'rooms': [
            {'id': 'room-72', 'displayName': '72'},
        ]}]},
        {'name': 'Корпус 2', 'floors': [{'rooms': [
            {'id': 'room-15', 'displayName': '15'},
        ]}]},
    ]}}
    rows, _ = bookings_to_rows(details, QS, QE, inventory)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'TL_tch_2026III_rooms.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)
    write_rooms_workbook(rows, rooms_path)

    core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP, TL_ROOMS_MAP,
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    wb = openpyxl.load_workbook(output_path, data_only=False)
    corpus_1 = wb['1']
    corpus_2 = wb['2']
    no_corpus = wb['нет корпуса']
    assert corpus_1['A2'].value == 'assigned-1'
    assert corpus_1['S2'].value == 72 and corpus_1['T2'].value == 1
    assert corpus_2['A2'].value == 'assigned-2'
    assert corpus_2['S2'].value == 15 and corpus_2['T2'].value == 2
    assert no_corpus['A2'].value == 'unassigned'
    assert no_corpus['S2'].value is None and no_corpus['T2'].value is None


def test_manual_room_file_matches_same_booking_by_exact_stay_dates(tmp_path):
    from datetime import datetime
    import openpyxl
    import core

    guest = (('Иван', 'Петров'),)
    first = _stay(arr='2026-07-10T14:00', dep='2026-07-12T12:00',
                  guests=guest)
    first.update({'roomId': None, 'pmsRoomStayId': 'stay-1'})
    second = _stay(arr='2026-07-12T14:00', dep='2026-07-14T12:00',
                   guests=guest)
    second.update({'roomId': None, 'pmsRoomStayId': 'stay-2'})
    rows, _ = bookings_to_rows(
        [_book(number='group-1', stays=[first, second], total=20000)],
        QS, QE)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'rooms.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)

    room_file = openpyxl.Workbook()
    sheet = room_file.active
    sheet.append(['ФИО', '№ комнаты', 'Корпус', '№ брони', 'Заезд', 'Выезд'])
    sheet.append(['Иван Петров', 72, 1, 'group-1',
                  datetime(2026, 7, 10, 14), datetime(2026, 7, 12, 12)])
    sheet.append(['Иван Петров', 15, 2, 'group-1',
                  datetime(2026, 7, 12, 14), datetime(2026, 7, 14, 12)])
    room_file.save(rooms_path)

    core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP,
        {'fio': 'ФИО', 'room': '№ комнаты', 'corpus': 'Корпус',
         'bron': '№ брони', 'checkin': 'Заезд', 'checkout': 'Выезд'},
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    report = openpyxl.load_workbook(output_path, data_only=True)['1&2']
    assert (report['S2'].value, report['T2'].value) == (72, 1)
    assert (report['S3'].value, report['T3'].value) == (15, 2)


def test_manual_rooms_fill_blank_direct_room_from_exported_data(tmp_path):
    import openpyxl
    import core

    stay = _stay(guests=(('Иван', 'Петров'),))
    stay.update({'roomId': None, 'pmsRoomStayId': 'stay-unassigned'})
    rows, _ = bookings_to_rows([_book(stays=[stay])], QS, QE)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'manual_rooms.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)

    manual_rooms = openpyxl.Workbook()
    sheet = manual_rooms.active
    sheet.title = 'Номера'
    sheet.append(['ФИО', '№ комнаты', 'Корпус'])
    sheet.append(['Иван Петров', 72, 1])
    manual_rooms.save(rooms_path)

    summary, _ = core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP,
        {'fio': 'ФИО', 'room': '№ комнаты', 'corpus': 'Корпус'},
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})

    report = openpyxl.load_workbook(output_path, data_only=False)['1']
    assert report['S2'].value == 72 and report['T2'].value == 1
    assert summary['Отсутствует № комнаты (КРАСНЫЙ)'] == 0


def test_manual_rooms_prefer_booking_number_when_guest_repeats(tmp_path):
    import openpyxl
    import core

    stay = _stay(guests=(('Иван', 'Петров'),))
    stay.update({'roomId': None, 'pmsRoomStayId': 'stay-unassigned'})
    data = _book(number='book-1', stays=[stay])
    rows, _ = bookings_to_rows([data], QS, QE)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'manual_rooms.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)

    manual_rooms = openpyxl.Workbook()
    sheet = manual_rooms.active
    sheet.append(['ФИО', '№ комнаты', 'Корпус', '№ брони'])
    sheet.append(['Иван Петров', 55, 2, 'other-book'])
    sheet.append(['Иван Петров', 72, 1, 'book-1'])
    manual_rooms.save(rooms_path)

    core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP,
        {'fio': 'ФИО', 'room': '№ комнаты', 'corpus': 'Корпус',
         'bron': '№ брони'},
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    report = openpyxl.load_workbook(output_path, data_only=False)['1']
    assert report['S2'].value == 72 and report['T2'].value == 1


def test_manual_rooms_can_join_by_booking_without_guest_column(tmp_path):
    import openpyxl
    import core

    stay = _stay(guests=(('Иван', 'Петров'),))
    stay.update({'roomId': None, 'pmsRoomStayId': 'stay-unassigned'})
    rows, _ = bookings_to_rows(
        [_book(number='book-91', stays=[stay])], QS, QE)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'manual_rooms.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)

    manual_rooms = openpyxl.Workbook()
    sheet = manual_rooms.active
    sheet.append(['№ брони', '№ комнаты', 'Корпус'])
    sheet.append(['book-91', 91, 2])
    manual_rooms.save(rooms_path)

    core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP,
        {'fio': None, 'room': '№ комнаты', 'corpus': 'Корпус',
         'bron': '№ брони'},
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    report = openpyxl.load_workbook(output_path, data_only=False)['2']
    assert report['S2'].value == 91 and report['T2'].value == 2


def test_manual_multi_room_booking_uses_unique_guest_after_booking_is_ambiguous(tmp_path):
    import openpyxl
    import core

    first = _stay(guests=(('Иван', 'Петров'),))
    first.update({'roomId': None, 'pmsRoomStayId': 'stay-1'})
    second = _stay(guests=(('Ольга', 'Сидорова'),))
    second.update({'roomId': None, 'pmsRoomStayId': 'stay-2'})
    rows, _ = bookings_to_rows(
        [_book(number='group-booking', stays=[first, second])], QS, QE)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'manual_rooms.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)

    manual_rooms = openpyxl.Workbook()
    sheet = manual_rooms.active
    sheet.append(['ФИО', '№ комнаты', 'Корпус', '№ брони'])
    sheet.append(['Иван Петров', 31, 1, 'group-booking'])
    sheet.append(['Ольга Сидорова', 82, 2, 'group-booking'])
    manual_rooms.save(rooms_path)

    core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP,
        {'fio': 'ФИО', 'room': '№ комнаты', 'corpus': 'Корпус',
         'bron': '№ брони'},
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    wb = openpyxl.load_workbook(output_path, data_only=False)
    assert wb['1']['S2'].value == 31
    assert wb['2']['S2'].value == 82


def test_arrivals_report_columns_auto_map_booking_and_stay_dates():
    mapping = guess_mapping(
        ['№ брони', 'ФИО гостей', 'Заезд', 'Выезд', '№ комнаты'],
        TCH_ROOMS_SCHEMA)
    assert mapping['bron'] == '№ брони'
    assert mapping['fio'] == 'ФИО гостей'
    assert mapping['room'] == '№ комнаты'
    assert mapping['checkin'] == 'Заезд'
    assert mapping['checkout'] == 'Выезд'


def test_legacy_room_mapping_gains_date_columns_without_overriding_choices():
    headers = ['№ брони', 'ФИО гостей', 'Заезд', 'Выезд', '№ комнаты']
    legacy = {'bron': '№ брони', 'fio': 'ФИО гостей',
              'room': '№ комнаты', 'corpus': None}
    updated = complete_mapping(legacy, headers, TCH_ROOMS_SCHEMA)
    assert updated['checkin'] == 'Заезд'
    assert updated['checkout'] == 'Выезд'

    legacy['checkin'] = None
    updated = complete_mapping(legacy, headers, TCH_ROOMS_SCHEMA)
    assert updated['checkin'] is None
    assert updated['checkout'] == 'Выезд'


def test_manual_multi_room_booking_uses_exact_stay_dates(tmp_path):
    import openpyxl
    import core

    first = _stay(arr='2026-07-10T14:00', dep='2026-07-12T12:00',
                  guests=(('Иван', 'Петров'),))
    first.update({'roomId': None, 'pmsRoomStayId': 'stay-1'})
    second = _stay(arr='2026-07-12T14:00', dep='2026-07-14T12:00',
                   guests=(('Ольга', 'Сидорова'),))
    second.update({'roomId': None, 'pmsRoomStayId': 'stay-2'})
    rows, _ = bookings_to_rows(
        [_book(number='group-dates', stays=[first, second])], QS, QE)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'arrivals.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)

    arrivals = openpyxl.Workbook()
    sheet = arrivals.active
    sheet.append(['№ брони', 'Заезд', 'Выезд', '№ комнаты'])
    sheet.append(['group-dates', '10.07.2026, 14:00',
                  '12.07.2026, 12:00', 31])
    sheet.append(['group-dates', '12.07.2026, 14:00',
                  '14.07.2026, 12:00', 82])
    arrivals.save(rooms_path)

    summary, _ = core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP,
        {'fio': None, 'room': '№ комнаты', 'corpus': None,
         'bron': '№ брони', 'checkin': 'Заезд', 'checkout': 'Выезд'},
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    ws = openpyxl.load_workbook(output_path, data_only=False)['1&2']
    rooms_by_guest = {
        ws.cell(row, 3).value: ws.cell(row, 19).value
        for row in range(2, ws.max_row)
    }
    corpora_by_guest = {
        ws.cell(row, 3).value: ws.cell(row, 20).value
        for row in range(2, ws.max_row)
    }
    assert rooms_by_guest == {'Иван Петров': 31, 'Ольга Сидорова': 82}
    # The data workbook has a Corpus column, but TravelLine left its cells
    # blank. Manual rooms still get literal corpora based on their room numbers.
    assert corpora_by_guest == {'Иван Петров': 1, 'Ольга Сидорова': 2}
    assert summary['Отсутствует корпус (КРАСНЫЙ)'] == 0


def test_manual_same_booking_and_stay_dates_remain_ambiguous(tmp_path):
    import openpyxl
    import core

    first = _stay(arr='2026-07-10T14:00', dep='2026-07-12T12:00',
                  guests=(('Иван', 'Петров'),))
    first.update({'roomId': None, 'pmsRoomStayId': 'stay-1'})
    second = _stay(arr='2026-07-10T14:00', dep='2026-07-12T12:00',
                   guests=(('Ольга', 'Сидорова'),))
    second.update({'roomId': None, 'pmsRoomStayId': 'stay-2'})
    rows, _ = bookings_to_rows(
        [_book(number='group-ambiguous', stays=[first, second])], QS, QE)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'arrivals.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)

    arrivals = openpyxl.Workbook()
    sheet = arrivals.active
    sheet.append(['№ брони', 'Заезд', 'Выезд', '№ комнаты'])
    sheet.append(['group-ambiguous', '10.07.2026, 14:00',
                  '12.07.2026, 12:00', 31])
    sheet.append(['group-ambiguous', '10.07.2026, 14:00',
                  '12.07.2026, 12:00', 82])
    arrivals.save(rooms_path)

    summary, _ = core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP,
        {'fio': None, 'room': '№ комнаты', 'corpus': None,
         'bron': '№ брони', 'checkin': 'Заезд', 'checkout': 'Выезд'},
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    ws = openpyxl.load_workbook(output_path, data_only=False)['1&2']
    assert [ws.cell(row, 19).value for row in range(2, ws.max_row)] == [None, None]
    assert summary['Отсутствует № комнаты (КРАСНЫЙ)'] == 2


def test_rooms_workbook_keeps_room_when_api_stay_has_no_guest_name(tmp_path):
    stay = _stay(guests=())
    stay.update({'roomId': 'room-44', 'pmsRoomStayId': 'stay-44'})
    inventory = {'accommodationInventory': {'buildings': [{
        'name': 'Корпус 1', 'floors': [{'rooms': [
            {'id': 'room-44', 'displayName': '44'},
        ]}],
    }]}}
    rows, _ = bookings_to_rows([_book(stays=[stay])], QS, QE, inventory)
    rooms_path = str(tmp_path / 'TL_tch_2026III_rooms.xlsx')
    write_rooms_workbook(rows, rooms_path)

    rooms = read_table(rooms_path)
    assert len(rooms) == 1
    assert rooms['ФИО'].isna().all()
    assert rooms.iloc[0]['№ комнаты'] == 44
    assert rooms.iloc[0]['Корпус'] == 1


def test_string_room_label_survives_excel_and_report_roundtrip(tmp_path):
    import openpyxl
    import core

    stay = _stay(guests=(('Иван', 'Петров'),))
    stay.update({'roomId': 'room-zero', 'pmsRoomStayId': 'stay-zero'})
    inventory = {'accommodationInventory': {'buildings': [{
        'name': 'Корпус 1', 'floors': [{'rooms': [
            {'id': 'room-zero', 'displayName': '001'},
        ]}],
    }]}}
    rows, _ = bookings_to_rows([_book(stays=[stay])], QS, QE, inventory)
    data_path = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rooms_path = str(tmp_path / 'TL_tch_2026III_rooms.xlsx')
    output_path = str(tmp_path / 'report.xlsx')
    write_data_workbook(rows, data_path)
    write_rooms_workbook(rows, rooms_path)

    data = read_table(data_path)
    assert data.iloc[0]['№ комнаты'] == ROOM_LABEL_TEXT_PREFIX + '001'
    core.process_tchaikovsky(
        data_path, rooms_path, output_path, TL_DATA_MAP, TL_ROOMS_MAP,
        extra_data=TL_EXTRA, source_map={'TravelLine YAX': (0, '')})
    report = openpyxl.load_workbook(output_path, data_only=False)['1']
    assert report['S2'].value == '001'
    assert report['T2'].value == 1


def test_workbooks_roundtrip_and_mapping(tmp_path):
    rows, _ = bookings_to_rows([_book()], QS, QE)
    dp = str(tmp_path / 'TL_tch_2026III_data.xlsx')
    rp = str(tmp_path / 'TL_tch_2026III_rooms.xlsx')
    write_data_workbook(rows, dp)
    write_rooms_workbook(rows, rp)
    import openpyxl
    data_sheet = openpyxl.load_workbook(dp, data_only=True)['Брони']
    assert data_sheet['Y1'].value == COL_TL_TAX_BASE
    assert data_sheet['Y2'].value == 10000.0
    assert is_tl_export(dp) and is_tl_export(rp)
    assert not is_tl_export(str(tmp_path / 'other.xlsx'))
    df = read_table(dp)
    assert list(df.columns) == ['Гость', 'Стоимость', 'Заезд', 'Выезд',
                                'Источник', 'Категория номера', '№ брони',
                                'Ночей', 'Тариф', 'Допуслуги', 'Наименование допуслуг',
                                'Сумма продления номера', 'Итог брони',
                                '№ комнаты', 'Корпус',
                                'Скидка TravelLine, ₽',
                                'Процент скидки TravelLine',
                                'Предполагаемые ночи со скидкой TravelLine',
                                'Сумма dailyRates', 'Сверка dailyRates',
                                'Название скидки/кампании TravelLine',
                                'Скидка добавлена к налоговой базе, ₽',
                                'Скидка сохранена в базе, ₽',
                                'Правило скидки TravelLine',
                                'Налоговая база TravelLine, ₽',
                                'Стоимость проживания за выбранный период TravelLine, ₽',
                                'Распределение по периоду TravelLine']
    assert df.iloc[0]['Стоимость'] == 10000.0
    assert df.iloc[0]['Гость'] == 'Иван Петров'
    headers = [str(c) for c in df.columns]
    assert mapping_valid(TL_DATA_MAP, headers, TCH_DATA_SCHEMA)
    assert mapping_valid(TL_DATA_MAP, headers, DNA_DATA_SCHEMA)
    assert all(c in headers for c in TL_EXTRA)
    dr = read_table(rp)
    assert list(dr.columns) == ['ФИО', '№ комнаты', 'Корпус',
                                'Здание TravelLine', '№ брони', 'Заезд', 'Выезд']
    assert len(dr) == 1
    assert dr.iloc[0]['ФИО'] == 'Иван Петров'
    assert dr.iloc[0]['Заезд'].date().isoformat() == '2026-07-10'
    assert dr.iloc[0]['Выезд'].date().isoformat() == '2026-07-12'
    assert dr['№ комнаты'].isna().all() and dr['Корпус'].isna().all()
    assert mapping_valid(TL_ROOMS_MAP, list(dr.columns), TCH_ROOMS_SCHEMA)


def _fake_http_factory(log):
    def fake(method, path, token=None, params=None, data=None, ctype=None):
        log.append(path)
        if path == '/auth/token':
            return 200, {'access_token': 'tok'}
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [{'number': 'n1'},
                                              {'number': 'n2'}],
                         'hasMoreData': False}
        return 200, {'booking': {'number': path.rsplit('/', 1)[-1],
                                 'status': 'Active'}}
    return fake


def test_fetch_details_progress(monkeypatch):
    import threading
    import tl_client
    log = []
    calls = []
    monkeypatch.setattr(tl_client, '_http', _fake_http_factory(log))
    monkeypatch.setattr(tl_client.time, 'sleep', lambda s: None)
    out = tl_client.fetch_details('id', 'sec', '8186', '2024-01-01',
                                  progress=lambda st, d, t: calls.append((st, d, t)),
                                  cancel_event=threading.Event())
    assert [b['number'] for b in out] == ['n1', 'n2']
    assert ('list', 2, 0) in calls
    assert ('details', 2, 2) in calls


def test_fetch_details_full_archive_batches_without_retaining_all_details(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)
    seen, batches = [], []

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/bookings'):
            seen.append(params)
            return 200, {'bookingSummaries': [{'number': f'n{i}'} for i in range(3)],
                         'hasMoreData': False}
        return 200, {'booking': {'number': path.rsplit('/', 1)[-1]}}

    monkeypatch.setattr(tl_client, '_http', fake)
    out = tl_client.fetch_details(
        'id', 'secret', '8186', depth=None, details_callback=batches.append,
        retain_details=False, details_batch_size=2)

    assert out == []
    assert len(batches) == 2
    assert [[b['number'] for b in batch] for batch in batches] == [
        ['n0', 'n1'], ['n2']]
    assert 'lastModification' not in seen[0]
    assert seen[0]['count'] == 1000


def test_fetch_details_uses_summary_number_when_detail_omits_it(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [{'number': 'TL-17'}],
                         'hasMoreData': False}
        return 200, {'booking': {'status': 'Active'}}

    monkeypatch.setattr(tl_client, '_http', fake)
    assert tl_client.fetch_details('id', 'secret', '8186', depth=None) == [
        {'status': 'Active', 'number': 'TL-17'}]


def test_fetch_details_rejects_missing_continuation_token(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client, '_http', lambda *args, **kwargs: (
        200, {'bookingSummaries': [], 'hasMoreData': True}))
    try:
        tl_client.fetch_details('id', 'secret', '8186', depth=None)
    except tl_client.TLApiError as exc:
        assert 'continueToken' in str(exc)
    else:
        raise AssertionError('an incomplete paginated response must not be archived')


def test_fetch_details_searches_only_selected_period_via_pms(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)
    progress = []
    search_calls, detail_calls = [], []
    search_path = '/api/pms/v2/properties/8186/reservations/search'

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path == search_path:
            search_calls.append(dict(params))
            state, page_token = params['state'], params.get('pageToken')
            if state == 'Active' and not page_token:
                return 200, {
                    'reservations': [{'number': 'n1'}, {'number': 'n2'}],
                    'hasNextPage': True, 'nextPageToken': 'active-page-2',
                }
            if state == 'Active':
                return 200, {
                    'reservations': [{'number': 'n3'}], 'hasNextPage': False,
                }
            return 200, {
                'reservations': [{'number': 'n2'}, {'number': 'n4'}],
                'hasNextPage': False,
            }
        if path.startswith('/api/read-reservation/v1/properties/8186/bookings/'):
            number = path.rsplit('/', 1)[-1]
            detail_calls.append(number)
            return 200, {'booking': {'number': number, 'status': 'Active'}}
        raise AssertionError(f'unexpected API request: {path}')

    monkeypatch.setattr(tl_client, '_http', fake)
    out = tl_client.fetch_details(
        'id', 'secret', '8186',
        period_start='2026-07-01', period_end_exclusive='2026-10-01',
        progress=lambda stage, done, total: progress.append((stage, done, total)))

    assert [booking['number'] for booking in out] == ['n1', 'n2', 'n3', 'n4']
    assert detail_calls == ['n1', 'n2', 'n3', 'n4']
    assert len(search_calls) == 3
    assert {call['state'] for call in search_calls} == {'Active', 'Cancelled'}
    assert all(call['startAffectPeriodDateTime'] == '2026-07-01T00:00'
               for call in search_calls)
    assert all(call['endAffectPeriodDateTime'] == '2026-10-01T00:00'
               for call in search_calls)
    assert all(call['maxPageSize'] == 100 for call in search_calls)
    assert all('lastModification' not in call for call in search_calls)
    assert search_calls[1]['pageToken'] == 'active-page-2'
    assert ('list', 4, 0) in progress
    assert ('details', 4, 4) in progress


def test_fetch_details_period_search_requires_next_page_token(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client, '_http', lambda *args, **kwargs: (
        200, {'reservations': [], 'hasNextPage': True}))
    try:
        tl_client.fetch_details(
            'id', 'secret', '8186', period_start='2026-07-01',
            period_end_exclusive='2026-10-01')
    except tl_client.TLApiError as exc:
        assert 'nextPageToken' in str(exc)
    else:
        raise AssertionError('period search must not silently drop a result page')


def test_fetch_details_period_search_403_names_the_required_api(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client, '_http', lambda *args, **kwargs: (
        (_ for _ in ()).throw(tl_client.TLApiError('HTTP 403: forbidden'))))
    try:
        tl_client.fetch_details(
            'id', 'secret', '8186', period_start='2026-07-01',
            period_end_exclusive='2026-10-01')
    except tl_client.TLApiError as exc:
        assert 'PMS API/PropertyReservation' in str(exc)
        assert 'secret' not in str(exc)
    else:
        raise AssertionError('period search should report missing PMS permission')


def test_retry_after_parser_accepts_seconds_and_http_dates():
    from datetime import datetime, timedelta, timezone
    from email.utils import format_datetime
    from tl_client import _retry_after_seconds

    assert _retry_after_seconds('7200') == 7200
    assert _retry_after_seconds('inf') is None
    value = format_datetime(datetime.now(timezone.utc) + timedelta(seconds=30),
                            usegmt=True)
    parsed = _retry_after_seconds(value)
    assert parsed is not None and 0 <= parsed <= 30


def test_fetch_details_retries_rate_limit_and_honors_retry_after(monkeypatch):
    import tl_client
    waits, detail_calls, progress = [], [], []
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', waits.append)

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [{'number': 'n1'}],
                         'hasMoreData': False}
        detail_calls.append(path)
        if len(detail_calls) == 1:
            raise tl_client.TLRateLimitError(
                'HTTP 429: API rate limit exceeded', retry_after_seconds=3)
        return 200, {'booking': {'number': 'n1', 'status': 'Active'}}

    monkeypatch.setattr(tl_client, '_http', fake)
    out = tl_client.fetch_details(
        'id', 'secret', '8186', progress=lambda *args: progress.append(args))

    assert out == [{'number': 'n1', 'status': 'Active'}]
    assert len(detail_calls) == 2
    assert waits == [3]
    assert ('rate_limit', 3, 1) in progress


def test_fetch_details_exponential_rate_limit_backoff(monkeypatch):
    import tl_client
    waits, detail_calls = [], []
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', waits.append)

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [{'number': 'n1'}],
                         'hasMoreData': False}
        detail_calls.append(path)
        if len(detail_calls) <= 2:
            raise tl_client.TLRateLimitError('HTTP 429: API rate limit exceeded')
        return 200, {'booking': {'number': 'n1'}}

    monkeypatch.setattr(tl_client, '_http', fake)
    assert tl_client.fetch_details('id', 'secret', '8186') == [
        {'number': 'n1'}]
    assert waits == [10.0] * 18  # 60 s + 120 s, cancellable in 10 s slices


def test_fetch_details_rate_limit_wait_can_be_cancelled(monkeypatch):
    import threading
    import tl_client
    cancel_event = threading.Event()
    detail_calls, stages, batches = [], [], []
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [{'number': 'n1'}, {'number': 'n2'}],
                         'hasMoreData': False}
        number = path.rsplit('/', 1)[-1]
        detail_calls.append(number)
        if number == 'n1':
            return 200, {'booking': {'number': number}}
        raise tl_client.TLRateLimitError('HTTP 429: API rate limit exceeded')

    def progress(stage, done, total):
        stages.append(stage)
        if stage == 'rate_limit':
            cancel_event.set()

    monkeypatch.setattr(tl_client, '_http', fake)
    try:
        tl_client.fetch_details(
            'id', 'secret', '8186', cancel_event=cancel_event,
            progress=progress, details_callback=batches.append,
            details_batch_size=25)
    except tl_client.TLCancelled:
        pass
    else:
        raise AssertionError('rate-limit backoff should be cancellable')

    assert detail_calls == ['n1', 'n2']
    assert 'rate_limit' in stages
    assert [[booking['number'] for booking in batch] for batch in batches] == [['n1']]


def test_fetch_details_reuses_cached_period_details_and_checkpoints_batches(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)
    detail_calls, batches = [], []

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path == '/api/pms/v2/properties/8186/reservations/search':
            return 200, {
                'reservations': [{'number': 'n1'}, {'number': 'n2'}],
                'hasNextPage': False,
            }
        if path.startswith('/api/read-reservation/v1/properties/8186/bookings/'):
            detail_calls.append(path.rsplit('/', 1)[-1])
            return 200, {'booking': {'number': detail_calls[-1],
                                     'status': 'Active'}}
        raise AssertionError(f'unexpected API request: {path}')

    monkeypatch.setattr(tl_client, '_http', fake)
    out = tl_client.fetch_details(
        'id', 'secret', '8186', period_start='2026-07-01',
        period_end_exclusive='2026-10-01',
        cached_details={'n1': {'number': 'n1', 'status': 'Active'}},
        details_callback=batches.append, details_batch_size=1)

    assert [booking['number'] for booking in out] == ['n1', 'n2']
    assert detail_calls == ['n2']
    assert [[booking['number'] for booking in batch] for batch in batches] == [['n2']]


def test_fetch_details_falls_back_to_paginated_property_rooms(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)
    requested_tokens = []

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/obtain-accommodation-inventory'):
            return 200, {'accommodationInventory': {'buildings': []}}
        if path.endswith('/rooms'):
            requested_tokens.append(params.get('pageToken'))
            if params.get('pageToken') is None:
                return 200, {
                    'rooms': [{'id': 'room-72', 'displayName': '72',
                               'floorId': 'floor-1'}],
                    'hasNextPage': True, 'nextPageToken': 'rooms-page-2',
                }
            return 200, {
                'rooms': [{'id': 'room-73', 'displayName': '73',
                           'floorId': 'floor-1'}],
                'hasNextPage': False,
            }
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [], 'hasMoreData': False}
        raise AssertionError(f'unexpected API request: {path}')

    monkeypatch.setattr(tl_client, '_http', fake)
    details, inventory = tl_client.fetch_details(
        'id', 'secret', '8186', include_room_inventory=True)

    assert details == []
    assert requested_tokens == [None, 'rooms-page-2']
    assert [item['room'] for item in
            tl_client.accommodation_room_map(inventory).values()] == [72, 73]


def test_fetch_details_flushes_downloaded_batch_before_failure(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)
    batches = []

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path == '/api/pms/v2/properties/8186/reservations/search':
            return 200, {
                'reservations': [{'number': 'n1'}, {'number': 'n2'}],
                'hasNextPage': False,
            }
        if path.endswith('/n1'):
            return 200, {'booking': {'number': 'n1'}}
        raise tl_client.TLApiError('HTTP 500: temporary API error')

    monkeypatch.setattr(tl_client, '_http', fake)
    try:
        tl_client.fetch_details(
            'id', 'secret', '8186', period_start='2026-07-01',
            period_end_exclusive='2026-10-01', details_callback=batches.append,
            details_batch_size=25)
    except tl_client.TLApiError:
        pass
    else:
        raise AssertionError('the test booking should fail')
    assert [[booking['number'] for booking in batch] for batch in batches] == [['n1']]


def test_month_week_quarter_presets_are_clipped_to_quarter():
    from tl_client import period_preset_options

    qstart, qend = '2026-04-01', '2026-07-01'
    months = period_preset_options('month', qstart, qend)
    assert [label for label, _, _ in months] == [
        'Апрель 2026', 'Май 2026', 'Июнь 2026']
    assert [(start, end) for _, start, end in months] == [
        ('2026-04-01', '2026-05-01'),
        ('2026-05-01', '2026-06-01'),
        ('2026-06-01', '2026-07-01')]

    weeks = period_preset_options('week', qstart, qend)
    assert len(weeks) == 14
    assert weeks[0][1:] == ('2026-04-01', '2026-04-06')
    assert weeks[-1][1:] == ('2026-06-29', '2026-07-01')
    assert all(qstart <= start < end <= qend for _, start, end in weeks)

    assert period_preset_options('quarter', qstart, qend) == [
        ('Весь квартал', qstart, qend)]


def test_fetch_details_can_also_request_pms_room_inventory(monkeypatch):
    import tl_client
    inventory = {'accommodationInventory': {'buildings': []}}
    seen, progress = [], []
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        seen.append(path)
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [], 'hasMoreData': False}
        if path.endswith('/obtain-accommodation-inventory'):
            return 200, inventory
        if path.endswith('/rooms'):
            return 200, {'rooms': [], 'hasNextPage': False}
        raise AssertionError(f'unexpected API request: {path}')

    monkeypatch.setattr(tl_client, '_http', fake)
    result = tl_client.fetch_details(
        'id', 'sec', '8186', include_room_inventory=True,
        progress=lambda stage, done, total: progress.append((stage, done, total)))

    assert result == ([], inventory)
    assert seen[0] == (
        '/api/pms/v2/properties/8186/obtain-accommodation-inventory')
    assert seen[1].endswith('/rooms')
    assert seen[2].endswith('/bookings')
    assert ('inventory', 0, 1) in progress
    assert ('inventory', 1, 1) in progress


def test_pms_room_inventory_403_explains_permission_without_secret(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [], 'hasMoreData': False}
        raise tl_client.TLApiError('HTTP 403: forbidden')

    monkeypatch.setattr(tl_client, '_http', fake)
    try:
        tl_client.fetch_details(
            'id', 'do-not-log-secret', '8186', include_room_inventory=True)
    except tl_client.TLApiError as exc:
        text = str(exc)
        assert 'PMS API/PropertyRoom' in text
        assert 'номерной фонд' in text.casefold()
        assert 'do-not-log-secret' not in text
    else:
        raise AssertionError('a denied PMS API request should fail explicitly')


def test_cancel_before_start():
    import threading
    import tl_client
    ev = threading.Event()
    ev.set()
    try:
        tl_client.fetch_details('id', 'sec', '8186',
                                cancel_event=ev)
    except tl_client.TLCancelled:
        pass
    else:
        raise AssertionError('установленный флаг обязан дать TLCancelled')


def test_cancel_midway(monkeypatch):
    import threading
    import tl_client
    ev = threading.Event()
    monkeypatch.setattr(tl_client, '_http', _fake_http_factory([]))
    monkeypatch.setattr(tl_client.time, 'sleep', lambda s: None)

    def progress(stage, done, total):
        if stage == 'details':
            ev.set()  # отмена после первой детали

    try:
        tl_client.fetch_details('id', 'sec', '8186', '2024-01-01',
                                progress=progress, cancel_event=ev)
    except tl_client.TLCancelled:
        pass
    else:
        raise AssertionError('отмена на середине обязана дать TLCancelled')


def test_list_401_refreshes_token_and_retries(monkeypatch):
    import tl_client
    tokens = iter(['expired', 'fresh'])
    monkeypatch.setattr(tl_client, 'get_token', lambda *a: next(tokens))
    calls = []

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        calls.append(token)
        if token == 'expired':
            raise tl_client.TLApiError('HTTP 401: {"message":"Unauthorized"}')
        return 200, {'bookingSummaries': [], 'hasMoreData': False}

    monkeypatch.setattr(tl_client, '_http', fake)
    out = tl_client.fetch_details('id', 'secret', '8186')
    assert out == []
    assert calls == ['expired', 'fresh']


def test_persistent_list_401_reports_stage_and_setup_hint(monkeypatch):
    import tl_client
    tokens = iter(['expired', 'fresh'])
    monkeypatch.setattr(tl_client, 'get_token', lambda *a: next(tokens))

    def fake(*args, **kwargs):
        raise tl_client.TLApiError('HTTP 401: {"message":"Unauthorized"}')

    monkeypatch.setattr(tl_client, '_http', fake)
    try:
        tl_client.fetch_details('id', 'do-not-log-secret', '8186')
    except tl_client.TLApiError as exc:
        assert 'Список броней (страница 1)' in str(exc)
        assert 'даже после выдачи нового токена' in str(exc)
        assert 'Read Reservation API' in str(exc)
        assert 'do-not-log-secret' not in str(exc)
    else:
        raise AssertionError('persisting HTTP 401 should raise a diagnostic error')


def test_booking_401_refreshes_token_and_retries(monkeypatch):
    import tl_client
    tokens = iter(['expired', 'fresh'])
    monkeypatch.setattr(tl_client, 'get_token', lambda *a: next(tokens))
    calls = []

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        calls.append((path, token))
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [{'number': 'n1'}],
                         'hasMoreData': False}
        if token == 'expired':
            raise tl_client.TLApiError('HTTP 401: {"message":"Unauthorized"}')
        return 200, {'booking': {'number': 'n1', 'status': 'Active'}}

    monkeypatch.setattr(tl_client, '_http', fake)
    monkeypatch.setattr(tl_client.time, 'sleep', lambda s: None)
    out = tl_client.fetch_details('id', 'secret', '8186')
    assert out == [{'number': 'n1', 'status': 'Active'}]
    assert [token for path, token in calls if path.endswith('/n1')] == [
        'expired', 'fresh']


def test_fetch_details_uses_requested_last_modification_date(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *a: 'tok')
    seen = []

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        seen.append((path, params))
        return 200, {'bookingSummaries': [], 'hasMoreData': False}

    monkeypatch.setattr(tl_client, '_http', fake)
    out = tl_client.fetch_details('id', 'secret', '8186', '2026-07-01')
    assert out == []
    assert seen[0][1]['lastModification'] == '2026-07-01T00:00:00Z'


def test_import_cache_roundtrip_preserves_rows_and_report():
    from datetime import datetime
    from tl_client import decode_tl_import_cache, encode_tl_import_cache

    rows = [{
        'Гость': 'Иван Петров', 'Стоимость': 10000.0,
        'Заезд': datetime(2026, 7, 10, 14, 0),
        'Выезд': datetime(2026, 7, 12, 12, 0),
        'Источник': 'TravelLine YAX', '№ брони': 'TL-42',
        '№ комнаты': 117, 'Корпус': 2,
    }]
    report = {'in_quarter': 1, 'rows': 1, 'cancelled': 0,
              'votes': {'A': 0, 'B': 0, '?': 0}, 'mismatch': []}

    payload = encode_tl_import_cache(rows, report)
    restored_rows, restored_report = decode_tl_import_cache(payload)
    assert restored_rows == rows
    assert isinstance(restored_rows[0]['Заезд'], datetime)
    assert isinstance(restored_rows[0]['Выезд'], datetime)
    assert restored_report == report


def test_import_cache_key_is_scoped_to_property_period_and_query():
    from tl_client import tl_import_cache_key

    common = ('tch', '8186', '2026-07-01', '2026-10-01', '2026-07-01')
    key = tl_import_cache_key(*common)
    assert key.startswith('tlcache-v7|')
    assert tl_import_cache_key(*common) == key
    assert tl_import_cache_key(*common, campaigns_with_discount=['Promo A']) != key
    assert tl_import_cache_key('tch', '479', *common[2:]) != key
    assert tl_import_cache_key('tch', '8186', '2026-10-01', '2027-01-01',
                               '2026-10-01') != key


def test_http_converts_read_timeout_to_diagnostic(monkeypatch):
    import socket
    import tl_client

    def timeout(*args, **kwargs):
        raise socket.timeout('The read operation timed out')

    monkeypatch.setattr(tl_client.urllib.request, 'urlopen', timeout)
    try:
        tl_client._http('GET', '/test')
    except tl_client.TLNetworkError as exc:
        assert 'Тайм-аут ответа TravelLine' in str(exc)
        assert '90 с' in str(exc)
    else:
        raise AssertionError('socket read timeout should be diagnosed')


def test_http_converts_429_and_preserves_retry_after(monkeypatch):
    import io
    import urllib.error
    import tl_client

    def rate_limited(*args, **kwargs):
        raise urllib.error.HTTPError(
            'https://partner.tlintegration.com/test', 429, 'Too Many Requests',
            {'Retry-After': '42'}, io.BytesIO(b'API rate limit exceeded'))

    monkeypatch.setattr(tl_client.urllib.request, 'urlopen', rate_limited)
    try:
        tl_client._http('GET', '/test')
    except tl_client.TLRateLimitError as exc:
        assert 'HTTP 429' in str(exc)
        assert exc.retry_after_seconds == 42
    else:
        raise AssertionError('HTTP 429 should retain rate-limit metadata')


def test_fetch_details_retries_a_transient_timeout(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)
    detail_calls = []

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [{'number': 'n1'}],
                         'hasMoreData': False}
        detail_calls.append(path)
        if len(detail_calls) == 1:
            raise tl_client.TLNetworkError('Тайм-аут ответа TravelLine (90 с)')
        return 200, {'booking': {'number': 'n1', 'status': 'Active'}}

    monkeypatch.setattr(tl_client, '_http', fake)
    out = tl_client.fetch_details('id', 'secret', '8186')
    assert out == [{'number': 'n1', 'status': 'Active'}]
    assert len(detail_calls) == 2


def test_repeated_timeout_reports_api_stage_without_secret(monkeypatch):
    import tl_client
    monkeypatch.setattr(tl_client, 'get_token', lambda *args: 'tok')
    monkeypatch.setattr(tl_client.time, 'sleep', lambda seconds: None)

    def fake(method, path, token=None, params=None, data=None, ctype=None):
        if path.endswith('/bookings'):
            return 200, {'bookingSummaries': [{'number': 'n1'}],
                         'hasMoreData': False}
        raise tl_client.TLNetworkError('Тайм-аут ответа TravelLine (90 с)')

    monkeypatch.setattr(tl_client, '_http', fake)
    try:
        tl_client.fetch_details('id', 'do-not-log-secret', '8186')
    except tl_client.TLApiError as exc:
        assert 'Бронь n1' in str(exc)
        assert '2 попыток' in str(exc)
        assert 'do-not-log-secret' not in str(exc)
    else:
        raise AssertionError('repeated timeout should fail with a stage diagnostic')
