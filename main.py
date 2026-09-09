import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import customtkinter as ctk
import pandas as pd
import re
import io
import os
import sys
import copy
import json
import shutil
import zipfile
import logging
import tempfile
import subprocess
import threading
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

from agents import (match_dona_agent, match_agent, DONA_AGENTS, migrate_agents,
                    default_agents, import_agents, export_agents, parse_dmy)
from tax import (tax_formula_daily_rate, tax_formula_method, tax_formula_tax,
                 parse_price, TaxParams, DEFAULT_TAX)
from db import HistoryDB, set_db_path, default_db_path
import ui_theme as ui
from schemas import TCH_DATA_SCHEMA, TCH_ROOMS_SCHEMA, DNA_DATA_SCHEMA

APP_VERSION = "3.0.1"
NOT_SET = "— не указано —"

# Тема фиксированно светлая (по требованию пользователя)
ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue")


def _app_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


logging.basicConfig(filename=os.path.join(_app_dir(), 'turnalog.log'),
                    level=logging.INFO,
                    format='%(asctime)s %(levelname)s %(message)s')
log = logging.getLogger('turnalog')

# Цвета подсветок
FILL_GUEST = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
FILL_DATE = PatternFill(start_color="FFCC99", end_color="FFCC99", fill_type="solid")
FILL_PRICE = PatternFill(start_color="99CCFF", end_color="99CCFF", fill_type="solid")
FILL_ROOM = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
FILL_DUP = PatternFill(start_color="CC99FF", end_color="CC99FF", fill_type="solid")
FILL_TOTALS = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
DATE_FMT = 'DD.MM.YYYY HH:MM'


# ================================================================
# УТИЛИТЫ
# ================================================================
def norm_name(s):
    """Нормализация ФИО: регистр, ё->е, лишние пробелы."""
    return re.sub(r'\s+', ' ', str(s).lower().replace('ё', 'е')).strip()


def norm_key(s):
    """Нормализация заголовков/строк для нечёткого сопоставления."""
    return re.sub(r'[^0-9a-zа-яё]+', '', str(s).lower().replace('ё', 'е'))


def clean_guest_names(text):
    if pd.isna(text):
        return ""
    return re.sub(r'\d+\.\s*', '', str(text)).strip()


def read_table(path, nrows=None):
    """Чтение .xlsx/.csv с автоподбором кодировки для CSV."""
    if str(path).lower().endswith('.csv'):
        last_err = None
        for enc in ('utf-8-sig', 'utf-8', 'cp1251', 'cp866'):
            try:
                return pd.read_csv(path, encoding=enc, nrows=nrows)
            except UnicodeDecodeError as e:
                last_err = e
        raise ValueError(f"Не удалось определить кодировку CSV: {last_err}")
    return pd.read_excel(path, nrows=nrows)


def read_headers(path):
    return [str(c) for c in read_table(path, nrows=0).columns]


def guess_mapping(headers, schema):
    """Автоподбор колонок файла под поля приложения (точное, затем по подстроке)."""
    nh = {h: norm_key(h) for h in headers}
    taken, mapping = set(), {}
    for f in schema:  # pass 1: точное совпадение синонима
        for h, k in nh.items():
            if h in taken:
                continue
            if any(norm_key(a) == k for a in f['aliases']):
                mapping[f['key']] = h
                taken.add(h)
                break
    for f in schema:  # pass 2: подстрока (самый длинный синоним)
        if f['key'] in mapping:
            continue
        best, bestlen = None, 0
        for h, k in nh.items():
            if h in taken or not k:
                continue
            for a in f['aliases']:
                na = norm_key(a)
                if na and (na in k or k in na) and len(na) > bestlen:
                    best, bestlen = h, len(na)
        if best:
            mapping[f['key']] = best
            taken.add(best)
    return mapping


def mapping_valid(mapping, headers, schema):
    if not mapping:
        return False
    hs = set(headers)
    return all(mapping.get(f['key']) in hs for f in schema if f['required'])


def open_in_system(path):
    try:
        if os.name == 'nt':
            os.startfile(path)
        elif sys.platform == 'darwin':
            subprocess.Popen(['open', path])
        else:
            subprocess.Popen(['xdg-open', path])
    except Exception as e:
        log.exception("open_in_system")
        messagebox.showerror("Ошибка", f"Не удалось открыть файл:\n{e}")


def pack_originals(paths):
    if len(paths) == 1:
        p = paths[0]
        return os.path.basename(p), open(p, 'rb').read()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            z.write(p, os.path.basename(p))
    return 'исходные_файлы.zip', buf.getvalue()


# ================================================================
# КОНФИГ (запоминание сопоставления колонок)
# ================================================================
def _config_path():
    return os.path.join(_app_dir(), 'turnalog_config.json')


def load_config():
    try:
        with open(_config_path(), encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def save_config(cfg):
    try:
        with open(_config_path(), 'w', encoding='utf-8') as f:
            json.dump(cfg, f, ensure_ascii=False, indent=1)
    except Exception:
        log.exception("save_config")


# ================================================================
# ДИАЛОГ СОПОСТАВЛЕНИЯ КОЛОНОК
# ================================================================
class MappingDialog(ctk.CTkToplevel):
    """Выбор соответствия «поле приложения -> колонка файла» + доп. колонки."""

    def __init__(self, master, title, headers, schema, mapping=None, extra=None):
        super().__init__(master)
        self.title(title)
        self.geometry("640x600")
        self.resizable(False, True)
        self.transient(master)
        self.headers = [str(h) for h in headers]
        self.schema = schema
        self.result = None

        ctk.CTkLabel(self, text="Сопоставление колонок",
                     font=ctk.CTkFont(size=16, weight="bold")).pack(pady=(12, 2))
        ctk.CTkLabel(self, text=f"{title}\n"
                                "Укажите, какой колонке файла соответствует каждое поле.\n"
                                "«*» — обязательные поля. Ненужные колонки просто не выбирайте.",
                     text_color="gray", justify="left").pack(padx=20, pady=(0, 8))

        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=15, pady=5)

        scroll = ctk.CTkScrollableFrame(body, fg_color=("gray88", "gray17"))
        scroll.pack(fill="both", expand=True)

        self.vars = {}
        for f in schema:
            row = ctk.CTkFrame(scroll, fg_color="transparent")
            row.pack(fill="x", pady=3, padx=6)
            lbl = f['label'] + ("  *" if f['required'] else "")
            ctk.CTkLabel(row, text=lbl, width=170, anchor="w",
                         font=ctk.CTkFont(size=13)).pack(side="left")
            cur = (mapping or {}).get(f['key'])
            if cur not in self.headers:
                cur = None
            var = tk.StringVar(value=cur or NOT_SET)
            self.vars[f['key']] = var
            ctk.CTkOptionMenu(row, variable=var, width=300,
                              values=[NOT_SET] + self.headers).pack(side="left", padx=6)

        ctk.CTkLabel(self, text="Дополнительные колонки файла — перенести в результат:",
                     anchor="w", font=ctk.CTkFont(size=13, weight="bold")).pack(padx=22, pady=(8, 2), fill="x")
        ex_wrap = ctk.CTkFrame(self, fg_color="transparent")
        ex_wrap.pack(fill="x", padx=20)
        self.extra_vars = {}
        extra = set(extra or [])
        for i, h in enumerate(self.headers):
            v = tk.BooleanVar(value=h in extra)
            self.extra_vars[h] = v
            cb = ctk.CTkCheckBox(ex_wrap, text=h, variable=v,
                                 checkbox_width=16, checkbox_height=16,
                                 font=ctk.CTkFont(size=12))
            cb.grid(row=i // 3, column=i % 3, sticky="w", padx=6, pady=2)
        ex_wrap.columnconfigure((0, 1, 2), weight=1)

        bar = ctk.CTkFrame(self, fg_color="transparent")
        bar.pack(pady=12)
        ctk.CTkButton(bar, text="Автоподбор", width=110, fg_color="gray",
                      command=self.auto).pack(side="left", padx=4)
        ctk.CTkButton(bar, text="OK", width=110, fg_color="green",
                      command=self.ok).pack(side="left", padx=4)
        ctk.CTkButton(bar, text="Отмена", width=110, fg_color="gray",
                      command=self.destroy).pack(side="left", padx=4)
        self.grab_set()
        self.lift()
        self.focus_force()
        self.attributes("-topmost", True)

    def auto(self):
        g = guess_mapping(self.headers, self.schema)
        for key, var in self.vars.items():
            var.set(g.get(key) or NOT_SET)

    def ok(self):
        mapping = {k: (v.get() if v.get() != NOT_SET else None)
                   for k, v in self.vars.items()}
        missing = [f['label'] for f in self.schema
                   if f['required'] and not mapping.get(f['key'])]
        if missing:
            messagebox.showerror("Обязательные поля",
                                 "Не сопоставлены обязательные поля:\n• " +
                                 "\n• ".join(missing), parent=self)
            return
        extra = [h for h, v in self.extra_vars.items() if v.get()]
        self.result = (mapping, extra)
        self.destroy()


class QuarterDialog(ctk.CTkToplevel):
    def __init__(self, master, hotel):
        super().__init__(master)
        self.title("Квартал отчёта")
        self.resizable(False, False)
        self.transient(master)
        self.result = None

        ctk.CTkLabel(self, text=f"Отель «{hotel}»\nЗа какой квартал этот отчёт?",
                     font=ctk.CTkFont(size=14, weight="bold")).pack(pady=(15, 10))
        now = datetime.now()
        self.year = ctk.StringVar(value=str(now.year))
        self.quarter = ctk.StringVar(value=f"Q{(now.month - 1) // 3 + 1}")
        row = ctk.CTkFrame(self, fg_color="transparent")
        row.pack(pady=5)
        ctk.CTkOptionMenu(row, width=110, variable=self.year,
                          values=[str(y) for y in range(now.year - 3, now.year + 2)]).pack(side="left", padx=5)
        ctk.CTkOptionMenu(row, width=110, variable=self.quarter,
                          values=["Q1", "Q2", "Q3", "Q4"]).pack(side="left", padx=5)
        btns = ctk.CTkFrame(self, fg_color="transparent")
        btns.pack(pady=10)
        ctk.CTkButton(btns, text="OK", width=100, fg_color="green",
                      command=self.ok).pack(side="left", padx=5)
        ctk.CTkButton(btns, text="Отмена", width=100, fg_color="gray",
                      command=self.destroy).pack(side="left", padx=5)
        autosize(self, 300)
        self.grab_set()
        self.lift()
        self.focus_force()
        self.attributes("-topmost", True)

    def ok(self):
        self.result = f"{self.year.get()}-{self.quarter.get()}"
        self.destroy()


# ================================================================
# ОКНО НАСТРОЕК (параметры налога по отелям + комиссии агентов)
# ================================================================
class SettingsDialog(ctk.CTkToplevel):
    """Изменяемые пользователем параметры: ставка/минимум/порог/шаблоны
    формул налога для каждого отеля и комиссии агентов Доны."""

    HOTELS = (('tch', 'Чайковский'), ('dna', 'Дона'), ('usd', 'Усадьба'))

    def __init__(self, master, cfg):
        super().__init__(master)
        self.title("Настройки")
        self.geometry("980x700")
        self.minsize(880, 580)
        self.transient(master)
        self.result = None
        self.cfg = cfg
        self.vars = {}
        self.configure(fg_color=ui.COLORS['bg'])

        head = ctk.CTkFrame(self, fg_color="transparent")
        head.pack(fill="x", padx=20, pady=(14, 4))
        ctk.CTkLabel(head, text="Настройки", font=ui.F_PAGE,
                     text_color=ui.COLORS['text']).pack(side="left")
        ctk.CTkLabel(head, text="применяются к новым обработкам",
                     text_color=ui.COLORS['muted'],
                     font=ui.F_SMALL).pack(side="left", padx=(12, 0), pady=(10, 0))

        foot = ctk.CTkFrame(self, fg_color=ui.COLORS['card'], corner_radius=0,
                            border_width=1, border_color=ui.COLORS['line'])
        foot.pack(fill="x", side="bottom")
        ctk.CTkButton(foot, text="Сохранить", width=150, height=38,
                      fg_color=ui.COLORS['green'],
                      hover_color=ui.COLORS['green_hover'],
                      font=ui.F_UI, command=self.ok).pack(side="left", padx=14, pady=10)
        ctk.CTkButton(foot, text="Отмена", width=120, height=38,
                      fg_color="transparent", border_width=1,
                      border_color=ui.COLORS['line'], text_color=ui.COLORS['muted'],
                      hover_color=ui.COLORS['line2'],
                      command=self.destroy).pack(side="left", padx=6, pady=10)

        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=14, pady=6)

        nav = ctk.CTkFrame(body, width=196, corner_radius=12,
                           fg_color=ui.COLORS['card'], border_width=1,
                           border_color=ui.COLORS['line'])
        nav.pack(side="left", fill="y", padx=(0, 8))
        nav.pack_propagate(False)
        self.nav_btns = {}

        def group(title):
            ctk.CTkLabel(nav, text=title, text_color=ui.COLORS['faint'],
                         font=("Segoe UI", 10, "bold"),
                         anchor="w").pack(fill="x", padx=14, pady=(12, 2))

        group("ОТЕЛИ")
        for hid, hname in self.HOTELS:
            self._nav_item(nav, hid, hname)
        group("СПРАВОЧНИКИ")
        self._nav_item(nav, 'agents', 'Агенты')
        if master.is_admin:
            self._nav_item(nav, 'users', 'Пользователи')
            group("СИСТЕМА")
            self._nav_item(nav, 'data', 'База данных')

        self.pane = ctk.CTkFrame(body, fg_color="transparent")
        self.pane.pack(side="left", fill="both", expand=True)

        # секции строятся один раз, переключение — показом/скрытием
        self.section_frames = {}
        for hid, _h in self.HOTELS:
            fr = ctk.CTkFrame(self.pane, fg_color="transparent")
            self._build_hotel_tab(fr, hid)
            self.section_frames[hid] = fr
        fr = ctk.CTkFrame(self.pane, fg_color="transparent")
        self._build_agents_tab(fr)
        self.section_frames['agents'] = fr
        if master.is_admin:
            fr = ctk.CTkFrame(self.pane, fg_color="transparent")
            self._build_users_tab(fr)
            self.section_frames['users'] = fr
            fr = ctk.CTkFrame(self.pane, fg_color="transparent")
            self._build_data_tab(fr)
            self.section_frames['data'] = fr

        self.select_section(self.HOTELS[0][0])
        self.grab_set()
        self.lift()
        self.focus_force()
        self.attributes("-topmost", True)

    def _nav_item(self, nav, sid, label):
        b = ctk.CTkButton(nav, text=label, anchor="w", height=34,
                          fg_color="transparent", text_color=ui.COLORS['text'],
                          hover_color=ui.COLORS['line2'], font=ui.F_UI,
                          command=lambda s=sid: self.select_section(s))
        b.pack(fill="x", padx=8, pady=1)
        self.nav_btns[sid] = b

    def select_section(self, sid):
        for fr in self.section_frames.values():
            fr.pack_forget()
        self.section_frames[sid].pack(fill="both", expand=True)
        for s, b in self.nav_btns.items():
            if s == sid:
                b.configure(fg_color=ui.COLORS['accent_soft'],
                            text_color=ui.COLORS['accent'])
            else:
                b.configure(fg_color="transparent",
                            text_color=ui.COLORS['text'])

    # ------------------------------------------------------------
    def _build_hotel_tab(self, tab, hid):
        cur = TaxParams(self.cfg.get('tax', {}).get(hid))
        v = {}
        self.vars[hid] = v
        wrap = ctk.CTkScrollableFrame(tab, fg_color="transparent")
        wrap.pack(fill="both", expand=True)

        ctk.CTkLabel(wrap, text=f"Отель «{dict(self.HOTELS)[hid]}» — параметры налога",
                     font=ui.F_CARD,
                     text_color=ui.COLORS['text']).pack(padx=6, pady=(4, 10), anchor="w")

        nums = ctk.CTkFrame(wrap, fg_color="transparent")
        nums.pack(fill="x", padx=6, pady=2)
        for i, (key, label, w) in enumerate((('rate', 'Ставка, %', 100),
                                             ('min', 'Минимум за ночь, ₽', 140),
                                             ('threshold', 'Порог, ₽/сутки', 140))):
            f = ctk.CTkFrame(nums, fg_color="transparent")
            f.grid(row=0, column=i, padx=(0, 20), sticky="w", pady=2)
            ctk.CTkLabel(f, text=label, text_color=ui.COLORS['muted'],
                         font=ui.F_SMALL, anchor="w").pack(anchor="w")
            var = tk.StringVar(value=str(getattr(cur, key)))
            v[key] = var
            e = ctk.CTkEntry(f, textvariable=var, width=w, height=34, font=ui.F_UI)
            e.pack()
            if key == 'threshold':
                v['threshold_entry'] = e

        for key, text in (('incl', 'Налог включён в стоимость (делитель 100+ставка)'),
                          ('auto', 'Порог — автоматически: минимум×(100+ставка)/ставка'),
                          ('use_full', 'Считать налог от стоимости БЕЗ комиссии (100%)')):
            var = tk.BooleanVar(value=getattr(cur, key))
            v[key] = var
            ctk.CTkCheckBox(wrap, text=text, variable=var, font=ui.F_UI,
                            fg_color=ui.COLORS['accent'],
                            hover_color=ui.COLORS['line2'],
                            command=self._sync_threshold if key == 'auto' else None
                            ).pack(padx=8, pady=4, anchor="w")
        self._sync_threshold(hid)

        fhead = ctk.CTkFrame(wrap, fg_color="transparent")
        fhead.pack(fill="x", padx=6, pady=(14, 2))
        btn = ctk.CTkButton(fhead, text="▸ Шаблоны Excel-формул", width=200, height=30,
                            fg_color="transparent", border_width=1,
                            border_color=ui.COLORS['line'],
                            text_color=ui.COLORS['muted'], font=ui.F_SMALL,
                            command=lambda h=hid: self._toggle_tpl(h))
        btn.pack(side="left")
        v['tpl_btn'] = btn
        bodyf = ctk.CTkFrame(wrap, fg_color=ui.COLORS['line2'], corner_radius=10)
        for key, label in (('tpl_daily', 'Сут. ставка'), ('tpl_method', 'Метод'),
                           ('tpl_tax', 'Налог')):
            fr = ctk.CTkFrame(bodyf, fg_color="transparent")
            fr.pack(fill="x", pady=2, padx=8)
            ctk.CTkLabel(fr, text=label, width=90, anchor="w",
                         font=ui.F_SMALL).pack(side="left")
            var = tk.StringVar(value=getattr(cur, key))
            v[key] = var
            ctk.CTkEntry(fr, textvariable=var, height=30,
                         font=ui.F_MONO).pack(side="left", fill="x", expand=True)
        ctk.CTkLabel(bodyf, text="Плейсхолдеры: {P} цена · {N} ночей · {R} сут.ставка · "
                                 "{M} метод · {T} порог · {min} · {rate} · {div}",
                     text_color=ui.COLORS['faint'], font=ui.F_TINY, anchor="w",
                     justify="left").pack(padx=8, pady=(2, 6), fill="x")
        bbar = ctk.CTkFrame(bodyf, fg_color="transparent")
        bbar.pack(padx=8, pady=(0, 8))
        ctk.CTkButton(bbar, text="Показать формулы", width=150, height=28,
                      fg_color=ui.COLORS['accent'],
                      hover_color=ui.COLORS['accent_hover'],
                      command=lambda h=hid: self.show_preview(h)).pack(side="left", padx=(0, 6))
        ctk.CTkButton(bbar, text="Сбросить по умолчанию", width=180, height=28,
                      fg_color="transparent", border_width=1,
                      border_color=ui.COLORS['line'], text_color=ui.COLORS['muted'],
                      hover_color=ui.COLORS['card'],
                      command=lambda h=hid: self.reset_tab(h)).pack(side="left")
        v['tpl_frame'] = bodyf
        v['prev'] = ctk.CTkLabel(wrap, text="", justify="left", anchor="w",
                                 font=ui.F_MONO, text_color=ui.COLORS['text'])
        v['prev'].pack(padx=8, pady=(6, 8), fill="x")

    def _sync_threshold(self, hid=None):
        for h, v in self.vars.items():
            if hid and h != hid:
                continue
            if 'threshold_entry' in v:
                v['threshold_entry'].configure(
                    state="disabled" if v['auto'].get() else "normal")

    def _toggle_tpl(self, hid):
        v = self.vars[hid]
        if v['tpl_frame'].winfo_manager():
            v['tpl_frame'].pack_forget()
            v['tpl_btn'].configure(text="▸ Шаблоны Excel-формул")
        else:
            v['tpl_frame'].pack(fill="x", padx=6, pady=4)
            v['tpl_btn'].configure(text="▾ Шаблоны Excel-формул")

    def _cur_params(self, hid):
        v = self.vars[hid]
        return TaxParams({
            'rate': float(v['rate'].get()),
            'min_per_night': float(v['min'].get()),
            'threshold': float(v['threshold'].get()),
            'incl': v['incl'].get(),
            'auto_threshold': v['auto'].get(),
            'use_full': v['use_full'].get(),
            'tpl_daily': v['tpl_daily'].get(),
            'tpl_method': v['tpl_method'].get(),
            'tpl_tax': v['tpl_tax'].get(),
        })

    def show_preview(self, hid):
        try:
            tp = self._cur_params(hid)
        except ValueError:
            messagebox.showerror("Настройки", "Числовые поля содержат нечисловые значения.",
                                 parent=self)
            return
        self.vars[hid]['prev'].configure(text=tp.preview())

    def reset_tab(self, hid):
        v = self.vars[hid]
        d = DEFAULT_TAX
        v['rate'].set(str(d['rate']))
        v['min'].set(str(d['min_per_night']))
        v['threshold'].set(str(d['threshold']))
        v['incl'].set(d['incl'])
        v['auto'].set(d['auto_threshold'])
        v['use_full'].set(d['use_full'])
        v['tpl_daily'].set(d['tpl_daily'])
        v['tpl_method'].set(d['tpl_method'])
        v['tpl_tax'].set(d['tpl_tax'])
        self.show_preview(hid)

    # ------------------------------------------------------------
    def _build_agents_tab(self, tab):
        import copy as _copy
        self.agents_edit = _copy.deepcopy(migrate_agents(self.cfg))
        self.cur_agent = -1

        ctk.CTkLabel(tab, text="Агенты: комиссии (в т.ч. по периодам), синонимы, отели",
                     font=ctk.CTkFont(size=13, weight="bold")).pack(padx=14, pady=(8, 2), anchor="w")

        bar = ctk.CTkFrame(tab, fg_color="transparent")
        bar.pack(fill="x", padx=14, pady=2)
        ctk.CTkButton(bar, text="➕ Добавить", width=100, height=28, fg_color="green",
                      font=ctk.CTkFont(size=11), command=self.add_agent).pack(side="left", padx=3)
        ctk.CTkButton(bar, text="✕ Удалить", width=90, height=28, fg_color="#B33",
                      font=ctk.CTkFont(size=11), command=self.del_agent).pack(side="left", padx=3)
        ctk.CTkButton(bar, text="Импорт…", width=90, height=28, fg_color="gray",
                      font=ctk.CTkFont(size=11), command=self.import_agents_file).pack(side="left", padx=3)
        ctk.CTkButton(bar, text="Экспорт…", width=90, height=28, fg_color="gray",
                      font=ctk.CTkFont(size=11), command=self.export_agents_file).pack(side="left", padx=3)
        ctk.CTkButton(bar, text="Проверить", width=90, height=28, fg_color="gray",
                      font=ctk.CTkFont(size=11), command=self.test_agent).pack(side="left", padx=3)

        tfr = ctk.CTkFrame(tab, fg_color="transparent")
        tfr.pack(fill="x", padx=14, pady=2)
        self.agent_test_var = tk.StringVar()
        ctk.CTkEntry(tfr, textvariable=self.agent_test_var, height=26,
                     font=ctk.CTkFont(size=11),
                     placeholder_text="вставьте «Источник» из отчёта…").pack(
            side="left", fill="x", expand=True, padx=(0, 6))
        self.agent_test_res = ctk.CTkLabel(tfr, text="", width=250, anchor="w",
                                           font=ctk.CTkFont(size=11, weight="bold"))
        self.agent_test_res.pack(side="left")

        body = ctk.CTkFrame(tab, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=14, pady=4)
        left = ctk.CTkFrame(body, fg_color="transparent")
        left.pack(side="left", fill="y")
        self.ag_list = tk.Listbox(left, width=33, height=22, font=("Segoe UI", 10),
                                  activestyle="none")
        ag_sb = ttk.Scrollbar(left, orient="vertical", command=self.ag_list.yview)
        self.ag_list.configure(yscrollcommand=ag_sb.set)
        self.ag_list.pack(side="left", fill="y", expand=True)
        ag_sb.pack(side="left", fill="y")
        self.ag_list.bind("<<ListboxSelect>>", lambda e: self.on_agent_select())

        right = ctk.CTkScrollableFrame(body, fg_color=("gray88", "gray17"))
        right.pack(side="left", fill="both", expand=True, padx=(8, 0))

        def lrow(label):
            fr = ctk.CTkFrame(right, fg_color="transparent")
            fr.pack(fill="x", pady=2, padx=6)
            ctk.CTkLabel(fr, text=label, width=105, anchor="w",
                         font=ctk.CTkFont(size=12)).pack(side="left")
            return fr

        fr = lrow("Название")
        self.ag_name_var = tk.StringVar()
        ctk.CTkEntry(fr, textvariable=self.ag_name_var, height=28,
                     font=ctk.CTkFont(size=12)).pack(side="left", fill="x", expand=True)

        fr = lrow("Применяется к")
        self.ag_hot_vars = {}
        for hid, hname in self.HOTELS:
            v = tk.BooleanVar()
            self.ag_hot_vars[hid] = v
            ctk.CTkCheckBox(fr, text=hname, variable=v, checkbox_width=15,
                            checkbox_height=15,
                            font=ctk.CTkFont(size=12)).pack(side="left", padx=4)

        fr = lrow("Синонимы")
        self.ag_alias_var = tk.StringVar()
        ctk.CTkEntry(fr, textvariable=self.ag_alias_var, height=28,
                     font=ctk.CTkFont(size=11)).pack(side="left", fill="x", expand=True)

        ctk.CTkLabel(right, text="Комиссия по периодам (пустые даты = базовая ставка; "
                                 "берётся период, содержащий дату заезда):",
                     anchor="w", font=ctk.CTkFont(size=11, weight="bold"),
                     justify="left").pack(padx=6, pady=(8, 2), fill="x")
        self.ag_rates_box = ctk.CTkFrame(right, fg_color="transparent")
        self.ag_rates_box.pack(fill="x", padx=6)
        self.ag_rate_vars = []
        ctk.CTkButton(right, text="+ период", width=90, height=24, fg_color="gray",
                      font=ctk.CTkFont(size=10),
                      command=lambda: self.add_rate_row()).pack(padx=6, pady=4, anchor="w")

        self.refresh_agent_list(select=0)

    # ------------------------------------------------------------
    # агент: список / детали
    # ------------------------------------------------------------
    def refresh_agent_list(self, select=None):
        self.ag_list.delete(0, "end")
        for a in self.agents_edit:
            self.ag_list.insert("end", a['name'])
        if select is not None and 0 <= select < len(self.agents_edit):
            self.ag_list.selection_set(select)
            self.load_agent(select)

    def on_agent_select(self):
        sel = self.ag_list.curselection()
        if not sel or sel[0] == self.cur_agent:
            return
        self.save_agent_details()
        self.load_agent(sel[0])

    def load_agent(self, idx):
        self.cur_agent = idx
        a = self.agents_edit[idx]
        self.ag_name_var.set(a['name'])
        for hid, v in self.ag_hot_vars.items():
            v.set(hid in (a.get('hotels') or []))
        self.ag_alias_var.set(", ".join(a.get('aliases') or []))
        for t in self.ag_rate_vars:
            t[3].destroy()
        self.ag_rate_vars = []
        for r in (a.get('rates') or [{'from': '', 'to': '', 'pct': 0}]):
            self.add_rate_row(r)

    def add_rate_row(self, r=None):
        r = r or {'from': '', 'to': '', 'pct': ''}
        fr = ctk.CTkFrame(self.ag_rates_box, fg_color="transparent")
        fr.pack(fill="x", pady=1)
        fv = tk.StringVar(value=r.get('from', '') or '')
        tv = tk.StringVar(value=r.get('to', '') or '')
        pv = tk.StringVar(value=str(r.get('pct', '')))
        ctk.CTkLabel(fr, text="с", font=ctk.CTkFont(size=11)).pack(side="left")
        ctk.CTkEntry(fr, textvariable=fv, width=88, height=24,
                     placeholder_text="дд.мм.гггг",
                     font=ctk.CTkFont(size=11)).pack(side="left", padx=2)
        ctk.CTkLabel(fr, text="по", font=ctk.CTkFont(size=11)).pack(side="left")
        ctk.CTkEntry(fr, textvariable=tv, width=88, height=24,
                     placeholder_text="дд.мм.гггг",
                     font=ctk.CTkFont(size=11)).pack(side="left", padx=2)
        ctk.CTkLabel(fr, text="%", font=ctk.CTkFont(size=11)).pack(side="left", padx=(4, 2))
        ctk.CTkEntry(fr, textvariable=pv, width=55, height=24,
                     font=ctk.CTkFont(size=11)).pack(side="left")
        ctk.CTkButton(fr, text="✕", width=26, height=24, fg_color="#B33",
                      command=lambda f=fr: self.del_rate_row(f)).pack(side="left", padx=4)
        self.ag_rate_vars.append((fv, tv, pv, fr))

    def del_rate_row(self, fr):
        if len(self.ag_rate_vars) <= 1:
            return
        self.ag_rate_vars = [t for t in self.ag_rate_vars if t[3] is not fr]
        fr.destroy()

    def save_agent_details(self):
        if self.cur_agent is None or self.cur_agent < 0:
            return
        a = self.agents_edit[self.cur_agent]
        old = a['name']
        a['name'] = self.ag_name_var.get().strip() or old
        a['hotels'] = [h for h, v in self.ag_hot_vars.items() if v.get()] or ['dna']
        a['aliases'] = [x.strip() for x in re.split(r'[,;]+', self.ag_alias_var.get())
                        if x.strip()]
        rates = []
        for fv, tv, pv, fr in self.ag_rate_vars:
            try:
                pct = float(pv.get().strip().replace(',', '.'))
            except ValueError:
                pct = 0.0
            rates.append({'from': fv.get().strip(), 'to': tv.get().strip(), 'pct': pct})
        a['rates'] = rates or [{'from': '', 'to': '', 'pct': 0}]
        if a['name'] != old:
            self.ag_list.delete(self.cur_agent)
            self.ag_list.insert(self.cur_agent, a['name'])
            self.ag_list.selection_set(self.cur_agent)

    def add_agent(self):
        self.save_agent_details()
        self.agents_edit.append({'name': 'Новый агент', 'aliases': [],
                                 'hotels': ['dna'],
                                 'rates': [{'from': '', 'to': '', 'pct': 0}]})
        self.cur_agent = -1
        self.refresh_agent_list(select=len(self.agents_edit) - 1)

    def del_agent(self):
        sel = self.ag_list.curselection()
        if not sel:
            return
        if not messagebox.askyesno("Удалить агента",
                                   f"Удалить «{self.agents_edit[sel[0]]['name']}»?",
                                   parent=self):
            return
        del self.agents_edit[sel[0]]
        self.cur_agent = -1
        self.refresh_agent_list(select=0)

    def import_agents_file(self):
        self.save_agent_details()
        path = filedialog.askopenfilename(parent=self, title="Импорт агентов",
                                          filetypes=[("Таблицы", "*.xlsx *.csv")])
        if not path:
            return
        try:
            imported = import_agents(path)
        except Exception as e:
            messagebox.showerror("Импорт", f"Не удалось импортировать:\n{e}", parent=self)
            return
        by_name = {a['name']: a for a in self.agents_edit}
        added = replaced = 0
        for a in imported:
            if a['name'] in by_name:
                by_name[a['name']].update(a)
                replaced += 1
            else:
                self.agents_edit.append(a)
                added += 1
        self.cur_agent = -1
        self.refresh_agent_list(select=0)
        messagebox.showinfo("Импорт", f"Добавлено: {added}, заменено: {replaced}",
                            parent=self)

    def export_agents_file(self):
        self.save_agent_details()
        path = filedialog.asksaveasfilename(
            parent=self, title="Экспорт агентов", defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv")])
        if not path:
            return
        try:
            export_agents(self.agents_edit, path)
            messagebox.showinfo("Экспорт", f"Сохранено:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Экспорт", f"Не удалось сохранить:\n{e}", parent=self)

    def test_agent(self):
        self.save_agent_details()
        src = self.agent_test_var.get()
        if not src.strip():
            return
        comm, name = match_agent(src, self.agents_edit)
        self.agent_test_res.configure(
            text=f"→ {name}, {comm}%" if name else "→ не распознан (0%)")

    # ------------------------------------------------------------
    # вкладка «Пользователи» (админ)
    # ------------------------------------------------------------
    def _build_users_tab(self, tab):
        ctk.CTkLabel(tab, text="Пользователи и права доступа",
                     font=ctk.CTkFont(size=13, weight="bold")).pack(
            padx=14, pady=(8, 2), anchor="w")
        ctk.CTkLabel(tab, text="Пользователь отеля видит только свой отель и его историю.",
                     text_color="gray", font=ctk.CTkFont(size=10)).pack(
            padx=14, anchor="w")
        bar = ctk.CTkFrame(tab, fg_color="transparent")
        bar.pack(padx=14, pady=6, fill="x")
        ctk.CTkButton(bar, text="➕ Добавить", width=110, height=30, fg_color="green",
                      font=ctk.CTkFont(size=12), command=self._add_user).pack(side="left", padx=3)
        ctk.CTkButton(bar, text="Сменить пароль", width=140, height=30,
                      font=ctk.CTkFont(size=12), command=self._change_pw).pack(side="left", padx=3)
        ctk.CTkButton(bar, text="✕ Удалить", width=100, height=30, fg_color="#B33",
                      font=ctk.CTkFont(size=12), command=self._del_user).pack(side="left", padx=3)

        cols = ("login", "role", "hotel")
        self.users_tree = ttk.Treeview(tab, columns=cols, show="headings", height=10)
        for c, h, w in (("login", "Логин", 180), ("role", "Роль", 140),
                        ("hotel", "Отель", 160)):
            self.users_tree.heading(c, text=h)
            self.users_tree.column(c, width=w, anchor="w")
        self.users_tree.pack(padx=14, pady=(0, 10), fill="both", expand=True)
        self._refresh_users()

    def _refresh_users(self):
        for i in self.users_tree.get_children():
            self.users_tree.delete(i)
        with HistoryDB() as db:
            for u in db.list_users():
                self.users_tree.insert("", "end", iid=str(u["id"]), values=(
                    u["login"],
                    "администратор" if u["role"] == "admin" else "пользователь",
                    HOTEL_LABELS.get(u["hotel"], "—")))

    def _sel_user(self):
        sel = self.users_tree.selection()
        return int(sel[0]) if sel else None

    def _add_user(self):
        dlg = ctk.CTkToplevel(self)
        dlg.title("Новый пользователь")
        dlg.transient(self)
        dlg.resizable(False, False)
        pad = dict(padx=24, fill="x")

        ctk.CTkLabel(dlg, text="Логин", font=ctk.CTkFont(size=12)).pack(
            anchor="w", pady=(14, 0), **pad)
        lv = tk.StringVar()
        ctk.CTkEntry(dlg, textvariable=lv, height=30).pack(**pad)
        ctk.CTkLabel(dlg, text="Пароль", font=ctk.CTkFont(size=12)).pack(
            anchor="w", pady=(8, 0), **pad)
        pv = tk.StringVar()
        ctk.CTkEntry(dlg, textvariable=pv, height=30, show="*").pack(**pad)
        ctk.CTkLabel(dlg, text="Роль", font=ctk.CTkFont(size=12)).pack(
            anchor="w", pady=(8, 0), **pad)
        rv = tk.StringVar(value="пользователь отеля")
        ctk.CTkOptionMenu(dlg, variable=rv, height=30,
                          values=["пользователь отеля", "администратор"]).pack(**pad)
        ctk.CTkLabel(dlg, text="Отель", font=ctk.CTkFont(size=12)).pack(
            anchor="w", pady=(8, 0), **pad)
        hv = tk.StringVar(value=HOTEL_LABELS['tch'])
        ctk.CTkOptionMenu(dlg, variable=hv, height=30,
                          values=list(HOTEL_LABELS.values())).pack(**pad)

        def ok():
            role = 'admin' if rv.get() == 'администратор' else 'hotel'
            hotel = None if role == 'admin' else \
                [k for k, v in HOTEL_LABELS.items() if v == hv.get()][0]
            with HistoryDB() as db:
                if not db.add_user(lv.get().strip(), pv.get(), role, hotel):
                    messagebox.showerror("Пользователи",
                                         "Логин занят или поля пустые.", parent=dlg)
                    return
            self._refresh_users()
            dlg.destroy()

        ctk.CTkButton(dlg, text="Создать", fg_color="green", height=32,
                      command=ok).pack(pady=14, **pad)
        dlg.grab_set()
        autosize(dlg, 350)
        dlg.lift()
        dlg.focus_force()

    def _change_pw(self):
        uid = self._sel_user()
        if uid is None:
            messagebox.showinfo("Пользователи", "Выберите пользователя в списке.",
                                parent=self)
            return
        dlg = ctk.CTkToplevel(self)
        dlg.title("Новый пароль")
        dlg.transient(self)
        dlg.resizable(False, False)
        pv = tk.StringVar()
        ctk.CTkEntry(dlg, textvariable=pv, height=30, show="*",
                     placeholder_text="Новый пароль").pack(padx=24, pady=(16, 6), fill="x")

        def ok():
            if not pv.get():
                return
            with HistoryDB() as db:
                db.set_password(uid, pv.get())
            dlg.destroy()

        ctk.CTkButton(dlg, text="Сохранить", fg_color="green", height=30,
                      command=ok).pack(padx=24, fill="x", pady=(0, 14))
        dlg.grab_set()
        autosize(dlg, 300)
        dlg.lift()
        dlg.focus_force()

    def _del_user(self):
        uid = self._sel_user()
        if uid is None:
            messagebox.showinfo("Пользователи", "Выберите пользователя в списке.",
                                parent=self)
            return
        login = self.users_tree.item(str(uid))['values'][0]
        if not messagebox.askyesno("Удалить пользователя", f"Удалить «{login}»?",
                                   parent=self):
            return
        with HistoryDB() as db:
            if not db.del_user(uid):
                messagebox.showerror("Пользователи",
                                     "Нельзя удалить последнего администратора.",
                                     parent=self)
                return
        self._refresh_users()

    # ------------------------------------------------------------
    # вкладка «Данные»: расположение базы
    # ------------------------------------------------------------
    def _build_data_tab(self, tab):
        ctk.CTkLabel(tab, text="База данных: расположение",
                     font=ctk.CTkFont(size=13, weight="bold")).pack(
            padx=14, pady=(8, 2), anchor="w")
        self.db_path_label = ctk.CTkLabel(tab, text=default_db_path(),
                                          text_color="gray", font=("Consolas", 11),
                                          anchor="w", justify="left", wraplength=640)
        self.db_path_label.pack(padx=14, fill="x")
        ctk.CTkLabel(tab, text="В базе — история отчётов, дубли кварталов, пользователи.\n"
                               "Общая сетевая папка = общая история и пользователи для всех ПК.\n"
                               "Смена пути применяется сразу, перезапуск не нужен.",
                     text_color="gray", font=ctk.CTkFont(size=10),
                     justify="left", anchor="w").pack(padx=14, pady=(4, 6), fill="x")
        bar = ctk.CTkFrame(tab, fg_color="transparent")
        bar.pack(padx=14, fill="x")
        ctk.CTkButton(bar, text="Сменить папку…", width=130, height=30,
                      command=self._change_db_dir).pack(side="left", padx=3)
        ctk.CTkButton(bar, text="Сбросить (рядом с приложением)", width=230,
                      height=30, fg_color="gray",
                      command=self._reset_db).pack(side="left", padx=3)

    def _change_db_dir(self):
        d = filedialog.askdirectory(parent=self, title="Папка для базы данных")
        if not d:
            return
        new = os.path.join(d, 'turnalog_history.db')
        cur = default_db_path()
        if os.path.normcase(new) == os.path.normcase(cur):
            return
        if not os.path.exists(new) and os.path.exists(cur):
            if messagebox.askyesno("Копирование",
                                   "Скопировать текущую базу (историю и пользователей)\n"
                                   "в новую папку?", parent=self):
                try:
                    shutil.copy(cur, new)
                except Exception as e:
                    messagebox.showerror("База данных",
                                         f"Не удалось скопировать базу:\n{e}", parent=self)
                    return
        try:
            t = HistoryDB(new)
            t.ensure_default_admin()
            t.close()
        except Exception as e:
            messagebox.showerror("База данных",
                                 f"Не удалось открыть/создать базу:\n{new}\n{e}",
                                 parent=self)
            return
        set_db_path(new)
        self.cfg['db_path'] = new
        save_config(self.cfg)
        self.db_path_label.configure(text=new)

    def _reset_db(self):
        self.cfg.pop('db_path', None)
        save_config(self.cfg)
        set_db_path(None)
        self.db_path_label.configure(text=default_db_path())

    # ------------------------------------------------------------
    def ok(self):
        tax = {}
        for hid, _hname in self.HOTELS:
            try:
                tp = self._cur_params(hid)
            except ValueError:
                messagebox.showerror("Настройки",
                                     f"«{_hname}»: числовые поля должны содержать числа.",
                                     parent=self)
                return
            for need, tname in (('{P}', 'цены {P}'), ('{N}', 'ночей {N}')):
                if need not in tp.tpl_daily:
                    messagebox.showerror("Настройки",
                                         f"«{_hname}»: в шаблоне сут.ставки нет {tname}.",
                                         parent=self)
                    return
            if '{M}' not in tp.tpl_tax or '{N}' not in tp.tpl_tax or '{P}' not in tp.tpl_tax:
                messagebox.showerror(
                    "Настройки",
                    "«" + _hname + "»: в шаблоне налога нужны плейсхолдеры "
                    "{M}, {N} и {P}.", parent=self)
                return
            tax[hid] = tp.as_dict()

        self.save_agent_details()
        for a in self.agents_edit:
            if not a.get('name', '').strip():
                messagebox.showerror("Настройки", "Есть агент без названия.", parent=self)
                return
            for r in a.get('rates', []):
                if r.get('from') and parse_dmy(r['from']) is None:
                    messagebox.showerror(
                        "Настройки",
                        f"Агент «{a['name']}»: неверная дата «с» ({r['from']}).",
                        parent=self)
                    return
                if r.get('to') and parse_dmy(r['to']) is None:
                    messagebox.showerror(
                        "Настройки",
                        f"Агент «{a['name']}»: неверная дата «по» ({r['to']}).",
                        parent=self)
                    return
        self.result = {'tax': tax, 'agents': self.agents_edit}
        self.destroy()


# ================================================================
# ДИАЛОГ НЕРАСПОЗНАННЫХ ИСТОЧНИКОВ
# ================================================================
HOTEL_LABELS = {'tch': 'Чайковский', 'dna': 'Дона', 'usd': 'Усадьба'}


def autosize(win, width, min_h=120):
    """Высота окна по содержимому (не больше экрана), ширина фиксированная."""
    win.update_idletasks()
    h = max(min_h, win.winfo_reqheight() + 4)
    h = min(h, win.winfo_screenheight() - 60)
    win.geometry(f"{width}x{h}")


NEW_AGENT = "➕ Новый агент…"


class UnrecognizedSourcesDialog(ctk.CTkToplevel):
    """Для каждого нераспознанного источника: сопоставить с агентом,
    создать нового агента (сохранится глобально) или оставить 0%."""

    def __init__(self, master, items, agent_names):
        super().__init__(master)
        self.title("Нераспознанные источники")
        self.geometry("720x540")
        self.resizable(False, True)
        self.transient(master)
        self.result = None

        ctk.CTkLabel(self, text="Нераспознанные источники",
                     font=ctk.CTkFont(size=16, weight="bold")).pack(pady=(12, 2))
        ctk.CTkLabel(self, text="Укажите, к кому отнести источники. Выбор сохранится "
                                "и в следующий раз вопрос не появится.\n«— не указано —» "
                                "= комиссия 0% без подсветки ошибки.",
                     text_color="gray").pack(padx=20)

        sc = ctk.CTkScrollableFrame(self, fg_color=("gray88", "gray17"))
        sc.pack(fill="both", expand=True, padx=15, pady=8)

        opts = [NOT_SET] + agent_names + [NEW_AGENT]
        self.rows = {}
        for src, cnt in items:
            fr = ctk.CTkFrame(sc, fg_color="transparent")
            fr.pack(fill="x", pady=3, padx=6)
            ctk.CTkLabel(fr, text=f"{src}   (строк: {cnt})", width=330, anchor="w",
                         justify="left", font=ctk.CTkFont(size=12)).pack(side="left")
            var = tk.StringVar(value=NOT_SET)
            ctk.CTkOptionMenu(fr, variable=var, values=opts, width=230,
                              command=lambda v, s=src: self._toggle(s)).pack(side="left", padx=6)
            new_fr = ctk.CTkFrame(fr, fg_color="transparent")
            name_var = tk.StringVar()
            comm_var = tk.StringVar()
            ctk.CTkEntry(new_fr, textvariable=name_var, width=140, height=26,
                         placeholder_text="Имя агента",
                         font=ctk.CTkFont(size=11)).pack(side="left", padx=(6, 3))
            ctk.CTkEntry(new_fr, textvariable=comm_var, width=50, height=26,
                         placeholder_text="%",
                         font=ctk.CTkFont(size=11)).pack(side="left")
            self.rows[src] = dict(var=var, new_fr=new_fr,
                                  name_var=name_var, comm_var=comm_var)

        bar = ctk.CTkFrame(self, fg_color="transparent")
        bar.pack(pady=(0, 12))
        ctk.CTkButton(bar, text="OK", width=120, fg_color="green",
                      command=self.ok).pack(side="left", padx=6)
        ctk.CTkButton(bar, text="Отмена", width=120, fg_color="gray",
                      command=self.destroy).pack(side="left", padx=6)
        self.grab_set()
        self.lift()
        self.focus_force()
        self.attributes("-topmost", True)

    def _toggle(self, src):
        r = self.rows[src]
        if r['var'].get() == NEW_AGENT:
            r['new_fr'].pack(side="left")
        else:
            r['new_fr'].pack_forget()

    def ok(self):
        result = {}
        for src, r in self.rows.items():
            choice = r['var'].get()
            if choice == NOT_SET:
                result[src] = ('none',)
            elif choice == NEW_AGENT:
                name = r['name_var'].get().strip()
                if not name:
                    messagebox.showerror("Новый агент",
                                         f"Для «{src}» укажите имя нового агента.",
                                         parent=self)
                    return
                try:
                    comm = float(r['comm_var'].get().strip().replace(',', '.'))
                except ValueError:
                    messagebox.showerror("Новый агент",
                                         f"Для «{name}» укажите комиссию числом.",
                                         parent=self)
                    return
                result[src] = ('new', name, comm)
            else:
                result[src] = ('agent', choice)
        self.result = result
        self.destroy()


# ================================================================
# ДИАЛОГ ПЕРВОГО ЗАПУСКА: ГДЕ ХРАНИТЬ БАЗУ
# ================================================================
class DbLocationDialog(ctk.CTkToplevel):
    def __init__(self, master):
        super().__init__(master)
        self.title("База данных")
        self.transient(master)
        self.resizable(False, False)
        self.configure(fg_color=ui.COLORS['bg'])
        self.result = None

        ctk.CTkLabel(self, text="Где хранить базу данных?",
                     font=ctk.CTkFont(size=15, weight="bold")).pack(pady=(16, 2))
        ctk.CTkLabel(self, text="В базе — история отчётов, дубли кварталов,\n"
                                "пользователи и пароли.\n\n"
                                "• «Рядом с приложением» — база только на этом ПК.\n"
                                "• «Выбрать папку» — например, общая сетевая папка:\n"
                                "   тогда история и пользователи будут общими.",
                     justify="left", font=ctk.CTkFont(size=12)).pack(padx=20)
        b = ctk.CTkFrame(self, fg_color="transparent")
        b.pack(pady=14)
        ctk.CTkButton(b, text="Рядом с приложением", width=170, height=34,
                      fg_color="green",
                      command=lambda: self.finish(None)).pack(side="left", padx=5)
        ctk.CTkButton(b, text="Выбрать папку…", width=140, height=34,
                      command=self.choose).pack(side="left", padx=5)
        autosize(self, 470)
        self.lift()
        self.focus_force()
        self.attributes("-topmost", True)

    def choose(self):
        d = filedialog.askdirectory(parent=self, title="Папка для базы данных")
        if d:
            self.finish(os.path.join(d, 'turnalog_history.db'))

    def finish(self, path):
        self.result = path
        self.destroy()


# ================================================================
# ОКНО ВХОДА
# ================================================================
class LoginDialog(ctk.CTkToplevel):
    def __init__(self, master, db):
        super().__init__(master)
        self.title("Вход")
        self.resizable(False, False)
        self.transient(master)
        self.configure(fg_color=ui.COLORS['bg'])
        self.db = db
        self.user = None

        ctk.CTkLabel(self, text="ТУРНАЛОГ",
                     font=ctk.CTkFont(size=20, weight="bold")).pack(pady=(16, 0))
        ctk.CTkLabel(self, text="Вход в систему", text_color="gray",
                     font=ctk.CTkFont(size=11)).pack()
        self.login_var = tk.StringVar()
        self.pass_var = tk.StringVar()
        ctk.CTkEntry(self, textvariable=self.login_var, height=32,
                     placeholder_text="Логин").pack(padx=60, pady=(10, 4), fill="x")
        ctk.CTkEntry(self, textvariable=self.pass_var, height=32, show="*",
                     placeholder_text="Пароль").pack(padx=60, pady=2, fill="x")
        self.err = ctk.CTkLabel(self, text="", text_color="#C33",
                                font=ctk.CTkFont(size=11), height=16)
        self.err.pack()
        ctk.CTkButton(self, text="Войти", fg_color=ui.COLORS['green'],
                      hover_color=ui.COLORS['green_hover'], height=34,
                      font=ui.F_UI,
                      command=self.ok).pack(padx=60, pady=4, fill="x")
        ctk.CTkLabel(self, text=f"БД: {default_db_path()}", text_color="gray",
                     font=ctk.CTkFont(size=8), wraplength=300).pack(pady=(2, 0))
        if db.check_user('admin', 'admin'):
            ctk.CTkLabel(self, text="Первый запуск: логин admin, пароль admin — "
                                     "смените пароль в «Пользователи»",
                         text_color="gray", font=ctk.CTkFont(size=9)).pack(pady=(4, 0))
        self.bind('<Return>', lambda e: self.ok())
        autosize(self, 340)
        self.lift()
        self.focus_force()
        self.attributes("-topmost", True)

    def ok(self):
        u = self.db.check_user(self.login_var.get().strip(), self.pass_var.get())
        if u is None:
            self.err.configure(text="Неверный логин или пароль")
            return
        self.user = u
        self.destroy()


# ================================================================
# ПРИЛОЖЕНИЕ
# ================================================================
class FileZone(ctk.CTkFrame):
    """Файл как объект: пустая зона → выбранный файл → ошибка чтения."""

    def __init__(self, parent, title, hint, on_pick, on_clear=None, on_cols=None):
        super().__init__(parent, corner_radius=10, border_width=1,
                         border_color=ui.COLORS['zone_border'],
                         fg_color=ui.COLORS['zone'])
        self.grid_columnconfigure(0, weight=1)
        self.on_clear = on_clear
        self._title = title
        self._hint = hint

        top = ctk.CTkFrame(self, fg_color="transparent")
        top.grid(row=0, column=0, sticky="ew", padx=12, pady=(10, 0))
        top.grid_columnconfigure(0, weight=1)
        self.tlabel = ctk.CTkLabel(top, text=title, font=("Segoe UI", 12, "bold"),
                                   text_color=ui.COLORS['text'], anchor="w")
        self.tlabel.grid(row=0, column=0, sticky="w")
        self.clear_btn = ctk.CTkButton(top, text="✕", width=24, height=24,
                                       fg_color="transparent",
                                       text_color=ui.COLORS['faint'],
                                       hover_color=ui.COLORS['red_soft'],
                                       command=self._do_clear)
        self.info = ctk.CTkLabel(self, text=hint, text_color=ui.COLORS['faint'],
                                 font=ui.F_SMALL, anchor="w")
        self.info.grid(row=1, column=0, sticky="ew", padx=12, pady=(2, 0))
        btns = ctk.CTkFrame(self, fg_color="transparent")
        btns.grid(row=2, column=0, sticky="ew", padx=12, pady=(6, 10))
        self.pick_btn = ctk.CTkButton(btns, text="Обзор…", width=96, height=30,
                                      fg_color=ui.COLORS['accent'],
                                      hover_color=ui.COLORS['accent_hover'],
                                      font=ui.F_SMALL, command=on_pick)
        self.pick_btn.pack(side="left")
        if on_cols:
            ctk.CTkButton(btns, text="Колонки", width=92, height=30,
                          fg_color="transparent", border_width=1,
                          border_color=ui.COLORS['line'],
                          text_color=ui.COLORS['accent'],
                          hover_color=ui.COLORS['accent_soft'],
                          font=ui.F_SMALL, command=on_cols).pack(side="left", padx=(6, 0))
        self.set_file(None)

    def _do_clear(self):
        if self.on_clear:
            self.on_clear()

    def set_file(self, path, info=""):
        if path:
            self.configure(fg_color=ui.COLORS['green_soft'],
                           border_color=ui.COLORS['green_border'])
            self.tlabel.configure(text=os.path.basename(path))
            self.info.configure(text=info or self._hint,
                                text_color=ui.COLORS['muted'])
            self.pick_btn.configure(text="Заменить")
            self.clear_btn.grid(row=0, column=1, sticky="e")
        else:
            self.configure(fg_color=ui.COLORS['zone'],
                           border_color=ui.COLORS['zone_border'])
            self.tlabel.configure(text=self._title)
            self.info.configure(text=self._hint, text_color=ui.COLORS['faint'])
            self.pick_btn.configure(text="Обзор…")
            self.clear_btn.grid_forget()

    def set_info(self, info, color=None):
        self.info.configure(text=info, text_color=color or ui.COLORS['muted'])


class App(ctk.CTk):
    SLOT_DEFS = {
        'tch': (('data', '«Данные» — брони',
                 'Выгрузка броней: даты, сумма, источник, гости',
                 'Колонки «Данные»'),
                ('rooms', '«Номера» — расселение',
                 'ФИО гостя + № комнаты', 'Колонки «Номера»')),
        'dna': (('data', '«Данные» — брони',
                 'Выгрузка броней с колонкой «Источник»', 'Колонки «Данные»'),),
        'usd': (('rep', 'Отчёт «Доход по номеру»',
                 'Выгрузка UCS Shelter, структура фиксированная', None),),
    }

    def __init__(self, user):
        super().__init__()
        self.user = user
        self.is_admin = user['role'] == 'admin'
        self.allowed_hotels = ({'tch', 'dna', 'usd'} if self.is_admin
                               else {user['hotel']})
        self.title("Турналог — обработка данных гостиниц")
        w, h = 1180, 780
        x = max(0, (self.winfo_screenwidth() - w) // 2)
        y = max(0, (self.winfo_screenheight() - h) // 2)
        self.geometry(f"{w}x{h}+{x}+{y}")
        self.minsize(1100, 700)
        self.configure(fg_color=ui.COLORS['bg'])
        ui.apply_treeview_style()
        self.cfg = load_config()

        self.slot_state = {}
        for hid, slots in (('tch', ('data', 'rooms')), ('dna', ('data',)),
                           ('usd', ('rep',))):
            self.slot_state[hid] = {s: {'path_var': tk.StringVar()} for s in slots}
        self.q_vars = {}
        self.dup_vars = {}
        self.zones = {}
        self.cta_btns = {}
        self.stat_lbs = {}
        self.nav_btns = {}

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)
        self._build_sidebar()

        self.right = ctk.CTkFrame(self, fg_color="transparent")
        self.right.grid(row=0, column=1, sticky="nsew", padx=12, pady=(10, 4))
        self.right.grid_columnconfigure(0, weight=1)
        self.right.grid_rowconfigure(0, weight=1)

        self.pages = {}
        for pid in ('tch', 'dna', 'usd'):
            if pid in self.allowed_hotels:
                self._build_hotel_page(pid)
        self._build_history_page()

        sb = ctk.CTkFrame(self, fg_color=ui.COLORS['card'], corner_radius=0,
                          border_width=1, border_color=ui.COLORS['line'])
        sb.grid(row=1, column=0, columnspan=2, sticky="ew")
        ctk.CTkLabel(sb, text="●", text_color=ui.COLORS['green'],
                     font=ui.F_SMALL).pack(side="left", padx=(12, 4))
        self.status_var = tk.StringVar(value="Готово к работе")
        ctk.CTkLabel(sb, textvariable=self.status_var, text_color=ui.COLORS['muted'],
                     font=ui.F_SMALL, anchor="w").pack(side="left", fill="x", expand=True)
        self.db_label = ctk.CTkLabel(sb, text=default_db_path(),
                                     text_color=ui.COLORS['faint'], font=ui.F_TINY)
        self.db_label.pack(side="right", padx=12)

        self.usd_var = self.slot_state['usd']['rep']['path_var']
        first = next(p for p in ('tch', 'dna', 'usd', 'hist') if p in self.pages)
        self.show_page(first)

    # ------------------------------------------------------------
    # сайдбар
    # ------------------------------------------------------------
    def _build_sidebar(self):
        sb = ctk.CTkFrame(self, width=232, corner_radius=0,
                          fg_color=ui.COLORS['sidebar'])
        sb.grid(row=0, column=0, sticky="nsew")
        sb.grid_propagate(False)
        ctk.CTkLabel(sb, text="ТУРНАЛОГ", font=("Segoe UI", 18, "bold"),
                     text_color="white").pack(padx=16, pady=(18, 0), anchor="w")
        ctk.CTkLabel(sb, text="налог по выгрузкам отелей", text_color="#8FA3C8",
                     font=ui.F_TINY).pack(padx=16, pady=(0, 14), anchor="w")

        ctk.CTkLabel(sb, text="ОТЕЛИ", text_color="#7286AC",
                     font=("Segoe UI", 10, "bold")).pack(padx=16, pady=(4, 2), anchor="w")
        for pid in ('tch', 'dna', 'usd'):
            if pid not in self.allowed_hotels:
                continue
            rowf = ctk.CTkFrame(sb, fg_color="transparent")
            rowf.pack(fill="x", padx=8, pady=1)
            ctk.CTkLabel(rowf, text="●", text_color=ui.HOTEL_DOT[pid],
                         font=ui.F_SMALL, width=16).pack(side="left", padx=(8, 0))
            b = ctk.CTkButton(rowf, text=ui.HOTEL_NAME[pid], anchor="w",
                              fg_color="transparent", text_color="#DCE6F5",
                              hover_color=ui.COLORS['sidebar_hover'], height=36,
                              font=ui.F_UI, command=lambda p=pid: self.show_page(p))
            b.pack(side="left", fill="x", expand=True, padx=(0, 6))
            self.nav_btns[pid] = b

        ctk.CTkFrame(sb, height=1, fg_color=ui.COLORS['sidebar_act']).pack(
            fill="x", padx=16, pady=10)
        b = ctk.CTkButton(sb, text="История", anchor="w", fg_color="transparent",
                          text_color="#DCE6F5", hover_color=ui.COLORS['sidebar_hover'],
                          height=36, font=ui.F_UI,
                          command=lambda: self.show_page('hist'))
        b.pack(fill="x", padx=8, pady=1)
        self.nav_btns['hist'] = b
        if self.is_admin:
            b = ctk.CTkButton(sb, text="⚙  Настройки", anchor="w",
                              fg_color="transparent", text_color="#DCE6F5",
                              hover_color=ui.COLORS['sidebar_hover'], height=36,
                              font=ui.F_UI, command=self.open_settings)
            b.pack(fill="x", padx=8, pady=1)

        low = ctk.CTkFrame(sb, fg_color="transparent")
        low.pack(side="bottom", fill="x", padx=14, pady=12)
        ctk.CTkLabel(low, text=self.user['login'][0].upper(), width=34, height=34,
                     corner_radius=17, fg_color=ui.COLORS['accent'],
                     text_color="white", font=("Segoe UI", 14, "bold")).pack(side="left")
        who = ctk.CTkFrame(low, fg_color="transparent")
        who.pack(side="left", padx=(8, 0))
        ctk.CTkLabel(who, text=self.user['login'], text_color="white",
                     font=ui.F_SMALL, anchor="w").pack(anchor="w")
        ctk.CTkLabel(who, text=("администратор" if self.is_admin
                                else ui.HOTEL_NAME.get(self.user['hotel'], ''))
                     + f" · v{APP_VERSION}",
                     text_color="#8FA3C8", font=ui.F_TINY, anchor="w").pack(anchor="w")

    def show_page(self, pid):
        for p, fr in self.pages.items():
            fr.grid_forget()
        self.pages[pid].grid(row=0, column=0, sticky="nsew")
        for p, b in self.nav_btns.items():
            if p == pid:
                b.configure(fg_color=ui.COLORS['sidebar_act'], text_color="white")
            else:
                b.configure(fg_color="transparent", text_color="#DCE6F5")
        if pid == 'hist':
            self.refresh_history()

    def set_status(self, text):
        self.status_var.set(text)
        self.update_idletasks()

    # ------------------------------------------------------------
    # страница отеля
    # ------------------------------------------------------------
    def _build_hotel_page(self, pid):
        pg = ctk.CTkScrollableFrame(self.right, fg_color="transparent",
                                    scrollbar_button_color=ui.COLORS['line'])
        self.pages[pid] = pg
        now = datetime.now()

        hdr = ctk.CTkFrame(pg, fg_color="transparent")
        hdr.pack(fill="x", padx=6, pady=(4, 8))
        ctk.CTkLabel(hdr, text=f"Отель «{ui.HOTEL_NAME[pid]}»", font=ui.F_PAGE,
                     text_color=ui.COLORS['text']).pack(side="left")
        bg, fg = ui.HOTEL_PILL[pid]
        sub = {'tch': 'корпуса 1 и 2 · листы 1&2', 'dna': 'без корпусов · агенты',
               'usd': 'без корпусов · UCS Shelter'}[pid]
        ctk.CTkLabel(hdr, text=sub, font=ui.F_TINY, fg_color=bg, text_color=fg,
                     corner_radius=6, height=22).pack(side="left", padx=(10, 0), pady=(8, 0))
        tp = TaxParams(self.cfg.get('tax', {}).get(pid))
        ctk.CTkLabel(hdr, text=f"ставка {tp.rate:g}% · минимум {tp.min:g} ₽ · порог {tp.threshold:g}",
                     text_color=ui.COLORS['muted'], font=ui.F_SMALL).pack(side="right", pady=(8, 0))

        steps = ctk.CTkFrame(pg, fg_color="transparent")
        steps.pack(fill="x", padx=4, pady=4)
        steps.grid_columnconfigure(0, weight=3, minsize=430)
        steps.grid_columnconfigure(1, weight=2, minsize=320)

        # --- шаг 1: файлы ---
        c1 = ui.card(steps)
        c1.grid(row=0, column=0, sticky="nsew", padx=(0, 6), pady=4)
        ctk.CTkLabel(c1, text="1 · Исходные файлы", font=ui.F_CARD,
                     text_color=ui.COLORS['text']).pack(padx=16, pady=(14, 8), anchor="w")
        for slot, title, hint, cols_label in self.SLOT_DEFS[pid]:
            zone = FileZone(c1, title, hint,
                            on_pick=lambda h=pid, s=slot: self.browse(h, s),
                            on_clear=lambda h=pid, s=slot: self.clear_file(h, s),
                            on_cols=(lambda h=pid, s=slot: self.open_mapping(h, s))
                            if cols_label else None)
            if cols_label:
                zone.pick_btn.configure(text="Обзор…")
            zone.pack(fill="x", padx=14, pady=5)
            self.zones[(pid, slot)] = zone
        ctk.CTkLabel(c1, text="", font=ui.F_TINY).pack()  # нижний отступ

        # --- шаг 2: обработка ---
        c2 = ui.card(steps)
        c2.grid(row=0, column=1, sticky="nsew", padx=(6, 0), pady=4)
        ctk.CTkLabel(c2, text="2 · Обработка", font=ui.F_CARD,
                     text_color=ui.COLORS['text']).pack(padx=16, pady=(14, 8), anchor="w")

        qrow = ctk.CTkFrame(c2, fg_color="transparent")
        qrow.pack(fill="x", padx=16, pady=2)
        ctk.CTkLabel(qrow, text="Квартал отчёта", text_color=ui.COLORS['muted'],
                     font=ui.F_SMALL).pack(anchor="w")
        qv = tk.StringVar(value=f"Q{(now.month - 1) // 3 + 1}")
        yv = tk.StringVar(value=str(now.year))
        self.q_vars[pid] = (yv, qv)
        q2 = ctk.CTkFrame(c2, fg_color="transparent")
        q2.pack(fill="x", padx=16, pady=(0, 6))
        ctk.CTkSegmentedButton(q2, values=["Q1", "Q2", "Q3", "Q4"], variable=qv,
                               width=190, height=32,
                               fg_color=ui.COLORS['line2'],
                               selected_color=ui.COLORS['accent'],
                               selected_hover_color=ui.COLORS['accent_hover'],
                               unselected_color=ui.COLORS['card'],
                               unselected_hover_color=ui.COLORS['line2'],
                               text_color=ui.COLORS['text'],
                               font=ui.F_SMALL).pack(side="left")
        ctk.CTkEntry(q2, textvariable=yv, width=70, height=32,
                     font=ui.F_UI, justify="center").pack(side="left", padx=(8, 0))

        self.dup_vars[pid] = tk.BooleanVar(value=True)
        ctk.CTkCheckBox(c2, text="Проверять дубли по истории",
                        variable=self.dup_vars[pid], font=ui.F_UI,
                        fg_color=ui.COLORS['accent'],
                        hover_color=ui.COLORS['line2']).pack(padx=16, pady=4, anchor="w")

        hintf = ctk.CTkFrame(c2, fg_color=ui.COLORS['purple_soft'], corner_radius=10)
        hintf.pack(fill="x", padx=16, pady=(6, 10))
        ctk.CTkLabel(hintf, text="Бронь на стыке кварталов попадёт в два отчёта — "
                                 "во втором она будет помечена фиолетовым.",
                     text_color=ui.COLORS['purple_text'], font=ui.F_SMALL,
                     justify="left", anchor="w").pack(padx=12, pady=8, fill="x")

        run = {'tch': self.run_tchaikovsky, 'dna': self.run_dona,
               'usd': self.run_usadba}[pid]
        btn = ctk.CTkButton(c2, text="ОБРАБОТАТЬ И СОХРАНИТЬ", height=44,
                            fg_color=ui.COLORS['green'],
                            hover_color=ui.COLORS['green_hover'],
                            font=("Segoe UI", 14, "bold"), command=run)
        btn.pack(fill="x", padx=16, pady=(0, 16))
        self.cta_btns[pid] = btn

        # --- мини-карточки последней обработки ---
        st = ctk.CTkFrame(pg, fg_color="transparent")
        st.pack(fill="x", padx=4, pady=6)
        self.stat_lbs[pid] = {}
        for i, (key, label) in enumerate((('date', 'Последняя обработка'),
                                          ('rows', 'Строк'), ('dups', 'Дублей'))):
            c = ui.card(st)
            c.pack(side="left", fill="x", expand=True,
                   padx=(0 if i == 0 else 6, 0 if i == 2 else 6))
            ctk.CTkLabel(c, text=label, text_color=ui.COLORS['muted'],
                         font=ui.F_TINY).pack(padx=14, pady=(10, 0), anchor="w")
            val = ctk.CTkLabel(c, text="—", text_color=ui.COLORS['text'],
                               font=("Segoe UI", 15, "bold"))
            val.pack(padx=14, pady=(0, 10), anchor="w")
            self.stat_lbs[pid][key] = val
        self.refresh_stats(pid)
        self._update_cta(pid)

    def refresh_stats(self, pid):
        try:
            with HistoryDB() as db:
                reps = db.list_reports(hotel=ui.HOTEL_NAME[pid])
        except Exception:
            reps = []
        d = self.stat_lbs.get(pid)
        if not d:
            return
        if reps:
            r = reps[0]
            d['date'].configure(text=r['processed_at'][:16].replace('T', ' '))
            d['rows'].configure(text=str(r['rows']))
            d['dups'].configure(text=str(r['dups']),
                                text_color=ui.COLORS['amber'] if r['dups'] else ui.COLORS['green'])
        else:
            d['date'].configure(text="ещё не было")
            d['rows'].configure(text="—")
            d['dups'].configure(text="—")

    # ------------------------------------------------------------
    # файлы
    # ------------------------------------------------------------
    def browse(self, hotel, slot, on_change=None):
        path = filedialog.askopenfilename(
            parent=self, title="Выберите файл",
            filetypes=[("Excel и CSV файлы", "*.xlsx *.csv")])
        if not path:
            return
        self.slot_state[hotel][slot]['path_var'].set(path)
        self.on_file_selected(hotel, slot)

    def clear_file(self, hotel, slot):
        st = self.slot_state[hotel][slot]
        st['path_var'].set('')
        st['mapping'], st['extra'], st['headers'] = None, [], None
        self.zones[(hotel, slot)].set_file(None)
        self._update_cta(hotel)

    def _update_cta(self, hotel):
        need = ('data', 'rooms') if hotel == 'tch' else ('data',) if hotel == 'dna' else ('rep',)
        ok = all(self.slot_state[hotel][s]['path_var'].get() for s in need)
        self.cta_btns[hotel].configure(
            state="normal" if ok else "disabled",
            fg_color=ui.COLORS['green'] if ok else ui.COLORS['green_dis'])

    def on_file_selected(self, hotel, slot):
        st = self.slot_state[hotel][slot]
        path = st['path_var'].get()
        zone = self.zones[(hotel, slot)]
        try:
            headers = read_headers(path)
        except Exception as e:
            st['headers'] = None
            zone.set_file(path, "⚠ файл не читается — проверьте формат")
            log.warning("read_headers %s: %s", path, e)
            self._update_cta(hotel)
            return
        st['headers'] = headers
        schema = self._schema(hotel, slot)
        saved = self.cfg.get(f"{hotel}_{slot}", {})
        if mapping_valid(saved.get('mapping'), headers, schema):
            st['mapping'], st['extra'] = saved['mapping'], saved.get('extra', [])
            info = f"{len(headers)} колонок · сопоставлено ✓"
            if hotel == 'dna':
                info += self._dona_unknown_note(path, st['mapping'])
            zone.set_file(path, info)
        else:
            st['mapping'], st['extra'] = None, []
            zone.set_file(path, f"{len(headers)} колонок — нажмите «Колонки»")
        self._update_cta(hotel)

    def _schema(self, hotel, slot):
        return {('tch', 'data'): TCH_DATA_SCHEMA, ('tch', 'rooms'): TCH_ROOMS_SCHEMA,
                ('dna', 'data'): DNA_DATA_SCHEMA}[hotel, slot]

    def ensure_mapping(self, hotel, slot):
        st = self.slot_state[hotel][slot]
        path = st['path_var'].get()
        if not path:
            return False
        if st.get('headers') is None:
            self.on_file_selected(hotel, slot)
        headers = st.get('headers')
        if not headers:
            messagebox.showerror("Ошибка", f"Файл не читается:\n{path}")
            return False
        schema = self._schema(hotel, slot)
        if mapping_valid(st.get('mapping'), headers, schema):
            return True
        dlg = MappingDialog(self, os.path.basename(path), headers, schema,
                            mapping=st.get('mapping'), extra=st.get('extra'))
        self.wait_window(dlg)
        if not dlg.result:
            return False
        st['mapping'], st['extra'] = dlg.result
        self.cfg[f"{hotel}_{slot}"] = {"mapping": dlg.result[0], "extra": dlg.result[1]}
        save_config(self.cfg)
        self.zones[(hotel, slot)].set_info(f"{len(headers)} колонок · сопоставлено ✓",
                                           ui.COLORS['green'])
        return True

    def open_mapping(self, hotel, slot):
        st = self.slot_state[hotel][slot]
        if not st['path_var'].get():
            messagebox.showinfo("Колонки", "Сначала выберите файл.")
            return
        self.ensure_mapping(hotel, slot)

    def _dona_unknown_note(self, path, mapping):
        try:
            scol = (mapping or {}).get('source')
            if not scol:
                return ""
            dfc = read_table(path)
            if scol not in dfc.columns:
                return ""
            cnt = dfc[scol].fillna('').astype(str).str.strip().value_counts()
            agents = migrate_agents(self.cfg)
            unk = [s for s, c in cnt.items()
                   if s and match_agent(s, agents, hotel='dna')[1] is None]
            return f" · нераспознанных: {len(unk)}" if unk else ""
        except Exception:
            log.exception("_dona_unknown_note")
            return ""

    def open_settings(self):
        dlg = SettingsDialog(self, self.cfg)
        self.wait_window(dlg)
        if dlg.result:
            self.cfg['tax'] = dlg.result['tax']
            self.cfg['agents'] = dlg.result['agents']
            save_config(self.cfg)
            ui.toast(self, "Настройки сохранены")
            self.set_status("Настройки сохранены")
            self.db_label.configure(text=default_db_path())
            for pid in self.pages:
                if pid in self.stat_lbs:
                    self.refresh_stats(pid)

    # ------------------------------------------------------------
    # история
    # ------------------------------------------------------------
    def _build_history_page(self):
        pg = ctk.CTkFrame(self.right, fg_color="transparent")
        self.pages['hist'] = pg

        hdr = ctk.CTkFrame(pg, fg_color="transparent")
        hdr.pack(fill="x", padx=6, pady=(4, 8))
        ctk.CTkLabel(hdr, text="История отчётов", font=ui.F_PAGE,
                     text_color=ui.COLORS['text']).pack(side="left")
        self.hist_search = tk.StringVar()
        self.hist_search.trace_add("write", lambda *a: self._apply_hist_filter())
        ctk.CTkEntry(hdr, textvariable=self.hist_search, width=220, height=32,
                     placeholder_text="Поиск: отель, квартал, файл…",
                     font=ui.F_SMALL).pack(side="right", pady=(8, 0))

        bar = ctk.CTkFrame(pg, fg_color="transparent")
        bar.pack(fill="x", padx=6, pady=(0, 6))
        for label, cmd, color in (("Обновить", self.refresh_history, None),
                                  ("Открыть оригинал",
                                   lambda: self.open_history_file('orig'), None),
                                  ("Открыть результат",
                                   lambda: self.open_history_file('out'), None),
                                  ("Удалить", self.delete_history_report,
                                   ui.COLORS['red'])):
            kw = dict(height=32, font=ui.F_SMALL)
            if color:
                kw.update(fg_color=color, hover_color="#C22F35")
            else:
                kw.update(fg_color=ui.COLORS['card'], border_width=1,
                          border_color=ui.COLORS['line'],
                          text_color=ui.COLORS['text'],
                          hover_color=ui.COLORS['line2'])
            ctk.CTkButton(bar, text=label, command=cmd, **kw).pack(side="left", padx=(0, 6))

        wrap = ui.card(pg)
        wrap.pack(fill="both", expand=True, padx=6, pady=2)
        cols = ("id", "hotel", "quarter", "processed_at", "rows", "dups",
                "original_name", "output_name")
        self.history_tree = ttk.Treeview(wrap, columns=cols, show="headings", height=14)
        heads = {"id": "№", "hotel": "Отель", "quarter": "Квартал",
                 "processed_at": "Обработан", "rows": "Строк", "dups": "Дубли",
                 "original_name": "Оригинал", "output_name": "Результат"}
        widths = {"id": 35, "hotel": 100, "quarter": 75, "processed_at": 130,
                  "rows": 55, "dups": 55, "original_name": 170, "output_name": 170}
        for c in cols:
            self.history_tree.heading(c, text=heads[c])
            self.history_tree.column(c, width=widths[c], anchor="w")
        for pid, color in ui.HOTEL_DOT.items():
            self.history_tree.tag_configure(f"h_{pid}", foreground=color)
        self.history_tree.tag_configure("dup", foreground=ui.COLORS['amber'])
        sb = ttk.Scrollbar(wrap, orient="vertical", command=self.history_tree.yview)
        self.history_tree.configure(yscrollcommand=sb.set)
        self.history_tree.pack(side="left", fill="both", expand=True, padx=(8, 0), pady=8)
        sb.pack(side="right", fill="y", pady=8)
        self.history_tree.bind("<<TreeviewSelect>>", lambda e: self.show_history_dups())
        self._hist_all = []

        ctk.CTkLabel(pg, text="Брони с переходом кварталов в выбранном отчёте",
                     text_color=ui.COLORS['muted'], font=ui.F_SMALL,
                     anchor="w").pack(padx=10, pady=(10, 4), fill="x")
        self.dup_text = tk.Text(pg, height=5, wrap="word", state="disabled",
                                bg=ui.COLORS['card'], relief="flat",
                                highlightthickness=1,
                                highlightbackground=ui.COLORS['line'],
                                fg=ui.COLORS['text'], font=("Segoe UI", 10))
        self.dup_text.tag_configure("muted", foreground=ui.COLORS['faint'])
        self.dup_text.pack(padx=6, pady=(0, 8), fill="x")
        self.refresh_history()

    def refresh_history(self):
        hotel = None if self.is_admin else ui.HOTEL_NAME[self.user['hotel']]
        try:
            with HistoryDB() as db:
                self._hist_all = [tuple(r) for r in db.list_reports(hotel)]
        except Exception:
            log.exception("refresh_history")
            self._hist_all = []
        self._apply_hist_filter()

    def _apply_hist_filter(self):
        if not hasattr(self, 'history_tree'):
            return
        for i in self.history_tree.get_children():
            self.history_tree.delete(i)
        q = (self.hist_search.get() or '').strip().lower()
        for r in self._hist_all:
            if q and q not in ' '.join(str(x) for x in r).lower():
                continue
            hid = {v: k for k, v in ui.HOTEL_NAME.items()}.get(r[1])
            tags = (f"h_{hid}",) if hid else ()
            if r[5]:
                tags += ("dup",)
            self.history_tree.insert("", "end", iid=str(r[0]), values=r, tags=tags)

    def _selected_report_id(self):
        sel = self.history_tree.selection()
        return int(sel[0]) if sel else None

    def show_history_dups(self):
        self.dup_text.configure(state="normal")
        self.dup_text.delete("1.0", "end")
        rid = self._selected_report_id()
        rows = []
        if rid is not None:
            try:
                with HistoryDB() as db:
                    rows = db.dup_bookings(rid)
            except Exception:
                log.exception("show_history_dups")
        if rows:
            for b in rows:
                self.dup_text.insert("end",
                    f"{b['guest'] or '(без ФИО)'}: {b['checkin'] or '—'} … "
                    f"{b['checkout'] or '—'} — {b['dup_note']}\n")
        else:
            self.dup_text.insert("end", "Выберите отчёт в таблице — здесь появятся "
                                        "строки, пересекающие границу квартала.", "muted")
        self.dup_text.configure(state="disabled")

    def open_history_file(self, kind):
        rid = self._selected_report_id()
        if rid is None:
            messagebox.showinfo("История", "Сначала выберите отчёт в списке.")
            return
        with HistoryDB() as db:
            name, blob = db.get_file(rid, kind)
        if not name or blob is None:
            messagebox.showwarning("История", "Файл не найден в БД.")
            return
        tmp = tempfile.mkdtemp(prefix="turnalog_")
        target = os.path.join(tmp, name)
        with open(target, 'wb') as f:
            f.write(blob)
        if name.lower().endswith('.zip'):
            with zipfile.ZipFile(target) as z:
                z.extractall(tmp)
            open_in_system(tmp)
        else:
            open_in_system(target)

    def delete_history_report(self):
        rid = self._selected_report_id()
        if rid is None:
            messagebox.showinfo("История", "Сначала выберите отчёт в списке.")
            return
        if messagebox.askyesno("Удалить отчёт",
                               "Удалить отчёт и его брони из истории?\n"
                               "Проверка дублей больше не будет его учитывать."):
            with HistoryDB() as db:
                db.delete_report(rid)
            self.refresh_history()
            self.show_history_dups()

    # ================================================================
    # ОБЩИЙ ЗАПУСК В ФОНЕ
    # ================================================================
    def _run_async(self, btn, work, on_done):
        btn.configure(state="disabled")
        self.set_status("Обработка…")

        def runner():
            try:
                result = work()
                err = None
            except Exception as e:
                log.exception("background work")
                result, err = None, e
            self.after(0, lambda: self._finish_async(btn, result, err, on_done))

        threading.Thread(target=runner, daemon=True).start()

    def _finish_async(self, btn, result, err, on_done):
        btn.configure(state="normal")
        if err:
            self.set_status("Ошибка обработки")
            messagebox.showerror("Ошибка", f"В процессе обработки произошла ошибка:\n{err}")
            return
        on_done(result)

    def _start_history(self, hotel, quarter):
        db = HistoryDB()
        old = db.find_same(hotel, quarter)
        if old:
            replace = messagebox.askyesno(
                "Отчёт уже есть в истории",
                f"За квартал {quarter} уже обработан отчёт от {old['processed_at'][:10]}\n"
                f"(файл: {old['original_name']}).\n\n"
                f"Заменить его? «Да» — старый удалится, дубли не помечаются.\n"
                f"«Нет» — старый останется, совпавшие брони будут помечены.")
            if replace:
                db.delete_report(old['id'])
        return db, (lambda rows: db.check_duplicates(hotel, rows))

    def _finish_history(self, db, hotel, quarter, input_paths, output_path,
                        bookings):
        orig_name, orig_blob = pack_originals(input_paths)
        with open(output_path, 'rb') as f:
            out_blob = f.read()
        dups = sum(1 for b in bookings if b.get('dup_note'))
        rid = db.add_report(hotel, quarter, orig_name, orig_blob,
                            os.path.basename(output_path), out_blob,
                            len(bookings), dups, "")
        db.add_bookings(rid, bookings)
        db.close()
        self.refresh_history()
        return dups

    def _show_summary(self, summary, output_path):
        total = sum(summary.values())
        if total > 0:
            msg = f"Файл сохранён, замечаний: {total}.\n\n"
            for err_name, count in summary.items():
                if count > 0:
                    msg += f"• {err_name}: {count}\n"
            msg += "\nПроблемные ячейки подсвечены цветом. Итоги добавлены в конец таблицы."
            messagebox.showwarning("Готово (есть замечания)", msg)
        else:
            messagebox.showinfo("Успех", f"Файл успешно сохранён без замечаний!\n{output_path}")

    # ================================================================
    # ЗАПУСК: ЧАЙКОВСКИЙ
    # ================================================================
    def run_tchaikovsky(self):
        s = self.slot_state['tch']
        data_path = s['data']['path_var'].get()
        rooms_path = s['rooms']['path_var'].get()
        if not data_path or not rooms_path:
            messagebox.showerror("Ошибка", "Выберите оба файла:\n«Данные» и «Номера»!")
            return
        if not self.ensure_mapping('tch', 'data'):
            return
        if not self.ensure_mapping('tch', 'rooms'):
            return
        md, mr = s['data']['mapping'], s['rooms']['mapping']
        ed = s['data'].get('extra', [])

        try:
            df_pre = read_table(data_path)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Файл не читается:\n{e}")
            return
        src_col = md.get('source')
        sources = list(df_pre[src_col].fillna('').astype(str))             if src_col and src_col in df_pre.columns else []
        res = self._resolve_unrecognized('tch', migrate_agents(self.cfg), sources)
        if res[0] is None:
            return
        agents, source_map = res

        output_path = filedialog.asksaveasfilename(
            parent=self, title="Сохранить итоговый файл как...",
            defaultextension=".xlsx", filetypes=[("Excel файлы", "*.xlsx")])
        if not output_path:
            return
        quarter = f"{self.q_vars['tch'][0].get()}-{self.q_vars['tch'][1].get()}"
        tax = self.cfg.get('tax', {}).get('tch')

        def work():
            db, checker = self._start_history("Чайковский", quarter)
            if not self.dup_vars['tch'].get():
                checker = None
            summary, bookings = self.process_tchaikovsky(
                data_path, rooms_path, output_path, md, mr,
                extra_data=ed, dup_checker=checker, tax=tax,
                agents=agents, source_map=source_map)
            dups = self._finish_history(db, "Чайковский", quarter,
                                        [data_path, rooms_path], output_path, bookings)
            return summary, output_path, len(bookings), dups

        def done(res):
            summary, out, n, dups = res
            self.set_status(f"Чайковский: готово, строк {n}, дублей {dups} → {os.path.basename(out)}")
            self.refresh_stats('tch')
            self._show_summary(summary, out)

        self._run_async(self.cta_btns['tch'], work, done)

    # ================================================================
    # ЗАПУСК: ДОНА
    # ================================================================
    def run_dona(self):
        s = self.slot_state['dna']
        data_path = s['data']['path_var'].get()
        if not data_path:
            messagebox.showerror("Ошибка", "Выберите файл «Данные»!")
            return
        if not self.ensure_mapping('dna', 'data'):
            return
        md, ed = s['data']['mapping'], s['data'].get('extra', [])

        try:
            df_pre = read_table(data_path)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Файл не читается:\n{e}")
            return
        src_col = md.get('source')
        sources = list(df_pre[src_col].fillna('').astype(str))             if src_col and src_col in df_pre.columns else []
        res = self._resolve_unrecognized('dna', migrate_agents(self.cfg), sources)
        if res[0] is None:
            return
        agents, source_map = res

        output_path = filedialog.asksaveasfilename(
            parent=self, title="Сохранить итоговый файл как...",
            defaultextension=".xlsx", filetypes=[("Excel файлы", "*.xlsx")])
        if not output_path:
            return
        quarter = f"{self.q_vars['dna'][0].get()}-{self.q_vars['dna'][1].get()}"
        tax = self.cfg.get('tax', {}).get('dna')

        def work():
            db, checker = self._start_history("Дона", quarter)
            if not self.dup_vars['dna'].get():
                checker = None
            summary, bookings = self.process_dona(
                data_path, output_path, md, extra_data=ed, dup_checker=checker,
                tax=tax, agents=agents, source_map=source_map)
            dups = self._finish_history(db, "Дона", quarter, [data_path],
                                        output_path, bookings)
            return summary, output_path, len(bookings), dups

        def done(res):
            summary, out, n, dups = res
            self.set_status(f"Дона: готово, строк {n}, дублей {dups} → {os.path.basename(out)}")
            self.refresh_stats('dna')
            self._show_summary(summary, out)

        self._run_async(self.cta_btns['dna'], work, done)

    # ================================================================
    # ЗАПУСК: УСАДЬБА
    # ================================================================
    def run_usadba(self):
        report_path = self.usd_var.get()
        if not report_path:
            messagebox.showerror("Ошибка", "Выберите файл отчёта «Доход по номеру»!")
            return
        try:
            sources = self._usadba_sources(report_path)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Файл не читается:\n{e}")
            return
        res = self._resolve_unrecognized('usd', migrate_agents(self.cfg), sources)
        if res[0] is None:
            return
        agents, source_map = res

        output_path = filedialog.asksaveasfilename(
            parent=self, title="Сохранить итоговый файл как...",
            defaultextension=".xlsx", filetypes=[("Excel файлы", "*.xlsx")])
        if not output_path:
            return
        quarter = f"{self.q_vars['usd'][0].get()}-{self.q_vars['usd'][1].get()}"
        tax = self.cfg.get('tax', {}).get('usd')

        def work():
            db, checker = self._start_history("Усадьба", quarter)
            if not self.dup_vars['usd'].get():
                checker = None
            stats, bookings = self.process_usadba(report_path, output_path,
                                                  dup_checker=checker, tax=tax,
                                                  agents=agents, source_map=source_map)
            dups = self._finish_history(db, "Усадьба", quarter, [report_path],
                                        output_path, bookings)
            return stats, output_path, dups

        def done(res):
            stats, out, dups = res
            self.set_status(f"Усадьба: готово, строк {stats['rows']}, дублей {dups}")
            self.refresh_stats('usd')
            msg = (f"Файл успешно сохранён!\n{out}\n\n"
                   f"Обработано записей: {stats['rows']}\n"
                   f"Номеров (комнат): {stats['rooms']}")
            warnings_list = []
            if stats['zero_price'] > 0:
                warnings_list.append(f"• Записей с нулевой стоимостью: {stats['zero_price']} (голубой)")
            if stats['bad_dates'] > 0:
                warnings_list.append(f"• Записей с нераспознанными датами: {stats['bad_dates']} (оранжевый)")
            if stats['still_living'] > 0:
                warnings_list.append(f"• Гостей «Сейчас проживает»: {stats['still_living']}")
            if stats.get('unknown_source', 0) > 0:
                warnings_list.append(f"• Нераспознанный источник (жёлтый): {stats['unknown_source']}")
            if dups > 0:
                warnings_list.append(f"• Переход кварталов (фиолетовые): {dups}")
            if warnings_list:
                msg += "\n\nОбратите внимание:\n" + "\n".join(warnings_list)
                messagebox.showwarning("Готово (есть замечания)", msg)
            else:
                messagebox.showinfo("Успех", msg)

        self._run_async(self.cta_btns['usd'], work, done)

    # ================================================================
    # ЛОГИКА: ЧАЙКОВСКИЙ
    # ================================================================
    def process_tchaikovsky(self, data_path, rooms_path, output_path,
                            map_data=None, map_rooms=None,
                            extra_data=None, dup_checker=None, tax=None,
                            agents=None, source_map=None):
        """Возвращает (errors_summary, bookings)."""
        tp = TaxParams(tax)
        log.info("Чайковский: %s / %s", data_path, rooms_path)
        df_data = read_table(data_path)
        df_rooms = read_table(rooms_path)

        headers_d = [str(c) for c in df_data.columns]
        headers_r = [str(c) for c in df_rooms.columns]
        map_data = map_data or guess_mapping(headers_d, TCH_DATA_SCHEMA)
        map_rooms = map_rooms or guess_mapping(headers_r, TCH_ROOMS_SCHEMA)
        miss = [f['label'] for f in TCH_DATA_SCHEMA if f['required'] and not map_data.get(f['key'])]
        if miss:
            raise ValueError(f"Не сопоставлены колонки: {', '.join(miss)}")
        if not (map_rooms.get('fio') and map_rooms.get('room')):
            raise ValueError("Не сопоставлены колонки файла «Номера» (ФИО, № комнаты)")
        extra_data = [c for c in (extra_data or []) if c in df_data.columns]

        room_mapping = {}
        for _, rr in df_rooms.iterrows():
            fio = rr[map_rooms['fio']]
            if pd.isna(fio):
                continue
            key = norm_name(fio)
            if key and key not in room_mapping:
                room_mapping[key] = rr[map_rooms['room']]

        def find_room(guest_string):
            if not guest_string:
                return None
            for guest in re.split(r'[,\n;]+', guest_string):
                key = norm_name(guest)
                if key in room_mapping:
                    room = room_mapping[key]
                    try:
                        return int(float(room))
                    except (ValueError, TypeError):
                        return room
            return None

        def get(row, key):
            col = map_data.get(key)
            return row[col] if col else None

        records = []
        for _, row in df_data.iterrows():
            guest = clean_guest_names(get(row, 'guest'))

            checkin = pd.to_datetime(get(row, 'checkin'), dayfirst=True, errors='coerce')
            checkout = pd.to_datetime(get(row, 'checkout'), dayfirst=True, errors='coerce')
            checkin = None if pd.isna(checkin) else checkin.to_pydatetime()
            checkout = None if pd.isna(checkout) else checkout.to_pydatetime()

            nights_v = get(row, 'nights')
            nights = None
            if nights_v is not None and not pd.isna(nights_v):
                try:
                    nights = int(float(nights_v))
                except (ValueError, TypeError):
                    nights = None
            if nights is None and checkin and checkout:
                nights = max(1, (checkout.date() - checkin.date()).days)

            category = get(row, 'category')
            category = '' if category is None or pd.isna(category) else str(category).strip()
            num_rooms = max(1, len([x for x in category.split('\n') if x.strip()])) if category else 1

            price = parse_price(get(row, 'price'))

            bron_v = get(row, 'bron')
            bron = '' if bron_v is None or pd.isna(bron_v) else str(bron_v).strip()

            source_v = get(row, 'source')
            source = '' if source_v is None or pd.isna(source_v) else str(source_v).strip()

            sm = (source_map or {}).get(source)
            if sm:
                commission, agent = float(sm[0]), sm[1]
            else:
                commission, agent = match_agent(
                    source, agents or default_agents(),
                    date=checkin.date() if checkin else None, hotel='tch')

            extra_vals = [row[c] if not pd.isna(row[c]) else '' for c in extra_data]

            records.append({
                'bron': bron, 'num_rooms': num_rooms, 'guest': guest, 'price': price,
                'commission': commission, 'agent': agent,
                'unknown_source': (not agent) and bool(source),
                'checkin': checkin, 'nights': nights, 'checkout': checkout,
                'source': source, 'category': category, 'room': find_room(guest),
                'extra_vals': extra_vals, 'dup_note': None,
            })

        if dup_checker:
            notes = dup_checker([{
                'bron': r['bron'], 'guest_norm': norm_name(r['guest']),
                'checkin': r['checkin'].date() if r['checkin'] else None,
                'checkout': r['checkout'].date() if r['checkout'] else None,
            } for r in records])
            for r, n in zip(records, notes):
                r['dup_note'] = n

        records.sort(key=lambda x: norm_name(x['guest']))

        errors_summary = {
            "Отсутствует № комнаты (КРАСНЫЙ)": 0,
            "Не указан Гость (ЖЕЛТЫЙ)": 0,
            "Отсутствуют Даты (ОРАНЖЕВЫЙ)": 0,
            "Отсутствует Стоимость (ГОЛУБОЙ)": 0,
            "Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)": 0,
            "Переход кварталов (ФИОЛЕТОВЫЙ)": 0,
        }

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        headers = ['№ брони', 'Номеров', 'Гость', 'Стоимость', 'Комиссия %', 'Агент',
                   'Стоимость без комиссии', 'Заезд', 'Ночей', 'Сут.ставка', 'Метод',
                   'Налог', 'Выезд', 'Источник', 'Категория номера', '№ комнаты',
                   'корпус', 'Примечание'] + extra_data
        n_cols = len(headers)
        font_bold = Font(bold=True)
        base = 'G' if tp.use_full else 'D'

        def write_sheet(name, recs, count_errors=False):
            ws = wb.create_sheet(name)
            for c, h in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=c, value=h)
                cell.font = font_bold
            for i, rec in enumerate(recs):
                r = i + 2
                ws.cell(row=r, column=1, value=rec['bron'])
                ws.cell(row=r, column=2, value=rec['num_rooms'])

                c_guest = ws.cell(row=r, column=3, value=rec['guest'])
                c_guest.alignment = Alignment(wrap_text=True, vertical='top')
                if not rec['guest']:
                    c_guest.fill = FILL_GUEST
                    if count_errors: errors_summary["Не указан Гость (ЖЕЛТЫЙ)"] += 1

                c_price = ws.cell(row=r, column=4, value=rec['price'])
                c_price.number_format = '#,##0.00'
                if rec['price'] is None:
                    c_price.fill = FILL_PRICE
                    if count_errors: errors_summary["Отсутствует Стоимость (ГОЛУБОЙ)"] += 1

                c_comm = ws.cell(row=r, column=5, value=rec['commission'])
                c_comm.number_format = '0.##'
                c_agent = ws.cell(row=r, column=6, value=rec['agent'] or '')
                if rec['unknown_source']:
                    c_comm.fill = FILL_GUEST
                    c_agent.fill = FILL_GUEST
                    if count_errors:
                        errors_summary["Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)"] += 1

                c_full = ws.cell(row=r, column=7,
                                 value=f'=ROUND(D{r}/(100-E{r})*100,2)')
                c_full.number_format = '#,##0.00'

                c_in = ws.cell(row=r, column=8, value=rec['checkin'])
                if rec['checkin']:
                    c_in.number_format = DATE_FMT
                else:
                    c_in.fill = FILL_DATE
                    if count_errors: errors_summary["Отсутствуют Даты (ОРАНЖЕВЫЙ)"] += 1

                ws.cell(row=r, column=9, value=rec['nights'])
                c_rate = ws.cell(row=r, column=10, value=tp.f_daily(base, 'I', r))
                c_rate.number_format = '#,##0.00'
                ws.cell(row=r, column=11, value=tp.f_method('J', r))
                c_tax = ws.cell(row=r, column=12, value=tp.f_tax('K', 'I', base, r))
                c_tax.number_format = '#,##0.00'

                c_out = ws.cell(row=r, column=13, value=rec['checkout'])
                if rec['checkout']:
                    c_out.number_format = DATE_FMT
                else:
                    c_out.fill = FILL_DATE
                    if count_errors: errors_summary["Отсутствуют Даты (ОРАНЖЕВЫЙ)"] += 1

                ws.cell(row=r, column=14, value=rec['source'])
                cat = ws.cell(row=r, column=15, value=rec['category'])
                cat.alignment = Alignment(wrap_text=True, vertical='top')

                c_room = ws.cell(row=r, column=16, value=rec['room'])
                if rec['room'] is None:
                    c_room.fill = FILL_ROOM
                    if count_errors: errors_summary["Отсутствует № комнаты (КРАСНЫЙ)"] += 1
                else:
                    ws.cell(row=r, column=17, value=f'=IF(P{r}<50,1,2)')

                c_note = ws.cell(row=r, column=18, value=rec['dup_note'] or '')
                if rec['dup_note']:
                    c_note.fill = FILL_DUP
                    if count_errors: errors_summary["Переход кварталов (ФИОЛЕТОВЫЙ)"] += 1

                for k, v in enumerate(rec['extra_vals']):
                    ws.cell(row=r, column=19 + k, value=v)

            if recs:
                last = len(recs) + 1
                tr = last + 1
                ws.cell(row=tr, column=3, value="ИТОГО:")
                ws.cell(row=tr, column=2, value=f'=SUM(B2:B{last})')
                ws.cell(row=tr, column=4, value=f'=SUM(D2:D{last})')
                ws.cell(row=tr, column=7, value=f'=SUM(G2:G{last})')
                ws.cell(row=tr, column=9, value=f'=SUM(I2:I{last})')
                ws.cell(row=tr, column=12, value=f'=SUM(L2:L{last})')
                for c in range(1, n_cols + 1):
                    cell = ws.cell(row=tr, column=c)
                    cell.font = font_bold
                    cell.fill = FILL_TOTALS

            widths = {'A': 24, 'B': 9, 'C': 35, 'D': 12, 'E': 10, 'F': 32, 'G': 20,
                      'H': 17, 'I': 8, 'J': 11, 'K': 12, 'L': 11, 'M': 17, 'N': 25,
                      'O': 30, 'P': 10, 'Q': 8, 'R': 45}
            for col, w in widths.items():
                ws.column_dimensions[col].width = w
            return ws

        def corpus_of(rec):
            try:
                return 1 if float(rec['room']) < 50 else 2
            except (ValueError, TypeError):
                return None

        write_sheet('1&2', records, count_errors=True)
        write_sheet('нет корпуса', [x for x in records if x['room'] is None])
        write_sheet('1', [x for x in records if corpus_of(x) == 1])
        write_sheet('2', [x for x in records if corpus_of(x) == 2])

        wb.save(output_path)
        log.info("Чайковский сохранён: %s", output_path)

        bookings = [{
            'guest': r['guest'], 'guest_norm': norm_name(r['guest']), 'bron': r['bron'],
            'checkin': r['checkin'].date().isoformat() if r['checkin'] else None,
            'checkout': r['checkout'].date().isoformat() if r['checkout'] else None,
            'price': r['price'],
            'price_full': round(r['price'] / (100 - r['commission']) * 100, 2)
                          if r['price'] is not None else None,
            'nights': r['nights'], 'dup_note': r['dup_note'],
        } for r in records]
        return errors_summary, bookings

    # ================================================================
    # ЛОГИКА: ДОНА
    # ================================================================
    @staticmethod
    def process_dona(data_path, output_path, map_data=None,
                     extra_data=None, dup_checker=None, tax=None,
                     agents=None, source_map=None):
        """Возвращает (summary, bookings)."""
        tp = TaxParams(tax)
        log.info("Дона: %s", data_path)
        df_data = read_table(data_path)
        headers = [str(c) for c in df_data.columns]
        map_data = map_data or guess_mapping(headers, DNA_DATA_SCHEMA)
        miss = [f['label'] for f in DNA_DATA_SCHEMA if f['required'] and not map_data.get(f['key'])]
        if miss:
            raise ValueError(f"Не сопоставлены колонки: {', '.join(miss)}")
        extra_data = [c for c in (extra_data or []) if c in df_data.columns]

        def get(row, key):
            col = map_data.get(key)
            return row[col] if col else None

        records = []
        for _, row in df_data.iterrows():
            guest = clean_guest_names(get(row, 'guest'))

            checkin = pd.to_datetime(get(row, 'checkin'), dayfirst=True, errors='coerce')
            checkout = pd.to_datetime(get(row, 'checkout'), dayfirst=True, errors='coerce')
            checkin = None if pd.isna(checkin) else checkin.to_pydatetime()
            checkout = None if pd.isna(checkout) else checkout.to_pydatetime()

            nights_v = get(row, 'nights')
            nights = None
            if nights_v is not None and not pd.isna(nights_v):
                try:
                    nights = int(float(nights_v))
                except (ValueError, TypeError):
                    nights = None
            if nights is None and checkin and checkout:
                nights = max(1, (checkout.date() - checkin.date()).days)

            category = get(row, 'category')
            category = '' if category is None or pd.isna(category) else str(category).strip()
            num_rooms = max(1, len([x for x in category.split('\n') if x.strip()])) if category else 1

            price = parse_price(get(row, 'price'))

            bron_v = get(row, 'bron')
            bron = '' if bron_v is None or pd.isna(bron_v) else str(bron_v).strip()

            source_v = get(row, 'source')
            source = '' if source_v is None or pd.isna(source_v) else str(source_v).strip()

            sm = (source_map or {}).get(source)
            if sm:
                commission, agent = float(sm[0]), sm[1]
            else:
                # комиссия берётся на дату заезда (плавающие проценты)
                commission, agent = match_agent(
                    source, agents or default_agents(),
                    date=checkin.date() if checkin else None, hotel='dna')
            extra_vals = [row[c] if not pd.isna(row[c]) else '' for c in extra_data]

            records.append({
                'bron': bron, 'num_rooms': num_rooms, 'guest': guest, 'price': price,
                'commission': commission, 'agent': agent, 'unknown_source': not agent,
                'checkin': checkin, 'nights': nights, 'checkout': checkout,
                'source': source, 'category': category,
                'extra_vals': extra_vals, 'dup_note': None,
            })

        if dup_checker:
            notes = dup_checker([{
                'bron': r['bron'], 'guest_norm': norm_name(r['guest']),
                'checkin': r['checkin'].date() if r['checkin'] else None,
                'checkout': r['checkout'].date() if r['checkout'] else None,
            } for r in records])
            for r, n in zip(records, notes):
                r['dup_note'] = n

        records.sort(key=lambda x: norm_name(x['guest']))

        summary = {
            "Не указан Гость (ЖЕЛТЫЙ)": 0,
            "Отсутствуют Даты (ОРАНЖЕВЫЙ)": 0,
            "Отсутствует Стоимость (ГОЛУБОЙ)": 0,
            "Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)": 0,
            "Переход кварталов (ФИОЛЕТОВЫЙ)": 0,
        }

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        headers_out = ['№ брони', 'Номеров', 'Гость', 'Стоимость', 'Комиссия %', 'Агент',
                       'Стоимость без комиссии', 'Заезд', 'Ночей', 'Сут.ставка', 'Метод',
                       'Налог', 'Выезд', 'Источник', 'Категория номера', 'Примечание'] + extra_data
        n_cols = len(headers_out)
        font_bold = Font(bold=True)

        ws = wb.create_sheet('Дона')
        for c, h in enumerate(headers_out, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = font_bold

        for i, rec in enumerate(records):
            r = i + 2
            ws.cell(row=r, column=1, value=rec['bron'])
            ws.cell(row=r, column=2, value=rec['num_rooms'])

            c_guest = ws.cell(row=r, column=3, value=rec['guest'])
            c_guest.alignment = Alignment(wrap_text=True, vertical='top')
            if not rec['guest']:
                c_guest.fill = FILL_GUEST
                summary["Не указан Гость (ЖЕЛТЫЙ)"] += 1

            c_price = ws.cell(row=r, column=4, value=rec['price'])
            c_price.number_format = '#,##0.00'
            if rec['price'] is None:
                c_price.fill = FILL_PRICE
                summary["Отсутствует Стоимость (ГОЛУБОЙ)"] += 1

            c_comm = ws.cell(row=r, column=5, value=rec['commission'])
            c_comm.number_format = '0'
            c_agent = ws.cell(row=r, column=6, value=rec['agent'] or '')
            if rec['unknown_source']:
                c_comm.fill = FILL_GUEST
                c_agent.fill = FILL_GUEST
                summary["Нераспознанный источник — комиссия 0% (ЖЕЛТЫЙ)"] += 1

            c_full = ws.cell(row=r, column=7, value=f'=ROUND(D{r}/(100-E{r})*100,2)')
            c_full.number_format = '#,##0.00'

            c_in = ws.cell(row=r, column=8, value=rec['checkin'])
            if rec['checkin']:
                c_in.number_format = DATE_FMT
            else:
                c_in.fill = FILL_DATE
                summary["Отсутствуют Даты (ОРАНЖЕВЫЙ)"] += 1

            ws.cell(row=r, column=9, value=rec['nights'])
            # налог от стоимости без комиссии (G) или как в документе (D)
            pcol = 'G' if tp.use_full else 'D'
            c_rate = ws.cell(row=r, column=10, value=tp.f_daily(pcol, 'I', r))
            c_rate.number_format = '#,##0.00'
            ws.cell(row=r, column=11, value=tp.f_method('J', r))
            c_tax = ws.cell(row=r, column=12, value=tp.f_tax('K', 'I', pcol, r))
            c_tax.number_format = '#,##0.00'

            c_out = ws.cell(row=r, column=13, value=rec['checkout'])
            if rec['checkout']:
                c_out.number_format = DATE_FMT
            else:
                c_out.fill = FILL_DATE
                summary["Отсутствуют Даты (ОРАНЖЕВЫЙ)"] += 1

            ws.cell(row=r, column=14, value=rec['source'])
            cat = ws.cell(row=r, column=15, value=rec['category'])
            cat.alignment = Alignment(wrap_text=True, vertical='top')

            c_note = ws.cell(row=r, column=16, value=rec['dup_note'] or '')
            if rec['dup_note']:
                c_note.fill = FILL_DUP
                summary["Переход кварталов (ФИОЛЕТОВЫЙ)"] += 1

            for k, v in enumerate(rec['extra_vals']):
                ws.cell(row=r, column=17 + k, value=v)

        if records:
            last = len(records) + 1
            tr = last + 1
            ws.cell(row=tr, column=3, value="ИТОГО:")
            ws.cell(row=tr, column=2, value=f'=SUM(B2:B{last})')
            ws.cell(row=tr, column=4, value=f'=SUM(D2:D{last})')
            ws.cell(row=tr, column=7, value=f'=SUM(G2:G{last})')
            ws.cell(row=tr, column=9, value=f'=SUM(I2:I{last})')
            ws.cell(row=tr, column=12, value=f'=SUM(L2:L{last})')
            for c in range(1, n_cols + 1):
                cell = ws.cell(row=tr, column=c)
                cell.font = font_bold
                cell.fill = FILL_TOTALS

        widths = {'A': 24, 'B': 9, 'C': 35, 'D': 12, 'E': 11, 'F': 30, 'G': 20,
                  'H': 17, 'I': 8, 'J': 11, 'K': 12, 'L': 11, 'M': 17, 'N': 35, 'O': 30, 'P': 45}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

        wb.save(output_path)
        log.info("Дона сохранена: %s", output_path)

        bookings = [{
            'guest': r['guest'], 'guest_norm': norm_name(r['guest']), 'bron': r['bron'],
            'checkin': r['checkin'].date().isoformat() if r['checkin'] else None,
            'checkout': r['checkout'].date().isoformat() if r['checkout'] else None,
            'price': r['price'],
            'price_full': round(r['price'] / (100 - r['commission']) * 100, 2)
                          if r['price'] is not None else None,
            'nights': r['nights'], 'dup_note': r['dup_note'],
        } for r in records]
        return summary, bookings

    # ================================================================
    # ЛОГИКА: УСАДЬБА
    # ================================================================
    @staticmethod
    def _load_workbook_safe(path):
        try:
            return openpyxl.load_workbook(path, data_only=True)
        except Exception:
            with zipfile.ZipFile(path) as z:
                items = {n: z.read(n) for n in z.namelist()}
            styles_name = 'xl/styles.xml'
            if styles_name in items:
                styles = items[styles_name].decode('utf-8')

                def fix_rgb(m):
                    v = m.group(1)
                    return f'rgb="{v.rjust(8, "0")}"' if len(v) < 8 else m.group(0)

                styles = re.sub(r'rgb="([0-9A-Fa-f]{1,8})"', fix_rgb, styles)
                items[styles_name] = styles.encode('utf-8')

            buf = io.BytesIO()
            with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
                for name, data in items.items():
                    z.writestr(name, data)
            buf.seek(0)
            return openpyxl.load_workbook(buf, data_only=True)

    def process_usadba(self, report_path, output_path, dup_checker=None,
                       tax=None, agents=None, source_map=None):
        """Возвращает (stats, bookings)."""
        tp = TaxParams(tax)
        log.info("Усадьба: %s", report_path)
        wb_in = self._load_workbook_safe(report_path)
        ws_in = wb_in[wb_in.sheetnames[0]]

        valid_states = ('Проживал', 'Сейчас проживает')

        def parse_dt(v):
            if isinstance(v, datetime):
                return v
            if v is None:
                return None
            s = str(v).strip()
            for fmt in ('%d.%m.%Y %H:%M', '%d.%m.%Y %H:%M:%S', '%d.%m.%Y'):
                try:
                    return datetime.strptime(s, fmt)
                except ValueError:
                    continue
            return None

        records = []
        rooms = set()
        current_room = None
        stats = {'rows': 0, 'rooms': 0, 'zero_price': 0, 'bad_dates': 0,
                 'still_living': 0, 'unknown_source': 0}

        for r in range(1, ws_in.max_row + 1):
            a_val = ws_in.cell(row=r, column=1).value
            if isinstance(a_val, str) and a_val.strip().startswith('Номер:'):
                current_room = a_val.split(':', 1)[1].strip()
                rooms.add(current_room)
                continue

            state = ws_in.cell(row=r, column=5).value
            fio = ws_in.cell(row=r, column=6).value
            if not fio or not state or str(state).strip() not in valid_states:
                continue

            checkin_raw = ws_in.cell(row=r, column=7).value
            checkout_raw = ws_in.cell(row=r, column=8).value
            price = parse_price(ws_in.cell(row=r, column=9).value)
            tariff = ws_in.cell(row=r, column=14).value
            org = ws_in.cell(row=r, column=16).value
            org_s = '' if org is None else str(org).strip()

            dt_in = parse_dt(checkin_raw)
            dt_out = parse_dt(checkout_raw)

            if dt_in and dt_out:
                nights = max(1, (dt_out.date() - dt_in.date()).days)
            else:
                nights = 1
                stats['bad_dates'] += 1

            if str(state).strip() == 'Сейчас проживает':
                stats['still_living'] += 1
            if not price:
                stats['zero_price'] += 1

            sm = (source_map or {}).get(org_s)
            if sm:
                commission, agent = float(sm[0]), sm[1]
            else:
                commission, agent = match_agent(
                    org_s, agents or default_agents(),
                    date=dt_in.date() if dt_in else None, hotel='usd')
            if not agent and org_s:
                stats['unknown_source'] += 1

            records.append({
                'fio': str(fio).strip(),
                'checkin': dt_in if dt_in else (str(checkin_raw or '').strip()),
                'checkout': dt_out if dt_out else (str(checkout_raw or '').strip()),
                'nights': nights,
                'price': price if price is not None else 0,
                'commission': commission, 'agent': agent,
                'unknown_source': (not agent) and bool(org_s),
                'tariff': str(tariff).strip() if tariff else '',
                'org': org_s,
                'room': current_room,
                'dup_note': None,
            })

        if not records:
            raise ValueError("В файле не найдено ни одной строки гостя.\n"
                             "Убедитесь, что выбран отчёт «Доход по номеру» из UCS Shelter.")

        if dup_checker:
            notes = dup_checker([{
                'bron': '', 'guest_norm': norm_name(rec['fio']),
                'checkin': rec['checkin'].date() if isinstance(rec['checkin'], datetime) else None,
                'checkout': rec['checkout'].date() if isinstance(rec['checkout'], datetime) else None,
            } for rec in records])
            for rec, n in zip(records, notes):
                rec['dup_note'] = n
        stats['dups'] = sum(1 for rec in records if rec['dup_note'])

        stats['rows'] = len(records)
        stats['rooms'] = len(rooms)

        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = 'Page 1'

        headers = ['Фамилия Имя Отчество', 'Заселение', 'ночей', 'Выселение ',
                   'Стоимость', 'Комиссия %', 'Агент', 'Стоимость без комиссии',
                   'Сут.ставка', 'Метод', 'Налог', 'Тариф',
                   'Организация\nзаказчик', 'Примечание']
        font_bold = Font(bold=True)
        for c, h in enumerate(headers, start=1):
            cell = ws_out.cell(row=1, column=c, value=h)
            cell.font = font_bold
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        base = 'H' if tp.use_full else 'E'

        for i, rec in enumerate(records):
            row = i + 2
            ws_out.cell(row=row, column=1, value=rec['fio'])

            c_in = ws_out.cell(row=row, column=2, value=rec['checkin'])
            c_out = ws_out.cell(row=row, column=4, value=rec['checkout'])
            if isinstance(rec['checkin'], datetime):
                c_in.number_format = DATE_FMT
            else:
                c_in.fill = FILL_DATE
            if isinstance(rec['checkout'], datetime):
                c_out.number_format = DATE_FMT
            else:
                c_out.fill = FILL_DATE

            ws_out.cell(row=row, column=3, value=rec['nights'])

            c_price = ws_out.cell(row=row, column=5, value=rec['price'])
            c_price.number_format = r'#,##0.00\ "₽"'
            if not rec['price']:
                c_price.fill = FILL_PRICE

            c_comm = ws_out.cell(row=row, column=6, value=rec['commission'])
            c_comm.number_format = '0.##'
            c_agent = ws_out.cell(row=row, column=7, value=rec['agent'] or '')
            if rec['unknown_source']:
                c_comm.fill = FILL_GUEST
                c_agent.fill = FILL_GUEST

            c_full = ws_out.cell(row=row, column=8,
                                 value=f'=ROUND(E{row}/(100-F{row})*100,2)')
            c_full.number_format = '#,##0.00'

            c_rate = ws_out.cell(row=row, column=9, value=tp.f_daily(base, 'C', row))
            c_rate.number_format = '#,##0.00'
            ws_out.cell(row=row, column=10, value=tp.f_method('I', row))
            c_tax = ws_out.cell(row=row, column=11, value=tp.f_tax('J', 'C', base, row))
            c_tax.number_format = '#,##0.00'
            ws_out.cell(row=row, column=12, value=rec['tariff'])
            ws_out.cell(row=row, column=13, value=rec['org'])

            c_note = ws_out.cell(row=row, column=14, value=rec['dup_note'] or '')
            if rec['dup_note']:
                c_note.fill = FILL_DUP

        last_row = len(records) + 1
        total_row = last_row + 1

        ws_out.cell(row=total_row, column=1, value="ИТОГО:")
        ws_out.cell(row=total_row, column=5, value=f'=SUM(E2:E{last_row})').number_format = r'#,##0.00\ "₽"'
        ws_out.cell(row=total_row, column=8, value=f'=SUM(H2:H{last_row})')
        ws_out.cell(row=total_row, column=11, value=f'=SUM(K2:K{last_row})')
        for c in range(1, 15):
            cell = ws_out.cell(row=total_row, column=c)
            cell.font = font_bold
            cell.fill = FILL_TOTALS

        widths = {'A': 38, 'B': 18, 'C': 8, 'D': 18, 'E': 14, 'F': 10, 'G': 30,
                  'H': 18, 'I': 12, 'J': 12, 'K': 12, 'L': 10, 'M': 40, 'N': 45}
        for col, w in widths.items():
            ws_out.column_dimensions[col].width = w

        ws2 = wb_out.create_sheet('Лист1')
        ws2.cell(row=1, column=1, value='Заселение').font = font_bold
        ws2.cell(row=1, column=2, value='Выселение ').font = font_bold
        for i, rec in enumerate(records):
            row = i + 2
            c_a = ws2.cell(row=row, column=1, value=rec['checkin'])
            c_b = ws2.cell(row=row, column=2, value=rec['checkout'])
            if isinstance(rec['checkin'], datetime):
                c_a.number_format = DATE_FMT
            if isinstance(rec['checkout'], datetime):
                c_b.number_format = DATE_FMT
            ws2.cell(row=row, column=4, value=f'=INT(A{row})')
            ws2.cell(row=row, column=5, value=f'=INT(B{row})')
            ws2.cell(row=row, column=6, value=f'=E{row}-D{row}')
        ws2.column_dimensions['A'].width = 18
        ws2.column_dimensions['B'].width = 18

        wb_out.save(output_path)
        log.info("Усадьба сохранена: %s", output_path)

        bookings = [{
            'guest': rec['fio'], 'guest_norm': norm_name(rec['fio']), 'bron': '',
            'checkin': rec['checkin'].date().isoformat() if isinstance(rec['checkin'], datetime) else None,
            'checkout': rec['checkout'].date().isoformat() if isinstance(rec['checkout'], datetime) else None,
            'price': rec['price'],
            'price_full': round(rec['price'] / (100 - rec['commission']) * 100, 2)
                          if rec['price'] else None,
            'nights': rec['nights'], 'dup_note': rec['dup_note'],
        } for rec in records]
        return stats, bookings



if __name__ == "__main__":
    log.info("Запуск приложения v%s", APP_VERSION)
    _cfg = load_config()
    _db_path = _cfg.get('db_path')
    if _db_path and not os.path.isdir(os.path.dirname(_db_path) or _app_dir()):
        log.warning("Папка БД недоступна (%s) — путь по умолчанию", _db_path)
        _db_path = None
    if not _db_path:  # первый запуск: предлагаем расположение базы
        _r0 = ctk.CTk()
        _r0.withdraw()
        _dlg0 = DbLocationDialog(_r0)
        _r0.wait_window(_dlg0)
        _db_path = _dlg0.result or os.path.join(_app_dir(), 'turnalog_history.db')
        _r0.destroy()
        _cfg['db_path'] = _db_path
        save_config(_cfg)
    set_db_path(_db_path)
    log.info("БД: %s", _db_path)
    _db = HistoryDB()
    _db.ensure_default_admin()
    _root = ctk.CTk()
    _root.withdraw()
    _login = LoginDialog(_root, _db)
    _root.wait_window(_login)
    _user = _login.user
    _root.destroy()
    _db.close()
    if _user is None:
        sys.exit(0)
    log.info("Вход: %s (%s)", _user['login'], _user['role'])
    app = App(_user)
    app.mainloop()
