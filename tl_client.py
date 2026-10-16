# -*- coding: utf-8 -*-
"""Импорт броней и номерного фонда из TravelLine API.

Покрывает Чайковский (8186) и Дону (479): загружает изменённые брони,
отбирает ночи выбранного периода внутри квартала, раскладывает проживание/
допуслуги/каналы и пишет Excel той же формы,
что ожидает обработка (канонический маппинг — диалог колонок не нужен).
Read Reservation возвращает идентификаторы комнат в roomStays[].roomId/room;
для Чайковского они сопоставляются с PMS-номерами и корпусами. Используется
вложенный инвентарь и, если там нет комнат, постраничный PropertyRoom /rooms.
Неоднозначные или неназначенные комнаты остаются пустыми; ID API не подменяется
номером комнаты.
"""
import hashlib
import json
import re
import time
from collections import defaultdict, deque
from email.utils import parsedate_to_datetime
import math
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone

import openpyxl

HOST = 'https://partner.tlintegration.com'
HTTP_TIMEOUT_SECONDS = 90
NETWORK_RETRIES = 1
NETWORK_RETRY_DELAY = 2
TL_DETAILS_REQUESTS_PER_HOUR = 3800  # ниже лимита TravelLine с запасом
TL_DETAILS_MIN_INTERVAL_SECONDS = 0.36
# Increment if the prepared-row/report cache format changes incompatibly.
TL_CACHE_VERSION = 7

# Канонические колонки выгрузки (имена пересекаются с алиасами схем,
# но маппинг задаётся явно — автоподбор не нужен).
COL_GUEST = 'Гость'
COL_PRICE = 'Стоимость'
COL_IN = 'Заезд'
COL_OUT = 'Выезд'
COL_SOURCE = 'Источник'
COL_CAT = 'Категория номера'
COL_BRON = '№ брони'
COL_NIGHTS = 'Ночей'
COL_TARIFF = 'Тариф'
COL_SVC = 'Допуслуги'
COL_SVC_NAME = 'Наименование допуслуг'
COL_SVC_EXTENSION = 'Сумма продления номера'
COL_TOTAL = 'Итог брони'
COL_ROOM = '№ комнаты'
COL_CORPUS = 'Корпус'
COL_BUILDING_NAME = 'Здание TravelLine'
COL_TL_DISCOUNT = 'Скидка TravelLine, ₽'
COL_TL_DISCOUNT_PERCENT = 'Процент скидки TravelLine'
COL_TL_DISCOUNT_DATES = 'Предполагаемые ночи со скидкой TravelLine'
COL_TL_DISCOUNT_CAMPAIGN = 'Название скидки/кампании TravelLine'
COL_TL_DISCOUNT_ADDED = 'Скидка добавлена к налоговой базе, ₽'
COL_TL_DISCOUNT_RETAINED = 'Скидка сохранена в базе, ₽'
COL_TL_DISCOUNT_RULE = 'Правило скидки TravelLine'
COL_TL_TAX_BASE = 'Налоговая база TravelLine, ₽'
COL_TL_PERIOD_PRICE = 'Стоимость проживания за выбранный период TravelLine, ₽'
COL_TL_PERIOD_STATUS = 'Распределение по периоду TravelLine'
COL_TL_DAILY_SUM = 'Сумма dailyRates'
COL_TL_DAILY_CHECK = 'Сверка dailyRates'
ROOM_LABEL_TEXT_PREFIX = '__TL_ROOM_LABEL__:'

TL_DATA_MAP = {'guest': COL_GUEST, 'price': COL_PRICE, 'checkin': COL_IN,
               'checkout': COL_OUT, 'source': COL_SOURCE,
               'category': COL_CAT, 'bron': COL_BRON, 'nights': COL_NIGHTS,
               'room': COL_ROOM, 'corpus': COL_CORPUS}
TL_EXTRA = [COL_TARIFF, COL_SVC, COL_SVC_NAME, COL_SVC_EXTENSION, COL_TOTAL,
            COL_TL_DISCOUNT, COL_TL_DISCOUNT_PERCENT,
            COL_TL_DISCOUNT_DATES, COL_TL_DAILY_SUM, COL_TL_DAILY_CHECK,
            COL_TL_DISCOUNT_CAMPAIGN, COL_TL_DISCOUNT_ADDED,
            COL_TL_DISCOUNT_RETAINED, COL_TL_DISCOUNT_RULE,
            COL_TL_TAX_BASE, COL_TL_PERIOD_PRICE, COL_TL_PERIOD_STATUS]
TL_ROOMS_MAP = {'fio': 'ФИО', 'room': COL_ROOM, 'corpus': COL_CORPUS,
                'building_name': COL_BUILDING_NAME, 'bron': COL_BRON,
                'checkin': COL_IN, 'checkout': COL_OUT}


class TLApiError(Exception):
    pass


class TLNetworkError(TLApiError):
    """Временная ошибка соединения/тайм-аут, которую можно повторить."""


class TLRateLimitError(TLApiError):
    """HTTP 429 с необязательным интервалом Retry-After от TravelLine."""

    def __init__(self, message, retry_after_seconds=None):
        super().__init__(message)
        self.retry_after_seconds = retry_after_seconds


class TLCancelled(Exception):
    """Загрузка отменена пользователем."""


def is_tl_export(path):
    """Файл, сгенерированный импортом TravelLine (канонический маппинг)."""
    try:
        base = str(path).replace('\\', '/').rsplit('/', 1)[-1]
    except Exception:
        return False
    return base.startswith('TL_') and (
        base.endswith('_data.xlsx') or base.endswith('_rooms.xlsx'))


def tl_export_paths(tmpdir, code, year, quarter):
    stem = f'TL_{code}_{year}{quarter}_'
    return (tmpdir + '/' + stem + 'data.xlsx',
            tmpdir + '/' + stem + 'rooms.xlsx')


def _campaign_name_key(value):
    return re.sub(r'\s+', ' ',
                  str(value or '').strip().casefold().replace('ё', 'е'))


def _campaign_name_list(values):
    if isinstance(values, str):
        values = values.splitlines() or [values]
    names, seen = [], set()
    for value in values or []:
        name = str(value or '').strip()
        key = _campaign_name_key(name)
        if key and key not in seen:
            seen.add(key)
            names.append(name)
    return names


def inclusive_period_bounds(start_value, end_value):
    """Convert inclusive YYYY-MM-DD dates to an ISO [start, end) interval."""
    try:
        start = date.fromisoformat(str(start_value)[:10])
        end = date.fromisoformat(str(end_value)[:10])
    except (TypeError, ValueError):
        raise ValueError('Даты периода должны быть в формате ГГГГ-ММ-ДД.')
    if end < start:
        raise ValueError('Дата окончания периода должна быть не раньше даты начала.')
    return start.isoformat(), (end + timedelta(days=1)).isoformat()


def period_preset_options(mode, quarter_start, quarter_end_exclusive):
    """Build month/week/quarter presets clipped to one report quarter.

    Weeks use Monday–Sunday (ISO-style) boundaries and are clipped at quarter
    edges so an import for one report quarter never includes nights from another.
    Each tuple is (label, start_iso, end_exclusive_iso).
    """
    def as_date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return date.fromisoformat(str(value)[:10])
        except (TypeError, ValueError):
            raise ValueError('Некорректные границы квартала.')

    start, end = as_date(quarter_start), as_date(quarter_end_exclusive)
    if end <= start:
        raise ValueError('Конец квартала должен быть позже его начала.')
    mode = str(mode or '').strip().casefold()
    if mode == 'quarter':
        return [('Весь квартал', start.isoformat(), end.isoformat())]

    options = []
    if mode == 'month':
        names = ('Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
                 'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь')
        cursor = start.replace(day=1)
        while cursor < end:
            if cursor.month == 12:
                next_month = date(cursor.year + 1, 1, 1)
            else:
                next_month = date(cursor.year, cursor.month + 1, 1)
            period_start, period_end = max(start, cursor), min(end, next_month)
            label = f'{names[cursor.month - 1]} {cursor.year}'
            options.append((label, period_start.isoformat(), period_end.isoformat()))
            cursor = next_month
        return options

    if mode == 'week':
        cursor = start - timedelta(days=start.weekday())
        while cursor < end:
            period_start = max(start, cursor)
            period_end = min(end, cursor + timedelta(days=7))
            week_number = cursor.isocalendar()[1]
            period_last = period_end - timedelta(days=1)
            label = (f'Неделя {week_number} · '
                     f'{period_start:%d.%m}–{period_last:%d.%m}')
            options.append((label, period_start.isoformat(), period_end.isoformat()))
            cursor += timedelta(days=7)
        return options

    raise ValueError('Тип периода должен быть month, week или quarter.')


def tl_import_cache_key(hotel_code, property_id, period_start, period_end,
                        modified_since, campaigns_with_discount=None):
    campaign_keys = sorted({_campaign_name_key(name) for name in
                            _campaign_name_list(campaigns_with_discount)})
    campaign_hash = hashlib.sha256(
        json.dumps(campaign_keys, ensure_ascii=False).encode('utf-8')
    ).hexdigest()[:12]
    return (f'tlcache-v{TL_CACHE_VERSION}|{hotel_code}|{property_id}|'
            f'{period_start}|{period_end}|{modified_since}|camp={campaign_hash}')


def encode_tl_import_cache(rows, report):
    """Сериализует только строки, нужные импорту (без сырого API JSON/контактов)."""
    def default(value):
        if hasattr(value, 'isoformat'):
            return value.isoformat()
        return str(value)
    return json.dumps({'rows': rows, 'report': report}, ensure_ascii=False,
                      default=default)


def decode_tl_import_cache(payload):
    """Восстанавливает строки кэша, включая Excel-совместимые даты."""
    data = json.loads(payload)
    rows, report = data['rows'], data['report']
    for row in rows:
        for col in (COL_IN, COL_OUT):
            value = row.get(col)
            if isinstance(value, str) and value:
                try:
                    row[col] = datetime.fromisoformat(value)
                except ValueError:
                    pass
    return rows, report


def _retry_after_seconds(value):
    if value in (None, ''):
        return None
    try:
        seconds = float(value)
        return max(0.0, seconds) if math.isfinite(seconds) else None
    except (TypeError, ValueError):
        try:
            retry_at = parsedate_to_datetime(str(value))
            if retry_at.tzinfo is None:
                retry_at = retry_at.replace(tzinfo=timezone.utc)
            return max(0.0, (retry_at - datetime.now(timezone.utc)).total_seconds())
        except (TypeError, ValueError, OverflowError):
            return None


def _http(method, path, token=None, params=None, data=None, ctype=None):
    url = HOST + path
    if params:
        url += '?' + urllib.parse.urlencode(params)
    body = data.encode('utf-8') if isinstance(data, str) else data
    req = urllib.request.Request(url, data=data if body is None else body,
                                 method=method)
    if token:
        req.add_header('Authorization', 'Bearer ' + token)
    if ctype:
        req.add_header('Content-Type', ctype)
    try:
        with urllib.request.urlopen(req, timeout=HTTP_TIMEOUT_SECONDS) as r:
            return r.status, json.loads(r.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        try:
            detail = e.read().decode('utf-8', 'replace')[:300]
        except Exception:
            detail = ''
        if e.code == 429:
            headers = getattr(e, 'headers', None)
            retry_after = _retry_after_seconds(
                headers.get('Retry-After') if headers else None)
            raise TLRateLimitError(
                f'HTTP 429: {detail}', retry_after_seconds=retry_after)
        raise TLApiError(f'HTTP {e.code}: {detail}')
    except (urllib.error.URLError, OSError) as e:
        reason = getattr(e, 'reason', e)
        reason_text = str(reason)
        if (isinstance(reason, TimeoutError)
                or 'timed out' in reason_text.lower()
                or 'time out' in reason_text.lower()):
            detail = f'Тайм-аут ответа TravelLine ({HTTP_TIMEOUT_SECONDS} с)'
        else:
            detail = f'Ошибка соединения с TravelLine: {reason_text}'
        raise TLNetworkError(detail) from e


def _http_with_retry(method, path, token=None, params=None, data=None,
                     ctype=None, cancel_event=None):
    """Повторяет один раз только временные сетевые ошибки, не HTTP-ошибки."""
    for attempt in range(NETWORK_RETRIES + 1):
        if cancel_event is not None and cancel_event.is_set():
            raise TLCancelled()
        try:
            return _http(method, path, token, params, data, ctype)
        except TLNetworkError as exc:
            if attempt >= NETWORK_RETRIES:
                tries = attempt + 1
                raise TLNetworkError(
                    f'{exc}; не удалось после {tries} попыток') from exc
            if cancel_event is not None:
                if cancel_event.wait(NETWORK_RETRY_DELAY):
                    raise TLCancelled()
            else:
                time.sleep(NETWORK_RETRY_DELAY)


def get_token(client_id, secret, cancel_event=None):
    try:
        _, data = _http_with_retry(
            'POST', '/auth/token',
            data=urllib.parse.urlencode(
                {'grant_type': 'client_credentials',
                 'client_id': client_id,
                 'client_secret': secret}),
            ctype='application/x-www-form-urlencoded',
            cancel_event=cancel_event)
        return data['access_token']
    except (TLApiError, KeyError) as e:
        raise TLApiError(f'Не удалось получить токен ({e}).')


def fetch_details(client_id, secret, property_id, depth='2024-01-01',
                  progress=None, cancel_event=None,
                  include_room_inventory=False, details_callback=None,
                  retain_details=True, details_batch_size=100,
                  period_start=None, period_end_exclusive=None,
                  cached_details=None):
    """Получает детали броней, при необходимости — номерной фонд.

    Если переданы ``period_start`` и ``period_end_exclusive``, список броней
    ищется через PMS PropertyReservation только в этом периоде, отдельно для
    Active/Cancelled. Это не запускает полную историю. Без этих аргументов
    действует прежний Read Reservation поток по ``lastModification``; при
    ``depth=None`` он запрашивает всю доступную историю.

    При ``include_room_inventory`` возвращается также ответ PMS API. Callback
    ``details_callback`` получает партии деталей; ``cached_details`` содержит
    уже сохранённые детали для возобновления загрузки без повторных запросов.
    ``retain_details=False`` отключает накопление ответов в памяти.
    progress(stage, done, total); HTTP 401 обновляет OAuth-токен один раз.
    Ошибки никогда не содержат client_secret.
    """
    def cancelled():
        return cancel_event is not None and cancel_event.is_set()

    def wait_with_progress(seconds, attempt=0):
        remaining = max(0.0, float(seconds))
        last_reported = None
        while remaining > 0:
            if cancelled():
                raise TLCancelled()
            seconds_left = math.ceil(remaining)
            if (progress and seconds_left != last_reported
                    and (last_reported is None or seconds_left % 10 == 0
                         or seconds_left <= 5)):
                progress('rate_limit', seconds_left, attempt)
                last_reported = seconds_left
            interval = min(10.0, remaining)
            if cancel_event is not None:
                if cancel_event.wait(interval):
                    raise TLCancelled()
            else:
                time.sleep(interval)
            remaining -= interval

    if cancelled():
        raise TLCancelled()
    token = get_token(client_id, secret, cancel_event)

    def api_get(path, params=None, stage='TravelLine API'):
        """Выполнить GET; на 401 получить свежий token и повторить один раз."""
        nonlocal token

        def contextual(error):
            text = str(error)
            if isinstance(error, TLNetworkError):
                return TLApiError(
                    f'{stage}: {text}. Проверьте соединение с интернетом; '
                    'TravelLine мог временно не ответить, повторите запрос позже.')
            if 'HTTP 403' in text:
                stage_key = stage.casefold()
                if 'номерной фонд' in stage_key:
                    permission = 'чтение номерного фонда (PMS API/PropertyRoom)'
                elif 'pms api' in stage_key or 'период' in stage_key:
                    permission = 'поиск броней за период (PMS API/PropertyReservation)'
                else:
                    permission = 'Read Reservation API'
                return TLApiError(
                    f'{stage}: доступ запрещён (403). В подключении TravelLine '
                    f'проверьте разрешение {permission} и доступ этого объекта. '
                    f'Ответ API: {text}')
            if 'HTTP 401' in text:
                return TLApiError(
                    f'{stage}: TravelLine не принял токен (401 Unauthorized). '
                    'Проверьте client_id/client_secret именно из активного '
                    'подключения этого объекта и повторите проверку. '
                    f'Ответ API: {text}')
            return TLApiError(f'{stage}: {text}')

        try:
            return _http_with_retry(
                'GET', path, token, params, cancel_event=cancel_event)
        except TLApiError as first:
            if 'HTTP 429' in str(first):
                # Сначала уважаем Retry-After, иначе используем растущий backoff.
                # Ждать можно отменить; детали предыдущих броней сохраняются batch-callback.
                rate_error = first
                max_retries = 8
                for retry in range(max_retries):
                    retry_after = getattr(rate_error, 'retry_after_seconds', None)
                    if retry_after is None:
                        delay = min(60 * (2 ** retry), 1800)
                    else:
                        delay = max(1.0, retry_after)
                    wait_with_progress(delay, retry + 1)
                    try:
                        return _http_with_retry(
                            'GET', path, token, params,
                            cancel_event=cancel_event)
                    except TLCancelled:
                        raise
                    except TLApiError as retry_error:
                        if 'HTTP 429' not in str(retry_error):
                            raise contextual(retry_error)
                        rate_error = retry_error
                raise contextual(rate_error)
            if 'HTTP 401' not in str(first):
                raise contextual(first)
            # Токен действует 15 минут. Обновляем его по 401 и повторяем
            # конкретный запрос один раз; не оставляем сырой ответ без этапа.
            try:
                token = get_token(client_id, secret, cancel_event)
            except TLApiError as auth_error:
                raise TLApiError(
                    f'{stage}: API вернул 401, а получить новый токен не удалось. '
                    f'{auth_error}')
            try:
                return _http_with_retry(
                    'GET', path, token, params, cancel_event=cancel_event)
            except TLApiError as retry_error:
                if 'HTTP 401' in str(retry_error):
                    stage_key = stage.casefold()
                    if 'номерной фонд' in stage_key:
                        permission = 'PMS API/PropertyRoom'
                    elif 'pms api' in stage_key or 'период' in stage_key:
                        permission = 'PMS API/PropertyReservation'
                    else:
                        permission = 'Read Reservation API'
                    raise TLApiError(
                        f'{stage}: HTTP 401 даже после выдачи нового токена. '
                        'TravelLine не принимает авторизацию этого API-запроса. '
                        'Проверьте пару client_id/client_secret из одного '
                        f'подключения и разрешение {permission}. '
                        f'Ответ API: {retry_error}')
                raise contextual(retry_error)

    inventory = None
    if include_room_inventory:
        if cancelled():
            raise TLCancelled()
        if progress:
            progress('inventory', 0, 1)
        inventory_error = None
        try:
            _, inventory = api_get(
                f'/api/pms/v2/properties/{property_id}'
                '/obtain-accommodation-inventory',
                stage='Номерной фонд (комнаты и корпуса)')
            if not isinstance(inventory, dict):
                inventory = {}
        except TLCancelled:
            raise
        except TLApiError as exc:
            inventory_error = exc
            inventory = {}

        # The documented PropertyRoom /rooms endpoint returns a flat, paginated
        # list. Use it when the legacy accommodation-inventory response has no
        # mappable room IDs (or is unavailable), and preserve building data if
        # the first endpoint did return it.
        if not accommodation_room_map(inventory):
            property_rooms = []
            page_token = None
            seen_page_tokens = set()
            page = 0
            try:
                while True:
                    if cancelled():
                        raise TLCancelled()
                    page += 1
                    params = {'maxPageSize': 100}
                    if page_token:
                        if page_token in seen_page_tokens:
                            raise TLApiError(
                                'TravelLine повторил pageToken номерного фонда.')
                        seen_page_tokens.add(page_token)
                        params['pageToken'] = page_token
                    _, data = api_get(
                        f'/api/pms/v2/properties/{property_id}/rooms', params,
                        f'Номерной фонд PropertyRoom (страница {page})')
                    if not isinstance(data, dict):
                        raise TLApiError(
                            'TravelLine вернул некорректный список комнат.')
                    page_rooms = data.get('rooms') or []
                    if not isinstance(page_rooms, list):
                        raise TLApiError(
                            'TravelLine вернул некорректный список комнат.')
                    property_rooms.extend(
                        room for room in page_rooms if isinstance(room, dict))
                    if not data.get('hasNextPage'):
                        break
                    page_token = data.get('nextPageToken')
                    if not page_token:
                        raise TLApiError(
                            'TravelLine сообщил о следующей странице комнат, '
                            'но не вернул nextPageToken.')
                    time.sleep(0.4)
            except TLCancelled:
                raise
            except TLApiError as rooms_error:
                if inventory_error is not None:
                    raise TLApiError(
                        f'{inventory_error}; список комнат PropertyRoom: '
                        f'{rooms_error}') from rooms_error
                raise
            if property_rooms:
                inventory = dict(inventory)
                inventory['propertyRooms'] = property_rooms
            elif inventory_error is not None:
                raise inventory_error

        if not isinstance(inventory, dict):
            inventory = {}
        if cancelled():
            raise TLCancelled()
        if progress:
            progress('inventory', 1, 1)

    summaries = []
    period_query = period_start is not None or period_end_exclusive is not None
    if period_query:
        if period_start is None or period_end_exclusive is None:
            raise ValueError('Для поиска по периоду нужны обе граничные даты.')
        try:
            from_day = date.fromisoformat(str(period_start)[:10])
            to_day = date.fromisoformat(str(period_end_exclusive)[:10])
        except (TypeError, ValueError):
            raise ValueError('Границы периода должны быть датами ГГГГ-ММ-ДД.')
        if to_day <= from_day:
            raise ValueError('Конец периода должен быть позже его начала.')
        if (to_day - from_day).days > 365:
            raise ValueError('TravelLine ограничивает поиск периода 365 днями.')

        seen_numbers = set()
        for state in ('Active', 'Cancelled'):
            page_token = None
            seen_page_tokens = set()
            page = 0
            while True:
                if cancelled():
                    raise TLCancelled()
                page += 1
                params = {
                    'state': state,
                    'startAffectPeriodDateTime': f'{from_day.isoformat()}T00:00',
                    'endAffectPeriodDateTime': f'{to_day.isoformat()}T00:00',
                    'maxPageSize': 100,
                }
                if page_token:
                    if page_token in seen_page_tokens:
                        raise TLApiError(
                            'TravelLine повторил pageToken поиска броней за период.')
                    seen_page_tokens.add(page_token)
                    params['pageToken'] = page_token
                _, data = api_get(
                    f'/api/pms/v2/properties/{property_id}/reservations/search',
                    params,
                    f'PMS API: список броней за период ({state}, страница {page})')
                if cancelled():
                    raise TLCancelled()
                if not isinstance(data, dict):
                    raise TLApiError(
                        'TravelLine вернул некорректный ответ поиска броней за период.')
                page_items = data.get('reservations') or []
                if not isinstance(page_items, list):
                    raise TLApiError(
                        'TravelLine вернул некорректный список броней за период.')
                for item in page_items:
                    if not isinstance(item, dict):
                        continue
                    number = (item.get('number') or item.get('bookingNumber'))
                    if number is None or str(number).strip() == '':
                        raise TLApiError(
                            'В поиске TravelLine за период отсутствует номер брони.')
                    number = str(number)
                    if number not in seen_numbers:
                        summaries.append({'number': number})
                        seen_numbers.add(number)
                if progress:
                    progress('list', len(summaries), 0)
                if not data.get('hasNextPage'):
                    break
                page_token = data.get('nextPageToken')
                if not page_token:
                    raise TLApiError(
                        'TravelLine сообщил о следующей странице, но не вернул nextPageToken.')
                time.sleep(0.4)
    else:
        cont, page = None, 0
        seen_tokens = set()
        while True:
            if cancelled():
                raise TLCancelled()
            page += 1
            params = {'count': 1000}
            if cont:
                if cont in seen_tokens:
                    raise TLApiError(
                        'TravelLine повторил continueToken списка броней.')
                seen_tokens.add(cont)
                params['continueToken'] = cont
            if depth:
                params['lastModification'] = str(depth)[:10] + 'T00:00:00Z'
            _, data = api_get(
                f'/api/read-reservation/v1/properties/{property_id}/bookings',
                params, f'Список броней (страница {page})')
            if cancelled():
                raise TLCancelled()
            summaries.extend(data.get('bookingSummaries', []))
            if progress:
                progress('list', len(summaries), 0)
            if not data.get('hasMoreData'):
                break
            cont = data.get('continueToken')
            if not cont:
                raise TLApiError(
                    'TravelLine сообщил о следующей странице, но не вернул continueToken.')
            time.sleep(0.4)
    details = []
    detail_batch = []
    batch_size = max(1, int(details_batch_size or 100))
    total = len(summaries)
    cached_by_number = {
        str(number): booking for number, booking in (cached_details or {}).items()
        if isinstance(booking, dict)
    }
    detail_request_times = deque()
    if progress:
        progress('details', 0, total)

    def wait_for_hourly_slot():
        while True:
            if cancelled():
                raise TLCancelled()
            now = time.monotonic()
            while (detail_request_times
                   and now - detail_request_times[0] >= 3600):
                detail_request_times.popleft()
            if len(detail_request_times) < TL_DETAILS_REQUESTS_PER_HOUR:
                return
            wait_with_progress(3600 - (now - detail_request_times[0]))

    def flush_detail_batch():
        nonlocal detail_batch
        if detail_batch and details_callback:
            batch, detail_batch = detail_batch, []
            details_callback(batch)

    try:
        for i, summary in enumerate(summaries, 1):
            if cancelled():
                raise TLCancelled()
            number = ((summary.get('number') or summary.get('bookingNumber'))
                      if isinstance(summary, dict) else None)
            if not number:
                raise TLApiError('В списке броней TravelLine отсутствует номер брони.')
            number = str(number)
            booking = cached_by_number.get(number)
            if booking is None:
                wait_for_hourly_slot()
                _, data = api_get(
                    f'/api/read-reservation/v1/properties/{property_id}'
                    f'/bookings/{number}',
                    stage=f'Бронь {number}')
                booking = data.get('booking', {})
                if not isinstance(booking, dict):
                    booking = {}
                if not booking.get('number'):
                    booking = dict(booking, number=number)
                detail_request_times.append(time.monotonic())
                if retain_details:
                    details.append(booking)
                if details_callback:
                    detail_batch.append(booking)
                    if len(detail_batch) >= batch_size:
                        flush_detail_batch()
                if i < total:
                    # Keep below the minute quota even when responses are fast.
                    time.sleep(TL_DETAILS_MIN_INTERVAL_SECONDS)
            elif retain_details:
                details.append(booking)
            if progress:
                progress('details', i, total)
    except Exception:
        # Checkpoint successful details even when a later request is rate-limited
        # or the user cancels; a retry will skip these booking numbers.
        flush_detail_batch()
        raise
    flush_detail_batch()
    return (details, inventory) if include_room_inventory else details


def _money(total):
    try:
        return float((total or {}).get('priceAfterTax') or 0)
    except (TypeError, ValueError):
        return 0.0


def _number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError, OverflowError):
        return None


def _discount_details(stay, campaigns_with_discount=None):
    """Разбирает скидки и решает, какие из них вернуть в стоимость.

    По принятому бизнес-правилу сумма скидки добавляется обратно в налоговую
    базу. Это расчётная корректировка, а не отдельная комиссия из API. Названия
    из ручного списка — исключения: для них цена остаётся
    со скидкой. Дату скидки по-прежнему выводим только при единственном
    совпадении агрегированной скидки с группой nightly rates.
    """
    total = stay.get('total') or {}
    discounts = total.get('discounts') or []
    if isinstance(discounts, dict):
        discounts = [discounts]
    discounts = [item for item in discounts if isinstance(item, dict)]
    if not discounts:
        return {
            'amount': 0.0, 'percent': '', 'dates': '', 'campaigns': '',
            'added': 0.0, 'retained': 0.0, 'rule': 'Скидки нет',
            'missing_campaign_name': False,
        }

    exception_keys = {_campaign_name_key(name) for name in
                      _campaign_name_list(campaigns_with_discount)}
    amount = added = retained = 0.0
    percent_labels, campaign_names = [], []
    missing_campaign_name = False
    for item in discounts:
        item_amount = _number(item.get('amount')) or 0.0
        amount += item_amount
        name = str(item.get('name') or '').strip()
        if name and name not in campaign_names:
            campaign_names.append(name)
        if name and _campaign_name_key(name) in exception_keys:
            retained += item_amount
        else:
            added += item_amount
            if exception_keys and not name:
                missing_campaign_name = True

        percent = _number(item.get('percent'))
        if percent is not None:
            label = f'{percent:g}%'
        elif str(item.get('discountType') or '').casefold() == 'fixed':
            label = 'фиксированная'
        else:
            label = str(item.get('discountType') or 'скидка')
        if label not in percent_labels:
            percent_labels.append(label)

    amount = round(amount, 2)
    added = round(added, 2)
    retained = round(retained, 2)
    if retained and added:
        rule = 'Часть скидки добавлена; кампания-исключение оставлена со скидкой'
    elif retained:
        rule = 'Кампания-исключение: скидка сохранена в налоговой базе'
    elif exception_keys and missing_campaign_name:
        rule = 'Название кампании не задано: применено правило добавления скидки'
    elif exception_keys:
        rule = 'Кампания не в списке исключений: скидка добавлена в налоговую базу'
    else:
        rule = 'Скидка добавлена к налоговой базе'

    dates = []
    daily_rates = stay.get('dailyRates') or []
    if len(discounts) == 1:
        item = discounts[0]
        percent = _number(item.get('percent'))
        discount_amount = _number(item.get('amount'))
        kind = str(item.get('discountType') or '').casefold()
        if (kind == 'percentage' and percent is not None
                and 0 < percent < 100 and discount_amount is not None
                and discount_amount > 0 and isinstance(daily_rates, list)):
            by_price = defaultdict(list)
            complete = True
            for rate in daily_rates:
                if not isinstance(rate, dict):
                    complete = False
                    break
                price = _number(rate.get('priceBeforeTax'))
                date = str(rate.get('date') or '')[:10]
                if price is None or len(date) != 10:
                    complete = False
                    break
                by_price[round(price, 2)].append(date)
            if complete and by_price:
                candidates = []
                for net_price, rate_dates in by_price.items():
                    undiscounted = net_price / (1 - percent / 100)
                    per_night = round(undiscounted - net_price, 2)
                    explained = round(per_night * len(rate_dates), 2)
                    tolerance = max(0.02, 0.01 * len(rate_dates))
                    if abs(explained - discount_amount) <= tolerance:
                        candidates.append(rate_dates)
                if len(candidates) == 1:
                    dates = sorted(set(candidates[0]))

    date_text = ', '.join(dates) if dates else (
        'Не распределена по датам' if amount > 0 else '')
    return {
        'amount': amount,
        'percent': '; '.join(percent_labels),
        'dates': date_text,
        'campaigns': '; '.join(campaign_names),
        'added': added,
        'retained': retained,
        'rule': rule,
        'missing_campaign_name': missing_campaign_name,
    }


def _daily_rate_audit(stay):
    """Сверяет сумму ночных тарифов с итогом проживания, не меняя его."""
    daily_rates = stay.get('dailyRates') or []
    if not isinstance(daily_rates, list) or not daily_rates:
        return None, 'Нет dailyRates'

    prices = []
    for rate in daily_rates:
        if not isinstance(rate, dict):
            return None, 'Неполные dailyRates'
        price = _number(rate.get('priceBeforeTax'))
        if price is None:
            return None, 'Неполные dailyRates'
        prices.append(price)
    daily_sum = round(sum(prices), 2)

    total = stay.get('total') or {}
    target_before_tax = _number(total.get('priceBeforeTax'))
    target_after_tax = _number(total.get('priceAfterTax'))
    # dailyRates[].priceBeforeTax — сравниваем с тем же налоговым базисом;
    # priceAfterTax используем только если API не прислал priceBeforeTax.
    target = (target_before_tax if target_before_tax is not None
              else target_after_tax)
    if target is None:
        return daily_sum, 'Нет итога для сверки'
    if abs(daily_sum - target) <= 0.02:
        return daily_sum, 'Сошлось по ночным тарифам'

    services = sum(_money(service.get('total'))
                   for service in stay.get('services') or []
                   if isinstance(service, dict))
    extras = stay.get('extraStayCharges') or {}
    if not isinstance(extras, dict):
        extras = {}
    extras = sum(_money(extras.get(key))
                 for key in ('earlyArrival', 'lateDeparture'))
    with_services = round(daily_sum + services + extras, 2)
    if ((services or extras)
            and abs(with_services - target) <= 0.02):
        return daily_sum, 'Сошлось с услугами'

    compared_sum = with_services if services or extras else daily_sum
    difference = round(compared_sum - target, 2)
    return daily_sum, f'Расхождение {difference:+.2f} ₽'


def _parse_day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value or '')[:10])
    except (TypeError, ValueError):
        return None


def _stay_period(stay_dates, qs, qe):
    """Return stay-night and selected-period bounds; checkout is exclusive."""
    stay_dates = stay_dates if isinstance(stay_dates, dict) else {}
    arrival = _parse_day(stay_dates.get('arrivalDateTime'))
    departure = _parse_day(stay_dates.get('departureDateTime'))
    period_start, period_end = _parse_day(qs), _parse_day(qe)
    if (not arrival or not departure or not period_start or not period_end
            or departure <= arrival or period_end <= period_start):
        return None
    nights = (departure - arrival).days
    selected_start = max(arrival, period_start)
    selected_end = min(departure, period_end)
    selected_nights = max(0, (selected_end - selected_start).days)
    if selected_nights <= 0:
        return None
    return {
        'arrival': arrival, 'departure': departure, 'nights': nights,
        'start': selected_start, 'end': selected_end,
        'period_nights': selected_nights,
        'full': selected_nights == nights,
    }


def _stay_overlaps(stay_dates, qs, qe):
    return _stay_period(stay_dates, qs, qe) is not None


def _complete_daily_rate_map(stay, arrival, departure, expected_total):
    """Use complete, reconciling nightly rates for a partial-period split."""
    rates = stay.get('dailyRates') or []
    if not isinstance(rates, list) or not rates:
        return None, 'Нет dailyRates'
    expected_dates = [arrival + timedelta(days=offset)
                      for offset in range((departure - arrival).days)]
    if len(rates) != len(expected_dates):
        return None, 'dailyRates неполные: число тарифов не равно числу ночей'

    by_date = {}
    for rate in rates:
        if not isinstance(rate, dict):
            return None, 'dailyRates неполные: некорректная строка тарифа'
        rate_date = _parse_day(rate.get('date'))
        if rate_date is None or rate_date in by_date:
            return None, 'dailyRates неполные: отсутствует или повторяется дата'
        by_date[rate_date] = rate
    if set(by_date) != set(expected_dates):
        return None, 'dailyRates не покрывают каждую ночь проживания'

    expected_total = _number(expected_total)
    if expected_total is None:
        return None, 'Не удалось определить полную стоимость проживания'
    # Prefer after-tax values when they reconcile to the unchanged stay total;
    # otherwise accept before-tax values only when they reconcile exactly too.
    for field in ('priceAfterTax', 'priceBeforeTax'):
        values = {}
        for rate_date, rate in by_date.items():
            price = _number(rate.get(field))
            if price is None:
                break
            values[rate_date] = price
        if len(values) == len(expected_dates):
            total = round(sum(values.values()), 2)
            if abs(total - expected_total) <= 0.02:
                return values, ''
    return None, 'Сумма dailyRates не совпадает со стоимостью проживания'


_SERVICE_DATE_KEYS = {
    'date', 'servicedate', 'servicedatetime', 'dateofservice',
    'operationdate', 'startdate', 'startdatetime', 'datetime',
    'performedat', 'servicedatefrom',
}


def _service_day(service):
    """Return a service date only when the API supplied a date field."""
    if not isinstance(service, dict):
        return None
    for key, value in service.items():
        normalized = re.sub(r'[^a-z]', '', str(key).casefold())
        if normalized in _SERVICE_DATE_KEYS:
            parsed = _parse_day(value)
            if parsed is not None:
                return parsed
    for value in service.values():
        if isinstance(value, dict):
            parsed = _service_day(value)
            if parsed is not None:
                return parsed
    return None


def _services_in_period(services, start, end):
    selected, undated = [], 0
    for service in services or []:
        if not isinstance(service, dict):
            continue
        service_date = _service_day(service)
        if service_date is None:
            if _money(service.get('total')) > 0:
                undated += 1
        elif start <= service_date < end:
            selected.append(service)
    return selected, undated


def _discount_for_period(discount, stay_period):
    """Allocate only discounts whose affected nights are identified."""
    if stay_period['full']:
        return {
            'amount': discount['amount'], 'added': discount['added'],
            'retained': discount['retained'], 'unallocated': False,
            'rule': discount['rule'],
        }
    if discount['amount'] <= 0:
        return {'amount': 0.0, 'added': 0.0, 'retained': 0.0,
                'unallocated': False, 'rule': discount['rule']}
    dates = []
    if discount.get('dates') and discount['dates'] != 'Не распределена по датам':
        dates = [parsed for parsed in
                 (_parse_day(value) for value in discount['dates'].split(', '))
                 if parsed is not None]
    if not dates:
        return {
            'amount': None, 'added': 0.0, 'retained': 0.0,
            'unallocated': True,
            'rule': ('Частичный период: дата агрегированной скидки не '
                     'определена; скидка не добавлена к базе'),
        }
    selected = [day for day in dates
                if stay_period['start'] <= day < stay_period['end']]
    if not selected:
        return {'amount': 0.0, 'added': 0.0, 'retained': 0.0,
                'unallocated': False,
                'rule': 'Ночи со скидкой находятся вне выбранного периода'}
    share = len(selected) / len(set(dates))
    if len(selected) == len(set(dates)):
        amount = discount['amount']
        added = discount['added']
        retained = discount['retained']
    else:
        amount = round(discount['amount'] * share, 2)
        added = round(discount['added'] * share, 2)
        retained = round(discount['retained'] * share, 2)
    return {
        'amount': amount, 'added': added, 'retained': retained,
        'unallocated': False,
        'rule': ('Ночи скидки определены по dailyRates; в базу добавлена '
                 'только часть выбранного периода'
                 if added else
                 'Кампания-исключение: скидка сохранена по подтверждённым ночам'),
    }


def _parse_dt(s):
    try:
        return datetime.strptime(s[:16], '%Y-%m-%dT%H:%M')
    except (TypeError, ValueError):
        return None


def _guest_name(g):
    return ' '.join(x for x in (g.get('firstName'), g.get('lastName'))
                    if x).strip()


def _room_number(value):
    """Номер из displayName; сохраняет буквенные номера, не roomId."""
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.casefold() in ('none', 'null', 'nan'):
        return None
    try:
        number = float(text)
    except ValueError:
        return text
    if number.is_integer():
        # Room labels are display values; don't strip meaningful leading zeroes.
        return text if re.fullmatch(r'0\d+', text) else int(number)
    return text


def _building_number(building, ordinal):
    """Нормализует номер корпуса, используя имя API или порядок зданий.

    PMS хранит корпуса как произвольные названия. Если в названии есть
    явный номер («Корпус 2»), используем его; для имён вроде «Main building»
    остаётся порядковый номер в accommodationInventory.buildings.
    """
    name = str((building or {}).get('name') or '').strip()
    patterns = (
        r'(?:корпус|корп\.?|здание|строение|building|block)\D{0,8}(\d+)',
        r'(\d+)\s*(?:корпус|корп\.?|здание|строение|building|block)',
    )
    for pattern in patterns:
        match = re.search(pattern, name, flags=re.IGNORECASE)
        if match:
            try:
                return int(match.group(1))
            except ValueError:
                pass
    if re.fullmatch(r'\d+', name):
        return int(name)
    return ordinal


def _room_id_key(value):
    if isinstance(value, dict):
        nested_id = value.get('id')
        if nested_id is None or (isinstance(nested_id, str)
                                 and not nested_id.strip()):
            nested_id = value.get('roomId')
        value = nested_id
    return '' if value is None else str(value).strip().casefold()


def _room_label_key(value):
    if value is None:
        return ''
    text = re.sub(r'\s+', ' ', str(value).strip()).casefold()
    return '' if text in ('', 'none', 'null', 'nan') else text


def _room_reference_candidates(value):
    """Return exact API IDs and display labels carried by a room reference."""
    ids, labels = [], []
    if isinstance(value, dict):
        for field in ('id', 'roomId'):
            key = _room_id_key(value.get(field))
            if key and key not in ids:
                ids.append(key)
        for field in ('displayName', 'roomNumber', 'number', 'name'):
            key = _room_label_key(value.get(field))
            if key and key not in labels:
                labels.append(key)
    else:
        key = _room_id_key(value)
        if key:
            ids.append(key)
            label = _room_label_key(value)
            if label:
                labels.append(label)
    return ids, labels


def _match_room_assignment(room_map, room_labels, raw_room_id, raw_room):
    id_candidates, label_candidates = [], []
    for source, value in (('roomId', raw_room_id), ('room', raw_room)):
        ids, labels = _room_reference_candidates(value)
        id_candidates.extend((key, source) for key in ids
                             if all(key != existing for existing, _ in id_candidates))
        label_candidates.extend((key, source) for key in labels
                                 if all(key != existing
                                        for existing, _ in label_candidates))
    for key, source in id_candidates:
        if key in room_map:
            return room_map[key], source, key, 'id'
    for key, source in label_candidates:
        if key in room_labels:
            return room_labels[key], f'{source}:displayName', key, 'label'
    if id_candidates:
        return None, id_candidates[0][1], id_candidates[0][0], 'id-unmatched'
    if label_candidates:
        return None, label_candidates[0][1], label_candidates[0][0], 'label-unmatched'
    return None, '', '', ''


def _accommodation_room_maps(response):
    """Return (ID map, unique display-name map) for supported PMS payloads.

    ``obtain-accommodation-inventory`` nests rooms under buildings/floors,
    whereas PropertyRoom ``/rooms`` returns a paginated top-level ``rooms``
    array. Both use room ``id``/``displayName``; the second form may be joined
    to a building through ``floorId`` when the building inventory is present.
    """
    if not isinstance(response, dict):
        return {}, {}
    inventory = response.get('accommodationInventory')
    if not isinstance(inventory, dict):
        inventory = response
    buildings = inventory.get('buildings')
    if not isinstance(buildings, list):
        buildings = []

    room_map = {}
    labels_by_key = {}
    ambiguous_labels = set()
    floor_buildings = {}

    def add_room(room, corpus=None, building_name=''):
        if not isinstance(room, dict):
            return
        room_id = _room_id_key(room.get('id'))
        if not room_id:
            room_id = _room_id_key(room.get('roomId'))
        if not room_id:
            return
        display_name = next((room.get(field) for field in
                             ('displayName', 'roomNumber', 'number', 'name')
                             if room.get(field) is not None
                             and str(room.get(field)).strip()), None)
        assignment = {
            'room': _room_number(display_name),
            'corpus': corpus,
            'building_name': building_name,
        }
        room_map.setdefault(room_id, assignment)
        label_key = _room_label_key(display_name)
        if label_key:
            previous = labels_by_key.get(label_key)
            if previous is None:
                labels_by_key[label_key] = assignment
            elif previous != assignment:
                ambiguous_labels.add(label_key)

    for building_index, building in enumerate(buildings, 1):
        if not isinstance(building, dict):
            continue
        corpus = _building_number(building, building_index)
        building_name = str(building.get('name') or '').strip()
        rooms = list(building.get('rooms') or [])
        for floor in building.get('floors') or []:
            if not isinstance(floor, dict):
                continue
            floor_id = _room_id_key(floor.get('id'))
            if floor_id:
                floor_buildings[floor_id] = (corpus, building_name)
            rooms.extend(floor.get('rooms') or [])
        for room in rooms:
            add_room(room, corpus, building_name)

    flat_rooms = response.get('propertyRooms')
    if not isinstance(flat_rooms, list):
        flat_rooms = inventory.get('rooms')
    if not isinstance(flat_rooms, list):
        flat_rooms = response.get('rooms')
    if isinstance(flat_rooms, list):
        buildings_by_id = {
            _room_id_key(building.get('id')): (
                _building_number(building, index),
                str(building.get('name') or '').strip())
            for index, building in enumerate(buildings, 1)
            if isinstance(building, dict) and _room_id_key(building.get('id'))
        }
        for room in flat_rooms:
            if not isinstance(room, dict):
                continue
            parent = buildings_by_id.get(_room_id_key(room.get('buildingId')))
            if parent is None:
                parent = floor_buildings.get(_room_id_key(room.get('floorId')))
            corpus, building_name = parent if parent is not None else (None, '')
            add_room(room, corpus, building_name)

    for key in ambiguous_labels:
        labels_by_key.pop(key, None)
    return room_map, labels_by_key


def accommodation_room_map(response):
    """TravelLine room ID -> physical room number/building (when available)."""
    return _accommodation_room_maps(response)[0]


def _service_amount_text(amount):
    return f'{amount:,.2f}'.replace(',', ' ').replace('.', ',')


def _service_name(service, fallback=''):
    if not isinstance(service, dict):
        return str(service or fallback or '').strip()
    name = next((str(service.get(key)).strip() for key in
                 ('name', 'displayName', 'serviceName', 'description',
                  'pmsSideCode', 'code', 'serviceCode')
                 if service.get(key)), '')
    return name or fallback or 'Дополнительная услуга'


def _service_description(service, fallback=''):
    """Читабельное название допуслуги с ценой, если API её передал."""
    name = _service_name(service, fallback)
    amount = _money(service.get('total')) if isinstance(service, dict) else 0.0
    if amount > 0:
        return f'{name} — {_service_amount_text(amount)} ₽'
    return name


def _is_extension_service(service):
    name = _service_name(service)
    normalized = ' '.join(name.casefold().replace('ё', 'е').split())
    return normalized == 'продление номера'


def _extension_amount(services):
    return sum(_money(service.get('total')) for service in services or []
               if isinstance(service, dict) and _is_extension_service(service))


def _service_descriptions(services):
    out = []
    for service in services or []:
        if not isinstance(service, dict):
            continue
        name = _service_description(service)
        if name:
            out.append(name)
    return out


def channel_label(source):
    """Источник брони человекочитаемо (идёт в колонку «Источник»)."""
    source = source or {}
    typ, code = source.get('type'), source.get('code')
    if typ == 'BookingEngine':
        return 'TravelLine сайт'
    if typ == 'PMS':
        return 'TravelLine ресепшен'
    if typ == 'Channel' and code:
        return f'TravelLine {code}'
    if typ or code:
        return f'TravelLine {typ or "?"}:{code or "?"}'
    return 'TravelLine ?'


def bookings_to_rows(details, qs, qe, accommodation_inventory=None,
                     campaigns_with_discount=None):
    """Детали броней -> (строки Excel, отчёт).

    ``qs``/``qe`` задают выбранные ночи полуоткрытым интервалом
    ``[qs, qe)``; дата выезда не считается ночью. Каждый ``roomStay`` —
    отдельная строка с собственными ценой, номером и налоговой базой.
    Если выбран лишь фрагмент проживания, сумма проживания берётся только
    из полного набора сверяющихся ``dailyRates``. Без него строка пропускается
    и попадает в предупреждения, без пропорциональной оценки.

    Исходные COL_PRICE и COL_TOTAL остаются API-значениями для комиссии и
    аудита. Периодная сумма и налоговая база хранятся отдельно.
    """
    period_start, period_end = _parse_day(qs), _parse_day(qe)
    room_map, room_labels = _accommodation_room_maps(accommodation_inventory)
    has_room_inventory = accommodation_inventory is not None
    rows, votes = [], {'A': 0, 'B': 0, '?': 0}
    rep = {
        'total': len(details), 'in_quarter': 0, 'cancelled': 0,
        'rows': 0, 'mismatch': [], 'channels': {},
        'period': {
            'start': qs, 'end_exclusive': qe,
            'end_inclusive': ((period_end - timedelta(days=1)).isoformat()
                              if period_end else ''),
        },
        'period_warnings': [],
        'period_allocation': {
            'partial_stays': 0, 'allocated_partial_stays': 0,
            'skipped_partial_stays': 0, 'undated_services': 0,
            'discounts_without_night_attribution': 0,
        },
        'discount_audit': {
            'stays_with_discount': 0, 'discount_amount': 0.0,
            'discount_base_added': 0.0,
            'discount_exception_retained': 0.0,
            'discount_exception_stays': 0,
            'discounts_missing_campaign_name': 0,
            'nights_with_discount_identified': 0,
            'discounts_without_night_attribution': 0,
            'daily_rates_checked': 0, 'daily_rates_matched': 0,
            'daily_rates_mismatched': 0, 'daily_rates_missing': 0,
        },
    }
    if has_room_inventory:
        rep['room_assignments'] = {
            'inventory_rooms': len(room_map), 'assigned': 0, 'unassigned': 0,
            'room_id_present': 0, 'room_id_missing': 0,
            'room_id_not_in_inventory': 0, 'inventory_no_display': 0,
            'assigned_by_room_id': 0, 'assigned_by_display_name': 0,
            'unknown_room_id_samples': [],
        }

    def add_period_warning(booking, stay, stay_period, kind, reason):
        stay_dates = stay.get('stayDates') or {}
        rep['period_warnings'].append({
            'number': booking.get('number') or '',
            'pmsRoomStayId': stay.get('pmsRoomStayId') or '',
            'arrival': stay_dates.get('arrivalDateTime') or '',
            'departure': stay_dates.get('departureDateTime') or '',
            'period_nights': stay_period['period_nights'],
            'stay_nights': stay_period['nights'],
            'type': kind,
            'reason': reason,
        })

    for b in details or []:
        if b.get('status') == 'Cancelled':
            rep['cancelled'] += 1
            continue
        stays = b.get('roomStays') or []
        room = svc_stay = svc_book = extra = 0.0
        for rs in stays:
            room += _money(rs.get('total'))
            for service in rs.get('services') or []:
                svc_stay += _money(service.get('total'))
            extra_charges = rs.get('extraStayCharges') or {}
            for key in ('earlyArrival', 'lateDeparture'):
                charge = extra_charges.get(key) or {}
                extra += _money(charge.get('total'))
        for service in b.get('services') or []:
            svc_book += _money(service.get('total'))
        total = _money(b.get('total'))
        paid_extras = svc_stay + svc_book + extra
        if paid_extras < 0.01:
            mode = 'A'
        elif abs(total - (room + paid_extras)) < 0.02:
            mode = 'A'
            votes['A'] += 1
        elif abs(total - room) < 0.02:
            mode = 'B'
            votes['B'] += 1
        else:
            mode = '?'
            votes['?'] += 1
            rep['mismatch'].append(
                {'number': b.get('number'), 'room': round(room, 2),
                 'svc': round(svc_stay + svc_book, 2),
                 'extra': round(extra, 2), 'total': round(total, 2)})

        hit = []
        for rs in stays:
            stay_period = _stay_period(rs.get('stayDates') or {}, qs, qe)
            if stay_period:
                hit.append((rs, stay_period))
        if not hit:
            continue
        rep['in_quarter'] += 1
        src = channel_label(b.get('source'))
        channel = rep['channels'].setdefault(src, [0, 0.0, 0.0])
        channel[0] += 1
        first = True

        for rs, stay_period in hit:
            stay_dates = rs.get('stayDates') or {}
            dt_in = _parse_dt(stay_dates.get('arrivalDateTime') or '')
            dt_out = _parse_dt(stay_dates.get('departureDateTime') or '')
            nights = stay_period['period_nights']
            stay_room = _money(rs.get('total'))
            stay_services = rs.get('services') or []
            stay_svc = sum(_money(service.get('total'))
                           for service in stay_services)
            discount = _discount_details(rs, campaigns_with_discount)
            daily_rate_sum, daily_rate_check = _daily_rate_audit(rs)
            discount_stats = rep['discount_audit']
            if daily_rate_sum is None or daily_rate_check.startswith(
                    ('Нет', 'Неполные')):
                discount_stats['daily_rates_missing'] += 1
            else:
                discount_stats['daily_rates_checked'] += 1
                if daily_rate_check.startswith('Сошлось'):
                    discount_stats['daily_rates_matched'] += 1
                else:
                    discount_stats['daily_rates_mismatched'] += 1

            full_row_svc = stay_svc + (svc_book + extra if first else 0.0)
            if mode == 'B':
                full_base = max(stay_room - full_row_svc, 0.0)
            else:
                full_base = stay_room

            period_status = 'Полное проживание в выбранном периоде'
            period_price = full_base
            row_svc = full_row_svc
            row_svc_extension = _extension_amount(stay_services)
            row_svc_names = _service_descriptions(stay_services)
            if first:
                booking_services = b.get('services') or []
                row_svc_extension += _extension_amount(booking_services)
                row_svc_names.extend(_service_descriptions(booking_services))
                for extra_stay in stays:
                    charges = extra_stay.get('extraStayCharges') or {}
                    for key, label in (('earlyArrival', 'Ранний заезд'),
                                       ('lateDeparture', 'Поздний выезд')):
                        charge = charges.get(key) or {}
                        if _is_extension_service(charge):
                            row_svc_extension += _money(charge.get('total'))
                        if _money(charge.get('total')) > 0:
                            row_svc_names.append(
                                _service_description(charge, label))

            period_discount = _discount_for_period(discount, stay_period)

            if not stay_period['full']:
                rep['period_allocation']['partial_stays'] += 1
                rate_map, rate_error = _complete_daily_rate_map(
                    rs, stay_period['arrival'], stay_period['departure'],
                    full_base)
                if rate_map is None:
                    rep['period_allocation']['skipped_partial_stays'] += 1
                    add_period_warning(
                        b, rs, stay_period, 'проживание',
                        f'Строка пропущена: {rate_error}')
                    continue
                period_price = round(sum(
                    price for day, price in rate_map.items()
                    if stay_period['start'] <= day < stay_period['end']), 2)
                rep['period_allocation']['allocated_partial_stays'] += 1
                period_status = 'Частичное проживание распределено по dailyRates'

                selected_stay_services, stay_undated = _services_in_period(
                    stay_services, period_start, period_end)
                selected_booking_services, booking_undated = [], 0
                if first:
                    selected_booking_services, booking_undated = (
                        _services_in_period(b.get('services') or [],
                                            period_start, period_end))
                selected_charges, extra_undated = [], 0
                charges = rs.get('extraStayCharges') or {}
                for key, label in (('earlyArrival', 'Ранний заезд'),
                                   ('lateDeparture', 'Поздний выезд')):
                    charge = charges.get(key) or {}
                    selected, undated = _services_in_period(
                        [charge] if charge else [], period_start, period_end)
                    for service in selected:
                        selected_charges.append((service, label))
                    extra_undated += undated

                selected_services = (selected_stay_services
                                     + selected_booking_services
                                     + [service for service, _ in selected_charges])
                row_svc = sum(_money(service.get('total'))
                              for service in selected_services)
                row_svc_extension = _extension_amount(selected_stay_services)
                row_svc_names = _service_descriptions(selected_stay_services)
                if first:
                    row_svc_extension += _extension_amount(
                        selected_booking_services)
                    row_svc_names.extend(
                        _service_descriptions(selected_booking_services))
                for service, label in selected_charges:
                    if _is_extension_service(service):
                        row_svc_extension += _money(service.get('total'))
                    row_svc_names.append(_service_description(service, label))

                undated_services = stay_undated + booking_undated + extra_undated
                if undated_services:
                    rep['period_allocation']['undated_services'] += (
                        undated_services)
                    period_status += '; допуслуги без даты исключены'
                    add_period_warning(
                        b, rs, stay_period, 'допуслуга',
                        f'{undated_services} допуслуг без даты не включены '
                        'в сумму частичного периода')

            if discount['amount'] > 0:
                discount_stats['stays_with_discount'] += 1
                if period_discount['amount'] is None:
                    discount_stats[
                        'discounts_without_night_attribution'] += 1
                    rep['period_allocation'][
                        'discounts_without_night_attribution'] += 1
                else:
                    discount_stats['discount_amount'] += (
                        period_discount['amount'])
                    if discount.get('dates') == 'Не распределена по датам':
                        discount_stats[
                            'discounts_without_night_attribution'] += 1
                discount_stats['discount_base_added'] += period_discount['added']
                discount_stats['discount_exception_retained'] += (
                    period_discount['retained'])
                if period_discount['retained'] > 0:
                    discount_stats['discount_exception_stays'] += 1
                if discount['missing_campaign_name']:
                    discount_stats['discounts_missing_campaign_name'] += 1
                if discount.get('dates') and discount.get('dates') != (
                        'Не распределена по датам'):
                    known_dates = [day for day in
                                   (_parse_day(value) for value in
                                    discount['dates'].split(', '))
                                   if day is not None]
                    selected_dates = [day for day in known_dates
                                      if stay_period['start'] <= day
                                      < stay_period['end']]
                    discount_stats['nights_with_discount_identified'] += len(
                        selected_dates)

            if period_discount['unallocated']:
                period_status += '; скидка без даты не восстановлена'
                add_period_warning(
                    b, rs, stay_period, 'скидка',
                    'Агрегированная скидка не добавлена к налоговой базе '
                    'частичного периода: API не указал затронутую ночь')
            tax_base = round(period_price + period_discount['added'], 2)

            guests = ', '.join(_guest_name(guest)
                               for guest in rs.get('guests') or []
                               if _guest_name(guest))
            room_type = rs.get('roomType') or {}
            rate_plans = rs.get('ratePlans') or []
            assignment, _, room_key, room_match_kind = (
                _match_room_assignment(
                    room_map, room_labels, rs.get('roomId'), rs.get('room')))
            room_number = assignment.get('room') if assignment else None
            corpus = assignment.get('corpus') if assignment else None
            building_name = (assignment.get('building_name', '')
                             if assignment else '')
            if has_room_inventory:
                room_stats = rep['room_assignments']
                room_stats['room_id_present' if room_key
                           else 'room_id_missing'] += 1
                if room_key and assignment is None:
                    room_stats['room_id_not_in_inventory'] += 1
                    samples = room_stats['unknown_room_id_samples']
                    if room_key not in samples and len(samples) < 20:
                        samples.append(room_key)
                elif assignment is not None and room_number is None:
                    room_stats['inventory_no_display'] += 1
                elif assignment is not None and room_match_kind == 'id':
                    room_stats['assigned_by_room_id'] += 1
                elif assignment is not None:
                    room_stats['assigned_by_display_name'] += 1
                bucket = 'assigned' if room_number is not None else 'unassigned'
                room_stats[bucket] += 1

            rows.append({
                COL_GUEST: guests,
                COL_PRICE: round(full_base, 2),
                COL_TL_PERIOD_PRICE: round(period_price, 2),
                COL_TL_PERIOD_STATUS: period_status,
                COL_IN: dt_in, COL_OUT: dt_out,
                COL_SOURCE: src,
                COL_CAT: room_type.get('name') or '',
                COL_BRON: b.get('number') or '',
                COL_NIGHTS: nights,
                COL_ROOM: room_number,
                COL_CORPUS: corpus if room_number is not None else None,
                COL_BUILDING_NAME: building_name if room_number is not None else '',
                COL_TARIFF: ', '.join(plan.get('name') or ''
                                      for plan in rate_plans),
                COL_SVC: round(row_svc, 2),
                COL_SVC_NAME: '; '.join(row_svc_names),
                COL_SVC_EXTENSION: round(row_svc_extension, 2),
                # Keep the actual group-booking total from the API unchanged.
                COL_TOTAL: round(total, 2),
                COL_TL_DISCOUNT: discount['amount'],
                COL_TL_DISCOUNT_PERCENT: discount['percent'],
                COL_TL_DISCOUNT_DATES: discount['dates'],
                COL_TL_DAILY_SUM: daily_rate_sum,
                COL_TL_DAILY_CHECK: daily_rate_check,
                COL_TL_DISCOUNT_CAMPAIGN: discount['campaigns'],
                COL_TL_DISCOUNT_ADDED: period_discount['added'],
                COL_TL_DISCOUNT_RETAINED: period_discount['retained'],
                COL_TL_DISCOUNT_RULE: period_discount['rule'],
                COL_TL_TAX_BASE: tax_base,
            })
            first = False
        channel[1] += room
        channel[2] += paid_extras

    rep['rows'] = len(rows)
    rep['votes'] = votes
    for key in ('discount_amount', 'discount_base_added',
                'discount_exception_retained'):
        rep['discount_audit'][key] = round(rep['discount_audit'][key], 2)
    return rows, rep

def write_data_workbook(rows, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Брони'
    cols = [COL_GUEST, COL_PRICE, COL_IN, COL_OUT, COL_SOURCE, COL_CAT,
            COL_BRON, COL_NIGHTS, COL_TARIFF, COL_SVC, COL_SVC_NAME,
            COL_SVC_EXTENSION, COL_TOTAL, COL_ROOM, COL_CORPUS,
            COL_TL_DISCOUNT, COL_TL_DISCOUNT_PERCENT,
            COL_TL_DISCOUNT_DATES, COL_TL_DAILY_SUM, COL_TL_DAILY_CHECK,
            COL_TL_DISCOUNT_CAMPAIGN, COL_TL_DISCOUNT_ADDED,
            COL_TL_DISCOUNT_RETAINED, COL_TL_DISCOUNT_RULE,
            COL_TL_TAX_BASE, COL_TL_PERIOD_PRICE, COL_TL_PERIOD_STATUS]
    ws.append(cols)
    for r in rows:
        values = []
        for column in cols:
            value = r.get(column)
            if column == COL_ROOM and isinstance(value, str):
                # pandas may coerce text labels like "001" to the integer 1
                # when it later reads the workbook; preserve the exact API label.
                value = ROOM_LABEL_TEXT_PREFIX + value
            values.append(value)
        ws.append(values)
    for col in ('C', 'D'):
        for cell in ws[col][1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = 'DD.MM.YYYY HH:MM'
    for col in ('P', 'S', 'V', 'W', 'Y', 'Z'):
        for cell in ws[col][1:]:
            cell.number_format = '#,##0.00'
    # Room assignment is consumed by the processor but kept out of the way
    # in the source workbook.
    for column in ('N', 'O'):
        ws.column_dimensions[column].hidden = True
    wb.save(path)


def write_rooms_workbook(rows, path):
    """Пишет доступные из API имя гостя, номер и корпус для проверки.

    Неприсвоенные и не найденные в номерном фонде roomId остаются пустыми.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Номера'
    ws.append(['ФИО', COL_ROOM, COL_CORPUS, COL_BUILDING_NAME, COL_BRON,
               COL_IN, COL_OUT])
    for row in rows or []:
        guests = [name.strip() for name in
                  re.split(r'[,;\n]+', str(row.get(COL_GUEST) or ''))
                  if name.strip()]
        if not guests:
            guests = ['']
        for guest in guests:
            ws.append([guest, row.get(COL_ROOM), row.get(COL_CORPUS),
                       row.get(COL_BUILDING_NAME) or '', row.get(COL_BRON),
                       row.get(COL_IN), row.get(COL_OUT)])
    for column in ('F', 'G'):
        for cell in ws[column][1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = 'DD.MM.YYYY HH:MM'
    wb.save(path)


def tl_inspection_path(tmpdir, code, year, quarter):
    """Путь к локальной книге для просмотра выгрузки/диагностики TravelLine."""
    return f'{tmpdir}/TL_{code}_{year}{quarter}_inspection.xlsx'


def _inspection_json(value, limit=32000):
    try:
        text = json.dumps(value, ensure_ascii=False, default=str)
    except (TypeError, ValueError):
        text = str(value)
    if len(text) > limit:
        return text[:limit - 24] + ' … [обрезано]'
    return text


def _inspection_cell(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list, tuple)):
        return _inspection_json(value)
    return value


def _inspection_discount_summary(stay):
    """Отдельно показывает названия скидок/кампаний, не раскрывая ФИО."""
    total = stay.get('total') or {}
    discounts = total.get('discounts') or []
    if isinstance(discounts, dict):
        discounts = [discounts]
    return [{
        'name': item.get('name') or '',
        'discountType': item.get('discountType') or '',
        'percent': item.get('percent'),
        'amount': item.get('amount'),
    } for item in discounts if isinstance(item, dict)]


def _inspection_redact(value):
    """Сохраняет техполя для отладки, убирая ФИО/контакты из сырого JSON."""
    private_parts = ('guest', 'name', 'email', 'phone', 'passport',
                     'document', 'contact')
    if isinstance(value, dict):
        return {
            key: ('<скрыто>' if any(part in str(key).casefold()
                                    for part in private_parts)
                  else _inspection_redact(item))
            for key, item in value.items()
        }
    if isinstance(value, (list, tuple)):
        return [_inspection_redact(item) for item in value]
    return value


def write_api_inspection_workbook(rows, report, path, details=None,
                                  accommodation_inventory=None):
    """Создаёт открытую пользователю книгу с экспортом и сырой диагностикой.

    В кэше TravelLine хранятся только подготовленные строки. Если ``details``
    отсутствует, книга честно помечает, что исходные API поля доступны только
    после выбора «Обновить из TravelLine».
    """
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = 'Сводка'
    stats = (report or {}).get('room_assignments') or {}
    discount_stats = (report or {}).get('discount_audit') or {}
    period_info = (report or {}).get('period') or {}
    period_allocation = (report or {}).get('period_allocation') or {}
    period_warnings = (report or {}).get('period_warnings') or []

    def stat_value(key):
        return stats.get(key, '—' if details is None else 0)

    summary_rows = [
        ('Источник', 'Свежий ответ TravelLine' if details is not None
         else 'Локальный кэш (без исходных ответов API)'),
        ('Брони в ответе API', len(details) if details is not None else 'нет в кэше'),
        ('Подготовлено строк для приложения', len(rows or [])),
        ('Выбранный период ночей',
         f"{period_info.get('start', '—')} — "
         f"{period_info.get('end_inclusive', '—')} (выездная дата не ночь)"),
        ('Частичных проживаний распределено по dailyRates',
         period_allocation.get('allocated_partial_stays', 0)),
        ('Частичных проживаний пропущено без точного распределения',
         period_allocation.get('skipped_partial_stays', 0)),
        ('Допуслуг без даты не включено в частичные периоды',
         period_allocation.get('undated_services', 0)),
        ('Комнат в извлечённом API-инвентаре', stats.get('inventory_rooms', 0)),
        ('Проживаний сопоставлено с номером', stats.get('assigned', 0)),
        ('Проживаний без номера', stats.get('unassigned', 0)),
        ('Проживаний с roomId/room', stat_value('room_id_present')),
        ('Проживаний без roomId/room', stat_value('room_id_missing')),
        ('ID не найден в инвентаре', stat_value('room_id_not_in_inventory')),
        ('Комната найдена, но нет отображаемого номера',
         stat_value('inventory_no_display')),
        ('Проживаний со скидкой TravelLine',
         discount_stats.get('stays_with_discount', 0)),
        ('Общая скидка TravelLine, ₽',
         discount_stats.get('discount_amount', 0.0)),
        ('Добавлено к стоимости для налоговой базы, ₽',
         discount_stats.get('discount_base_added', 0.0)),
        ('Оставлено со скидкой по исключениям, ₽',
         discount_stats.get('discount_exception_retained', 0.0)),
        ('Проживаний с исключением кампании',
         discount_stats.get('discount_exception_stays', 0)),
        ('Скидок без названия кампании',
         discount_stats.get('discounts_missing_campaign_name', 0)),
        ('Предполагаемых ночей со скидкой (вывод по dailyRates)',
         discount_stats.get('nights_with_discount_identified', 0)),
        ('Скидок без однозначной даты',
         discount_stats.get('discounts_without_night_attribution', 0)),
        ('Сверок dailyRates совпало',
         discount_stats.get('daily_rates_matched', 0)),
        ('Сверок dailyRates с расхождением',
         discount_stats.get('daily_rates_mismatched', 0)),
        ('Сверок dailyRates без данных',
         discount_stats.get('daily_rates_missing', 0)),
    ]
    if details is None:
        summary_rows.append((
            'Важно', 'Кэш содержит только готовые строки; выберите '
            '«Обновить из TravelLine», чтобы получить исходные roomId и инвентарь.'))
    for row_index, item in enumerate(summary_rows, 1):
        summary.append(item)
        if item[0].endswith('₽'):
            summary.cell(row=row_index, column=2).number_format = '#,##0.00'
    if stats.get('unknown_room_id_samples'):
        summary.append(('Примеры ID без совпадения',
                        ', '.join(map(str, stats['unknown_room_id_samples']))))
    summary.column_dimensions['A'].width = 48
    summary.column_dimensions['B'].width = 100
    summary.freeze_panes = 'A1'

    period_sheet = wb.create_sheet('Предупреждения периода')
    period_sheet.append([
        '№ брони', 'pmsRoomStayId', 'Заезд API', 'Выезд API',
        'Ночей в проживании', 'Ночей в периоде', 'Тип', 'Причина',
    ])
    for warning in period_warnings:
        period_sheet.append([
            warning.get('number', ''), warning.get('pmsRoomStayId', ''),
            warning.get('arrival', ''), warning.get('departure', ''),
            warning.get('stay_nights', ''), warning.get('period_nights', ''),
            warning.get('type', ''), warning.get('reason', ''),
        ])
    if not period_warnings:
        period_sheet.append(['Предупреждений по распределению нет.'])
    period_sheet.freeze_panes = 'A2'
    for column, width in zip('ABCDEFGH', (26, 22, 22, 22, 18, 18, 18, 78)):
        period_sheet.column_dimensions[column].width = width

    data_sheet = wb.create_sheet('Данные')
    data_columns = [COL_GUEST, COL_PRICE, COL_IN, COL_OUT, COL_SOURCE, COL_CAT,
                    COL_BRON, COL_NIGHTS, COL_TARIFF, COL_SVC, COL_SVC_NAME,
                    COL_SVC_EXTENSION, COL_TOTAL, COL_ROOM, COL_CORPUS,
                    COL_BUILDING_NAME, COL_TL_DISCOUNT,
                    COL_TL_DISCOUNT_PERCENT, COL_TL_DISCOUNT_DATES,
                    COL_TL_DAILY_SUM, COL_TL_DAILY_CHECK,
                    COL_TL_DISCOUNT_CAMPAIGN, COL_TL_DISCOUNT_ADDED,
                    COL_TL_DISCOUNT_RETAINED, COL_TL_DISCOUNT_RULE,
                    COL_TL_TAX_BASE, COL_TL_PERIOD_PRICE,
                    COL_TL_PERIOD_STATUS]
    data_sheet.append(data_columns)
    for row in rows or []:
        data_sheet.append([row.get(column) for column in data_columns])
    for col in ('Q', 'T', 'W', 'X', 'Z', 'AA'):
        for cell in data_sheet[col][1:]:
            cell.number_format = '#,##0.00'

    rooms_sheet = wb.create_sheet('Номера')
    rooms_sheet.append(['ФИО', COL_ROOM, COL_CORPUS,
                        COL_BUILDING_NAME, COL_BRON, COL_IN, COL_OUT])
    for row in rows or []:
        guests = [name.strip() for name in
                  re.split(r'[,;\n]+', str(row.get(COL_GUEST) or ''))
                  if name.strip()]
        if not guests:
            guests = ['']
        for guest in guests:
            rooms_sheet.append([guest, row.get(COL_ROOM), row.get(COL_CORPUS),
                                row.get(COL_BUILDING_NAME) or '',
                                row.get(COL_BRON), row.get(COL_IN),
                                row.get(COL_OUT)])

    inventory_sheet = wb.create_sheet('Номерной фонд API')
    inventory_headers = [
        '№ корпуса', 'ID корпуса', 'Корпус API', 'ID этажа', 'Этаж API',
        'room.id', 'room.roomId', 'displayName', 'roomNumber', 'number', 'name',
        'Поля комнаты (JSON)',
    ]
    inventory_sheet.append(inventory_headers)
    raw_inventory = accommodation_inventory
    inventory = raw_inventory
    if isinstance(inventory, dict) and isinstance(
            inventory.get('accommodationInventory'), dict):
        inventory = inventory['accommodationInventory']
    buildings = inventory.get('buildings') if isinstance(inventory, dict) else None
    floor_buildings = {}
    building_records = {}
    if isinstance(buildings, list):
        for building_index, building in enumerate(buildings, 1):
            if not isinstance(building, dict):
                continue
            building_id = _room_id_key(building.get('id'))
            if building_id:
                building_records[building_id] = (building_index, building)
            room_entries = [(None, room) for room in (building.get('rooms') or [])]
            for floor in building.get('floors') or []:
                if isinstance(floor, dict):
                    floor_id = _room_id_key(floor.get('id'))
                    if floor_id:
                        floor_buildings[floor_id] = (building_index, building, floor)
                    room_entries.extend((floor, room)
                                        for room in (floor.get('rooms') or []))
            for floor, room in room_entries:
                if not isinstance(room, dict):
                    continue
                inventory_sheet.append([
                    _building_number(building, building_index),
                    _inspection_cell(building.get('id')),
                    building.get('name') or '',
                    _inspection_cell(floor.get('id')) if floor else '',
                    floor.get('name') or '' if floor else '',
                    _inspection_cell(room.get('id')),
                    _inspection_cell(room.get('roomId')),
                    _inspection_cell(room.get('displayName')),
                    _inspection_cell(room.get('roomNumber')),
                    _inspection_cell(room.get('number')),
                    _inspection_cell(room.get('name')),
                    _inspection_json(_inspection_redact(room)),
                ])
    flat_rooms = (raw_inventory.get('propertyRooms')
                  if isinstance(raw_inventory, dict) else None)
    if not isinstance(flat_rooms, list) and isinstance(inventory, dict):
        flat_rooms = inventory.get('rooms')
    if not isinstance(flat_rooms, list) and isinstance(raw_inventory, dict):
        flat_rooms = raw_inventory.get('rooms')
    for room in flat_rooms or []:
        if not isinstance(room, dict):
            continue
        parent = building_records.get(_room_id_key(room.get('buildingId')))
        floor = None
        if parent is None:
            parent = floor_buildings.get(_room_id_key(room.get('floorId')))
            if parent is not None:
                building_index, building, floor = parent
            else:
                building_index, building = None, {}
        else:
            building_index, building = parent
        inventory_sheet.append([
            building_index if building_index is not None else '',
            _inspection_cell(building.get('id')),
            building.get('name') or '',
            _inspection_cell(room.get('floorId')),
            _inspection_cell(floor.get('name')) if floor else '',
            _inspection_cell(room.get('id')),
            _inspection_cell(room.get('roomId')),
            _inspection_cell(room.get('displayName')),
            _inspection_cell(room.get('roomNumber')),
            _inspection_cell(room.get('number')),
            _inspection_cell(room.get('name')),
            _inspection_json(_inspection_redact(room)),
        ])
    if inventory_sheet.max_row == 1:
        inventory_sheet.append(['Нет сырого номерного фонда в этой выгрузке.'])

    stays_sheet = wb.create_sheet('Проживания API')
    stays_sheet.append([
        '№ брони', '№ проживания', 'pmsRoomStayId', 'roomId', 'room',
        'roomTypeId', 'roomType', 'Заезд API', 'Выезд API', 'Все ключи roomStay',
        'Поля размещения (JSON)', 'roomId использован из',
        'Результат сопоставления', 'Номер из инвентаря', 'Корпус',
        'Скидки API: название/тип/процент/сумма',
    ])
    room_map, room_labels = _accommodation_room_maps(accommodation_inventory)
    if details is not None:
        for booking in details or []:
            if not isinstance(booking, dict):
                continue
            booking_number = booking.get('number') or ''
            for index, stay in enumerate(booking.get('roomStays') or [], 1):
                if not isinstance(stay, dict):
                    stays_sheet.append([booking_number, index, '', '', '', '', '',
                                        '', '', '', _inspection_json(stay), '',
                                        'roomStay не является объектом', '', '', ''])
                    continue
                raw_room_id = stay.get('roomId')
                raw_room = stay.get('room')
                assignment, room_id_source, room_key, room_match_kind = (
                    _match_room_assignment(
                        room_map, room_labels, raw_room_id, raw_room))
                if not room_key:
                    match_status = 'roomId/room отсутствует'
                elif assignment is None and room_match_kind == 'label-unmatched':
                    match_status = 'Подпись комнаты не найдена в инвентаре'
                elif assignment is None:
                    match_status = 'ID не найден в инвентаре'
                elif assignment.get('room') is None:
                    match_status = 'ID найден, display name отсутствует'
                elif room_match_kind == 'label':
                    match_status = 'Сопоставлено по подписи комнаты'
                else:
                    match_status = 'Сопоставлено по ID'
                room_type = stay.get('roomType') or {}
                if isinstance(room_type, dict):
                    room_type = room_type.get('name') or room_type.get('displayName') \
                        or _inspection_json(room_type)
                stay_dates = stay.get('stayDates') or {}
                safe_stay = _inspection_redact(stay)
                stays_sheet.append([
                    booking_number, index,
                    _inspection_cell(stay.get('pmsRoomStayId')),
                    _inspection_cell(raw_room_id), _inspection_cell(raw_room),
                    _inspection_cell(stay.get('roomTypeId')),
                    _inspection_cell(room_type),
                    _inspection_cell(stay_dates.get('arrivalDateTime'))
                    if isinstance(stay_dates, dict) else '',
                    _inspection_cell(stay_dates.get('departureDateTime'))
                    if isinstance(stay_dates, dict) else '',
                    ', '.join(map(str, stay.keys())),
                    _inspection_json(safe_stay), room_id_source, match_status,
                    assignment.get('room') if assignment else None,
                    assignment.get('corpus') if assignment else None,
                    _inspection_json(_inspection_discount_summary(stay)),
                ])
    if details is None:
        stays_sheet.append([
            'Сырые проживания не сохранены в кэше. Нажмите «Обновить из TravelLine».'])
    elif stays_sheet.max_row == 1:
        stays_sheet.append(['В API-ответе нет roomStays.'])

    default_widths = {
        'Данные': [26, 16, 20, 20, 26, 28, 26, 10, 28, 14, 48, 20,
                   16, 16, 12, 28, 18, 24, 38, 16, 34, 42, 16, 16, 58, 20],
        'Номера': [32, 16, 12, 28, 28, 22, 22],
        'Номерной фонд API': [12, 18, 28, 18, 24, 28, 28, 22, 22, 18, 26, 72],
        'Проживания API': [28, 12, 28, 32, 32, 24, 32, 24, 24, 54, 72, 18, 36, 20, 14, 60],
    }
    for worksheet in wb.worksheets:
        worksheet.freeze_panes = 'A2' if worksheet.max_row > 1 else 'A1'
        if worksheet.max_row > 1 and worksheet.title != 'Сводка':
            worksheet.auto_filter.ref = worksheet.dimensions
        for index, width in enumerate(default_widths.get(worksheet.title, []), 1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    wb.save(path)


def write_rooms_stub(path):
    """Совместимая пустая заглушка для импорта отелей без номерного фонда."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Номера'
    ws.append(['ФИО', COL_ROOM])
    wb.save(path)
